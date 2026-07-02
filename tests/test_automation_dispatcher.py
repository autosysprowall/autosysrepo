from __future__ import annotations

import io
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.automation_dispatcher import (
    AutomationService,
    ControlRecord,
    GanttFileState,
    Notification,
    SharePointBackend,
    VersionArtifact,
    WorkbookMetadata,
    apply_notification_delivery_mode,
    due_tracking_kind,
    extract_workbook_metadata,
    run_system2,
    sent_notification_updates,
)
from src.gantt_versioning import (
    decide_version,
    snapshot_gantt,
    workbook_with_status,
)


NOW = datetime(2026, 7, 1, 12, 0, tzinfo=timezone.utc)


def test_sent_notification_updates_maps_assignment() -> None:
    assert sent_notification_updates(
        "AsignacionGantt",
        "2026-07-01T21:27:47Z",
    ) == {
        "CorreoAsignacionEnviado": True,
        "FechaCorreoAsignacion": "2026-07-01T21:27:47Z",
    }


def test_sent_notification_updates_ignores_unknown_kind() -> None:
    assert sent_notification_updates("Advertencia3", "2026-07-01T21:27:47Z") == {}


class FakeBackend:
    def __init__(self, metadata: WorkbookMetadata | None = None) -> None:
        self.metadata = metadata or WorkbookMetadata()
        self.patches: list[tuple[str, dict]] = []
        self.grants: list[tuple[str, str]] = []
        self.notifications: list[tuple[str, Notification]] = []
        self.notification_keys: set[tuple[str, str]] = set()
        self.permission_error: Exception | None = None
        self.notification_error: Exception | None = None
        self.metadata_error: Exception | None = None
        self.version_error: Exception | None = None
        self.versions: list[str] = []
        self.resolved_gantt_link = (
            "https://contoso.sharepoint.com/Proyectos/Proyectos%20Activos/"
            "gantt_working.xlsx"
        )

    def patch_control(self, item_id: str, updates: dict) -> None:
        self.patches.append((item_id, updates))

    def load_metadata(self, record: ControlRecord, *, prefer_gantt: bool = True) -> WorkbookMetadata:
        if self.metadata_error:
            raise self.metadata_error
        return self.metadata

    def load_gantt_file_state(
        self,
        record: ControlRecord,
    ) -> GanttFileState:
        if self.metadata_error:
            raise self.metadata_error
        return GanttFileState(
            metadata=self.metadata,
            etag=record.gantt_etag or '"etag-current"',
            modified_at=NOW,
            modified_by_email="ingeniero@example.com",
        )

    def grant_edit_access(self, record: ControlRecord, email: str) -> None:
        if self.permission_error:
            raise self.permission_error
        self.grants.append((record.item_id, email))

    def resolve_gantt_editor_link(self, record: ControlRecord) -> str:
        return self.resolved_gantt_link

    def notification_exists(self, item_id: str, kind: str) -> bool:
        return (item_id, kind.casefold()) in self.notification_keys

    def queue_notification(self, record: ControlRecord, notification: Notification) -> None:
        if self.notification_error:
            raise self.notification_error
        key = (record.item_id, notification.kind.casefold())
        if key in self.notification_keys:
            raise AssertionError(f"duplicate notification: {key}")
        self.notification_keys.add(key)
        self.notifications.append((record.item_id, notification))

    def create_initial_version(self, record: ControlRecord) -> VersionArtifact:
        if self.version_error:
            raise self.version_error
        self.versions.append(record.item_id)
        return VersionArtifact(
            identifier="version-item-id",
            web_url="https://contoso.sharepoint.com/versionados/2026-001_gantt_v1.0.xlsx",
            file_name="2026-001_gantt_v1.0.xlsx",
            created=True,
            source_etag='"etag-versioned"',
        )


def record(**changes) -> ControlRecord:
    defaults = {
        "item_id": "10",
        "project_id": "2026-001",
        "project_name": "Proyecto Prueba",
        "budget_link": "https://contoso.sharepoint.com/Presupuesto.xlsx",
        "gantt_link": "https://contoso.sharepoint.com/Proyectos/Proyectos%20Activos/gantt_working.xlsx",
        "engineer_email": "ingeniero@example.com",
        "supervisors_email": "supervisor@example.com",
        "state": "Pendiente de asignación",
    }
    defaults.update(changes)
    return ControlRecord(**defaults)


class AssignmentTests(unittest.TestCase):
    def test_valid_assignment_grants_permission_and_queues_one_email(self) -> None:
        backend = FakeBackend()
        result = AutomationService(backend, NOW).assign(record())

        self.assertTrue(result.permission_granted)
        self.assertEqual("En Progreso", result.state)
        self.assertEqual(NOW + timedelta(days=9), result.deadline)
        self.assertEqual([("10", "ingeniero@example.com")], backend.grants)
        self.assertEqual(["AsignacionGantt"], [item.kind for _, item in backend.notifications])
        notification = backend.notifications[0][1]
        self.assertEqual("ingeniero@example.com", notification.to)
        self.assertEqual("supervisor@example.com", notification.cc)
        self.assertEqual(
            "Asignación Cronograma Proyecto Proyecto Prueba",
            notification.subject,
        )
        self.assertIn("El departamento de comercial", notification.body)
        self.assertIn(record().gantt_link, notification.body)
        self.assertIn(
            f'<a href="{record().gantt_link}">',
            notification.body,
        )
        self.assertIn("Abrir Gantt con permiso de edición</a>", notification.body)
        self.assertNotIn("Guía PDF Ingenieros + Planta", notification.body)
        merged = {key: value for _, patch in backend.patches for key, value in patch.items()}
        self.assertEqual("En Progreso", merged["EstadoGantt"])
        self.assertTrue(merged["PermisoGanttOtorgado"])

    def test_identifier_only_assignment_resolves_editor_link_for_email(self) -> None:
        backend = FakeBackend()
        result = AutomationService(backend, NOW).assign(
            record(gantt_link="", gantt_identifier="drive-item-id")
        )

        self.assertEqual(backend.resolved_gantt_link, result.gantt_link)
        self.assertIn(
            backend.resolved_gantt_link,
            backend.notifications[0][1].body,
        )

    def test_missing_email_does_not_share_or_queue(self) -> None:
        backend = FakeBackend(WorkbookMetadata())
        with patch.dict(
            "os.environ",
            {"NOTIFICATION_DELIVERY_MODE": "live"},
        ):
            result = AutomationService(backend, NOW).assign(
                record(engineer_email="")
            )

        self.assertEqual("Pendiente de asignación", result.state)
        self.assertEqual([], backend.grants)
        self.assertEqual([], backend.notifications)
        self.assertTrue(
            any("vacío o inválido" in patch.get("UltimoErrorTracking", "") for _, patch in backend.patches)
        )

    def test_missing_email_queues_autosys_preview_only_in_test_mode(self) -> None:
        backend = FakeBackend(WorkbookMetadata())
        with patch.dict(
            "os.environ",
            {"NOTIFICATION_DELIVERY_MODE": "test"},
        ):
            result = AutomationService(backend, NOW).assign(
                record(engineer_email="")
            )
        self.assertEqual("Pendiente de asignación", result.state)
        self.assertEqual([], backend.grants)
        self.assertEqual(
            ["AsignacionGantt"],
            [item.kind for _, item in backend.notifications],
        )

    def test_permission_failure_does_not_mark_permission_or_queue_email(self) -> None:
        backend = FakeBackend()
        backend.permission_error = RuntimeError("Graph 403")
        AutomationService(backend, NOW).assign(record())

        self.assertEqual([], backend.notifications)
        self.assertFalse(
            any(patch.get("PermisoGanttOtorgado") is True for _, patch in backend.patches)
        )
        self.assertTrue(
            any("Graph 403" in patch.get("UltimoErrorTracking", "") for _, patch in backend.patches)
        )

    def test_rerun_does_not_duplicate_permission_or_assignment_email(self) -> None:
        backend = FakeBackend()
        backend.notification_keys.add(("10", "asignaciongantt"))
        current = record(
            state="Asignado",
            permission_granted=True,
            assignment_date=NOW,
            deadline=NOW + timedelta(days=9),
        )
        AutomationService(backend, NOW).assign(current)
        AutomationService(backend, NOW).assign(current)

        self.assertEqual([], backend.grants)
        self.assertEqual([], backend.notifications)

    def test_existing_assignment_notification_prevents_regrant_without_legacy_flag(self) -> None:
        backend = FakeBackend()
        backend.notification_keys.add(("10", "asignaciongantt"))
        current = record(
            state="Asignado",
            permission_granted=False,
            assignment_date=NOW,
            deadline=NOW + timedelta(days=9),
        )
        AutomationService(backend, NOW).assign(current)
        self.assertEqual([], backend.grants)

    def test_assignment_preserves_en_progreso_state(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En progreso",
            permission_granted=True,
            assignment_date=NOW - timedelta(days=1),
            deadline=NOW + timedelta(days=8),
        )
        result = AutomationService(backend, NOW).assign(current)
        self.assertEqual("En progreso", result.state)

    def test_existing_engineer_is_assigned_even_if_optional_metadata_fails(self) -> None:
        backend = FakeBackend()
        backend.metadata_error = RuntimeError("Datos unreadable")
        result = AutomationService(backend, NOW).assign(record(supervisors_email=""))
        self.assertEqual("En Progreso", result.state)
        self.assertEqual([("10", "ingeniero@example.com")], backend.grants)

    def test_supervisors_are_enriched_when_engineer_already_exists(self) -> None:
        backend = FakeBackend(WorkbookMetadata(supervisors_email="jefe@example.com"))
        result = AutomationService(backend, NOW).assign(record(supervisors_email=""))
        self.assertEqual("jefe@example.com", result.supervisors_email)
        self.assertTrue(
            any(patch.get("SupervisoresEmail") == "jefe@example.com" for _, patch in backend.patches)
        )

    def test_proyectos_terminados_path_is_rejected(self) -> None:
        backend = FakeBackend()
        result = AutomationService(backend, NOW).assign(
            record(gantt_link="https://contoso/Proyectos/Proyectos%20Terminados/gantt.xlsx")
        )
        self.assertEqual("Requiere revisión manual", result.state)
        self.assertEqual([], backend.grants)
        self.assertEqual([], backend.notifications)


class TrackingTests(unittest.TestCase):
    def test_day_3_queues_warning_1_once(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Asignado",
            permission_granted=True,
            assignment_email_sent=True,
            assignment_date=NOW - timedelta(days=3),
        )
        service = AutomationService(backend, NOW)
        service.track(current)
        service.track(current)
        self.assertEqual(["Advertencia1"], [item.kind for _, item in backend.notifications])
        notification = backend.notifications[0][1]
        self.assertEqual("ingeniero@example.com", notification.to)
        self.assertEqual(
            "supervisor@example.com;jaime.madrid@prowallpanama.com;"
            "enrique.correa@prowallpanama.com",
            notification.cc,
        )
        self.assertIn("Han pasado 3 días", notification.body)

    def test_day_6_queues_warning_2_with_supervisor_cc(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En progreso",
            permission_granted=True,
            assignment_email_sent=True,
            warning1_sent=True,
            assignment_date=NOW - timedelta(days=6),
        )
        AutomationService(backend, NOW).track(current)
        notification = backend.notifications[0][1]
        self.assertEqual("Advertencia2", notification.kind)
        self.assertEqual(
            "supervisor@example.com;jaime.madrid@prowallpanama.com;"
            "enrique.correa@prowallpanama.com",
            notification.cc,
        )
        self.assertIn("Han pasado 6 días", notification.body)

    def test_day_9_marks_expired_and_queues_escalation_without_third_warning(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Asignado",
            permission_granted=True,
            assignment_email_sent=True,
            warning1_sent=True,
            warning2_sent=True,
            assignment_date=NOW - timedelta(days=9),
        )
        result = AutomationService(backend, NOW).track(current)
        self.assertEqual("Vencido", result.state)
        self.assertEqual(["Vencimiento"], [item.kind for _, item in backend.notifications])
        self.assertTrue(any(patch.get("EstadoGantt") == "Vencido" for _, patch in backend.patches))
        closed = record(
            state="Vencido",
            assignment_date=NOW - timedelta(days=20),
            expiration_notified=True,
        )
        self.assertEqual("", due_tracking_kind(closed, NOW))

    def test_day_9_does_not_mark_expired_when_queueing_fails(self) -> None:
        backend = FakeBackend()
        backend.notification_error = RuntimeError("SharePoint unavailable")
        current = record(
            state="En progreso",
            permission_granted=True,
            assignment_email_sent=True,
            warning1_sent=True,
            warning2_sent=True,
            assignment_date=NOW - timedelta(days=9),
        )
        with self.assertRaisesRegex(RuntimeError, "SharePoint unavailable"):
            AutomationService(backend, NOW).track(current)
        self.assertFalse(any(patch.get("EstadoGantt") == "Vencido" for _, patch in backend.patches))

    def test_existing_expiration_notification_recovers_vencido_state(self) -> None:
        backend = FakeBackend()
        backend.notification_keys.add(("10", "vencimiento"))
        current = record(
            state="En progreso",
            assignment_date=NOW - timedelta(days=9),
            warning1_sent=True,
            warning2_sent=True,
        )
        result = AutomationService(backend, NOW).track(current)
        self.assertEqual("Vencido", result.state)
        self.assertTrue(any(patch.get("EstadoGantt") == "Vencido" for _, patch in backend.patches))


class StatusAndMetadataTests(unittest.TestCase):
    def test_review_status_requests_automatic_versioning(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En progreso",
            assignment_date=NOW - timedelta(days=4),
        )
        result = AutomationService(backend, NOW).sync_excel_status(
            current,
            WorkbookMetadata(status="En revisión inicial"),
        )
        self.assertEqual("Entregar", result.state)
        merged = {key: value for _, patch in backend.patches for key, value in patch.items()}
        self.assertEqual(4, merged["DiasParaCompletar"])
        self.assertTrue(merged["SolicitarVersionado"])
        self.assertTrue(result.version_requested)

    def test_extracts_engineer_supervisors_and_status_from_datos(self) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Datos"
        ws["A1"] = "Ingeniero residente"
        ws["B1"] = "Ingeniero@Example.com"
        ws["A2"] = "Supervisores"
        ws["B2"] = "uno@example.com, dos@example.com"
        ws["A3"] = "EstadoGantt"
        ws["B3"] = "En revisión inicial"
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        metadata = extract_workbook_metadata(stream.getvalue())
        self.assertEqual("ingeniero@example.com", metadata.engineer_email)
        self.assertEqual("uno@example.com;dos@example.com", metadata.supervisors_email)
        self.assertEqual("En revisión inicial", metadata.status)

    def test_visible_gantt_status_takes_priority_over_datos(self) -> None:
        workbook = Workbook()
        gantt = workbook.active
        gantt.title = "Gantt"
        gantt["A6"] = "Estado general del Gantt"
        gantt["B6"] = "En revisión inicial"
        datos = workbook.create_sheet("Datos")
        datos["A1"] = "EstadoGantt"
        datos["B1"] = "En progreso"
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        metadata = extract_workbook_metadata(stream.getvalue())
        self.assertEqual("En revisión inicial", metadata.status)

    def test_human_edit_moves_actual_to_en_progreso_and_starts_timer(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-old"',
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Actual"),
                etag='"etag-human"',
                modified_at=NOW,
                modified_by_email="ingeniero@example.com",
            ),
        )
        self.assertEqual("En Progreso", result.state)
        self.assertEqual(NOW, result.progress_started_at)
        merged = {
            key: value
            for _, patch_fields in backend.patches
            for key, value in patch_fields.items()
        }
        self.assertEqual("En Progreso", merged["StatusExcelDeseado"])
        self.assertEqual("Pendiente", merged["EstadoSyncExcel"])

    def test_automation_etag_does_not_reopen_actual_version(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-old"',
            automation_etag='"etag-automation"',
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Actual"),
                etag='"etag-automation"',
                modified_at=NOW,
            ),
        )
        self.assertEqual("Actual", result.state)

    def test_pending_power_automate_etag_is_confirmed_without_reopening(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-old"',
            automation_etag="__POWER_AUTOMATE_PENDING_ETAG__",
            desired_excel_status="Actual",
            excel_sync_state="Sincronizado",
            last_excel_sync_at=NOW,
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Actual"),
                etag='"etag-new"',
                modified_at=NOW - timedelta(seconds=10),
            ),
        )
        self.assertEqual("Actual", result.state)
        merged = {
            key: value
            for _, patch_fields in backend.patches
            for key, value in patch_fields.items()
        }
        self.assertEqual('"etag-new"', merged["UltimoETagAutomatizacion"])
        self.assertEqual("Sincronizado", merged["EstadoSyncExcel"])

    def test_interrupted_sync_is_recovered_after_script_changed_actual(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-old"',
            desired_excel_status="Actual",
            excel_sync_state="Procesando",
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Actual"),
                etag='"etag-new"',
                modified_at=NOW,
            ),
        )
        self.assertEqual("Actual", result.state)
        merged = {
            key: value
            for _, patch_fields in backend.patches
            for key, value in patch_fields.items()
        }
        self.assertEqual('"etag-new"', merged["UltimoETagAutomatizacion"])

    def test_stale_actual_label_does_not_close_active_progress(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En Progreso",
            current_version="v1.0",
            gantt_etag='"etag-current"',
            progress_started_at=NOW - timedelta(minutes=30),
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Actual"),
                etag='"etag-current"',
                modified_at=NOW,
            ),
        )
        self.assertEqual("En Progreso", result.state)
        self.assertEqual(
            NOW - timedelta(minutes=30),
            result.progress_started_at,
        )

    def test_entregar_has_priority_over_edit_detection(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-old"',
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(status="Entregar"),
                etag='"etag-human"',
                modified_at=NOW,
            ),
        )
        self.assertEqual("Entregar", result.state)
        self.assertTrue(result.version_requested)

    def test_legacy_review_label_does_not_reopen_closed_actual(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Actual",
            current_version="v1.0",
            gantt_etag='"etag-current"',
        )
        result = AutomationService(backend, NOW).observe_gantt_file(
            current,
            GanttFileState(
                metadata=WorkbookMetadata(
                    status="En revisión inicial"
                ),
                etag='"etag-current"',
                modified_at=NOW,
            ),
        )
        self.assertEqual("Actual", result.state)
        self.assertFalse(result.version_requested)


class VersioningTests(unittest.TestCase):
    @staticmethod
    def gantt_bytes(
        cost: float,
        activity_end: datetime | None = None,
        baseline_cost: float | None = None,
    ) -> bytes:
        workbook = Workbook()
        gantt = workbook.active
        gantt.title = "Gantt"
        gantt.append([])
        for _ in range(8):
            gantt.append([])
        gantt.append(
            [
                "Actividad",
                "Fecha de Inicio",
                "Fecha de Fin",
                "Costo Total",
            ]
        )
        gantt.append(["Actividad 1", None, activity_end, cost])
        datos = workbook.create_sheet("Datos")
        datos.append(["Fecha Final", datetime(2026, 12, 31)])
        if baseline_cost is not None:
            baseline = workbook.create_sheet("AutosysVersionBaseline")
            baseline.append(["AUTOSYS_VERSION_BASELINE", "Valor"])
            baseline.append(["Fecha Final contractual", datetime(2026, 12, 31)])
            baseline.append(["Costo Total", baseline_cost])
            baseline.sheet_state = "hidden"
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_version_copy_status_is_actual(self) -> None:
        workbook = Workbook()
        gantt = workbook.active
        gantt.title = "Gantt"
        gantt["A6"] = "Estado general del Gantt"
        gantt["B6"] = "Entregar"
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        updated = workbook_with_status(stream.getvalue(), "Actual")
        reopened = load_workbook(io.BytesIO(updated), data_only=True)
        try:
            self.assertEqual("Actual", reopened["Gantt"]["B6"].value)
            validation = next(
                item
                for item in reopened["Gantt"].data_validations.dataValidation
                if "B6" in str(item.sqref)
            )
            self.assertEqual(
                '"Actual,En Progreso,Entregar"',
                validation.formula1,
            )
        finally:
            reopened.close()

    def test_requested_review_creates_v1_and_closes_control(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En revisión inicial",
            version_requested=True,
        )

        result = AutomationService(backend, NOW).version(current)

        self.assertEqual(["10"], backend.versions)
        self.assertEqual("v1.0", result.current_version)
        self.assertEqual("Actual", result.state)
        self.assertFalse(result.version_requested)
        merged = {key: value for _, patch in backend.patches for key, value in patch.items()}
        self.assertEqual("v1.0", merged["VersionActual"])
        self.assertEqual("version-item-id", merged["GanttVersionIdentifier"])
        self.assertFalse(merged["SolicitarVersionado"])
        self.assertIn("FechaAprobacion", merged)

    def test_successful_delivery_closes_progress_timer(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Entregar",
            version_requested=True,
            progress_started_at=NOW - timedelta(minutes=90),
            progress_minutes_accumulated=10,
        )

        result = AutomationService(backend, NOW).version(current)

        self.assertEqual("Actual", result.state)
        self.assertEqual(100, result.progress_minutes_accumulated)
        self.assertIsNone(result.progress_started_at)
        merged = {
            key: value
            for _, patch_fields in backend.patches
            for key, value in patch_fields.items()
        }
        self.assertEqual(0, merged["MinutosEnProgresoActual"])
        self.assertEqual(100, merged["MinutosEnProgresoAcumulados"])
        self.assertEqual("Actual", merged["StatusExcelDeseado"])
        self.assertEqual("Pendiente", merged["EstadoSyncExcel"])

    def test_existing_v1_is_reassessed_before_reapproval(self) -> None:
        backend = FakeBackend()
        current = record(
            state="En revisión inicial",
            version_requested=True,
            current_version="v1.0",
            version_identifier="version-item-id",
        )

        result = AutomationService(backend, NOW).version(current)

        self.assertEqual(["10"], backend.versions)
        self.assertFalse(result.version_requested)
        self.assertEqual("Actual", result.state)

    def test_cost_increase_selects_v2(self) -> None:
        decision = decide_version(
            self.gantt_bytes(1200),
            self.gantt_bytes(1000),
        )
        self.assertEqual("v2.0", decision.version)
        self.assertTrue(decision.cost_increase)

    def test_price_total_is_not_counted_as_project_cost(self) -> None:
        content = self.gantt_bytes(1000, baseline_cost=1000)
        workbook = load_workbook(io.BytesIO(content))
        gantt = workbook["Gantt"]
        header_row = next(
            row
            for row in range(1, gantt.max_row + 1)
            if gantt.cell(row, 1).value == "Actividad"
        )
        gantt.cell(header_row, 5, "Precio Total")
        gantt.cell(header_row + 1, 5, 5000)
        baseline = workbook["AutosysVersionBaseline"]
        baseline.append(["Precio Total", 5000])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        snapshot = snapshot_gantt(stream.getvalue())
        self.assertEqual(1000, snapshot.total_cost)
        self.assertEqual(1000, snapshot.baseline_total_cost)

    def test_first_approval_uses_embedded_budget_baseline(self) -> None:
        decision = decide_version(
            self.gantt_bytes(1200, baseline_cost=1000),
            None,
        )
        self.assertEqual("v2.0", decision.version)
        self.assertTrue(decision.cost_increase)

    def test_legacy_gantt_without_baseline_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "Regenere el WORKING"):
            decide_version(self.gantt_bytes(1000), None)

    def test_cost_reduction_stays_v1(self) -> None:
        decision = decide_version(
            self.gantt_bytes(900),
            self.gantt_bytes(1000),
        )
        self.assertEqual("v1.0", decision.version)
        self.assertFalse(decision.cost_increase)

    def test_activity_after_contractual_end_selects_v2(self) -> None:
        decision = decide_version(
            self.gantt_bytes(1000, datetime(2027, 1, 2)),
            self.gantt_bytes(1000),
        )
        self.assertEqual("v2.0", decision.version)
        self.assertTrue(decision.schedule_overrun)

    def test_minor_versions_accumulate_without_major_change(self) -> None:
        decision = decide_version(
            self.gantt_bytes(900, datetime(2026, 12, 20)),
            self.gantt_bytes(1000, datetime(2026, 12, 15)),
            current_version="v2.2",
            prior_version_contents=(
                self.gantt_bytes(1000, datetime(2027, 1, 15)),
                self.gantt_bytes(1000, datetime(2026, 12, 15)),
            ),
        )
        self.assertEqual("v2.3", decision.version)
        self.assertFalse(decision.cost_increase)
        self.assertFalse(decision.schedule_overrun)

    def test_cost_increase_advances_major_and_resets_minor(self) -> None:
        decision = decide_version(
            self.gantt_bytes(1001, datetime(2026, 12, 20)),
            self.gantt_bytes(1000, datetime(2026, 12, 20)),
            current_version="v2.3",
        )
        self.assertEqual("v3.0", decision.version)
        self.assertTrue(decision.cost_increase)

    def test_new_latest_date_advances_major_again(self) -> None:
        latest_approved = datetime(2027, 2, 1)
        decision = decide_version(
            self.gantt_bytes(900, datetime(2027, 2, 2)),
            self.gantt_bytes(1000, datetime(2027, 1, 15)),
            current_version="v2.3",
            prior_version_contents=(
                self.gantt_bytes(1100, latest_approved),
                self.gantt_bytes(1000, datetime(2027, 1, 15)),
            ),
        )
        self.assertEqual("v3.0", decision.version)
        self.assertFalse(decision.cost_increase)
        self.assertTrue(decision.schedule_overrun)

    def test_major_change_uses_or_between_cost_and_schedule(self) -> None:
        cost_only = decide_version(
            self.gantt_bytes(1200, datetime(2026, 12, 20)),
            self.gantt_bytes(1000, datetime(2026, 12, 20)),
            current_version="v1.4",
        )
        date_only = decide_version(
            self.gantt_bytes(900, datetime(2027, 1, 2)),
            self.gantt_bytes(1000, datetime(2026, 12, 20)),
            current_version="v1.4",
        )
        self.assertEqual("v2.0", cost_only.version)
        self.assertEqual("v2.0", date_only.version)

    def test_contractual_end_prefers_embedded_generation_baseline(self) -> None:
        content = self.gantt_bytes(
            1000,
            datetime(2027, 1, 2),
            baseline_cost=1000,
        )
        workbook = load_workbook(io.BytesIO(content))
        workbook["Datos"]["B1"] = datetime(2030, 1, 1)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        snapshot = snapshot_gantt(stream.getvalue())
        self.assertEqual(datetime(2026, 12, 31).date(), snapshot.contractual_end)
        self.assertTrue(snapshot.schedule_overrun)

    def test_version_requires_review_state(self) -> None:
        backend = FakeBackend()
        current = record(
            state="Asignado",
            version_requested=True,
            version_attempts=2,
        )

        result = AutomationService(backend, NOW).version(current)

        self.assertEqual([], backend.versions)
        self.assertTrue(result.version_requested)
        merged = {key: value for _, patch in backend.patches for key, value in patch.items()}
        self.assertEqual(3, merged["VersionadoIntentos"])
        self.assertIn("Entregar", merged["UltimoErrorVersionado"])

    def test_version_failure_keeps_request_and_records_error(self) -> None:
        backend = FakeBackend()
        backend.version_error = RuntimeError("Graph upload 503")
        current = record(
            state="En revisión inicial",
            version_requested=True,
        )

        result = AutomationService(backend, NOW).version(current)

        self.assertTrue(result.version_requested)
        self.assertNotEqual("Aprobado / Versionado", result.state)
        merged = {key: value for _, patch in backend.patches for key, value in patch.items()}
        self.assertEqual(1, merged["VersionadoIntentos"])
        self.assertIn("Graph upload 503", merged["UltimoErrorVersionado"])

    def test_sharepoint_backend_creates_version_in_sibling_folder(self) -> None:
        backend = SharePointBackend.__new__(SharePointBackend)
        backend.token = "token"
        backend.site_id = "site-id"
        backend.resolve_drive_item = lambda link, identifier: {
            "id": "working-id",
            "name": "2026-001_gantt_WORKING.xlsx",
            "webUrl": "https://contoso/Proyectos/Proyectos%20Activos/P1/gantts/working/file.xlsx",
            "parentReference": {
                "driveId": "drive-id",
                "path": (
                    "/drives/drive-id/root:/Proyectos/Proyectos Activos/"
                    "2026-001 Proyecto/gantts/working"
                ),
            },
        }
        backend._download_drive_item = lambda item: b"xlsx-content"
        uploaded = {
            "id": "version-id",
            "name": "2026-001_gantt_v1.0.xlsx",
            "webUrl": "https://contoso/versionados/2026-001_gantt_v1.0.xlsx",
        }

        with (
            patch("src.automation_dispatcher.path_exists", return_value=False),
            patch(
                "src.automation_dispatcher.decide_version",
                return_value=type(
                    "Decision",
                    (),
                    {
                        "version": "v1.0",
                        "reasons": ("Sin aumento.",),
                    },
                )(),
            ),
            patch(
                "src.automation_dispatcher.workbook_with_status",
                return_value=b"version-content",
            ),
            patch("src.automation_dispatcher.ensure_drive_folder") as ensure_folder,
            patch(
                "src.automation_dispatcher.graph_put_bytes",
                return_value=uploaded,
            ) as put_bytes,
        ):
            artifact = backend.create_initial_version(record())

        self.assertTrue(artifact.created)
        self.assertEqual("version-id", artifact.identifier)
        ensure_folder.assert_called_once_with(
            "token",
            "site-id",
            (
                "Proyectos/Proyectos Activos/2026-001 Proyecto/"
                "gantts/versionados"
            ),
        )
        self.assertIn(
            "2026-001_gantt_v1.0.xlsx",
            put_bytes.call_args.args[1],
        )

    def test_sharepoint_backend_never_replaces_existing_version(self) -> None:
        backend = SharePointBackend.__new__(SharePointBackend)
        backend.token = "token"
        backend.site_id = "site-id"
        backend.resolve_drive_item = lambda link, identifier: {
            "id": "working-id",
            "webUrl": "https://contoso/Proyectos/Proyectos%20Activos/P1/gantts/working/file.xlsx",
            "parentReference": {
                "path": (
                    "/drives/drive-id/root:/Proyectos/Proyectos Activos/"
                    "2026-001 Proyecto/gantts/working"
                )
            },
        }
        existing = {
            "id": "existing-version-id",
            "name": "2026-001_gantt_v1.0.xlsx",
            "webUrl": "https://contoso/versionados/2026-001_gantt_v1.0.xlsx",
        }
        backend._download_drive_item = lambda item: b"xlsx-content"

        with (
            patch("src.automation_dispatcher.path_exists", return_value=True),
            patch(
                "src.automation_dispatcher.graph_get",
                return_value={"value": [{**existing, "file": {}}]},
            ),
            patch(
                "src.automation_dispatcher.decide_version",
                return_value=type(
                    "Decision",
                    (),
                    {
                        "version": "v1.0",
                        "reasons": ("Sin aumento.",),
                    },
                )(),
            ),
            patch(
                "src.automation_dispatcher.workbook_with_status",
                return_value=b"version-content",
            ),
            patch("src.automation_dispatcher.ensure_drive_folder"),
            patch("src.automation_dispatcher.graph_put_bytes") as put_bytes,
        ):
            with self.assertRaisesRegex(RuntimeError, "ya existe"):
                backend.create_initial_version(record())

        put_bytes.assert_not_called()


class DispatcherTests(unittest.TestCase):
    def test_empty_lists_finish_with_zero_errors(self) -> None:
        class EmptyBackend(FakeBackend):
            def process_status_events(self, now: datetime, max_items: int = 20) -> int:
                return 0

            def control_records(self) -> list[ControlRecord]:
                return []

        summary = run_system2(EmptyBackend(), NOW, 20)
        self.assertEqual(0, summary["control_items"])
        self.assertEqual(0, summary["errors"])
        self.assertEqual(0, summary["versioned"])

    def test_dispatcher_processes_requested_version_before_closed_state_filter(self) -> None:
        class VersionBackend(FakeBackend):
            def process_status_events(self, now: datetime, max_items: int = 20) -> int:
                return 0

            def control_records(self) -> list[ControlRecord]:
                return [
                    record(
                        state="En revisión inicial",
                        version_requested=True,
                    )
                ]

        backend = VersionBackend()
        summary = run_system2(backend, NOW, 20)
        self.assertEqual(1, summary["versioned"])
        self.assertEqual(0, summary["version_errors"])
        self.assertEqual(["10"], backend.versions)

    def test_dispatcher_versions_review_state_without_supervisor_flag(self) -> None:
        class ReviewBackend(FakeBackend):
            def process_status_events(self, now: datetime, max_items: int = 20) -> int:
                return 0

            def control_records(self) -> list[ControlRecord]:
                return [
                    record(
                        state="En revisión inicial",
                        version_requested=False,
                    )
                ]

        backend = ReviewBackend()
        summary = run_system2(backend, NOW, 20)
        self.assertEqual(1, summary["versioned"])
        self.assertEqual(["10"], backend.versions)
        self.assertTrue(
            any(
                patch.get("SolicitarVersionado") is True
                for _, patch in backend.patches
            )
        )

    def test_dispatcher_polls_excel_and_versions_without_status_event(self) -> None:
        class PollingBackend(FakeBackend):
            def __init__(self) -> None:
                super().__init__(WorkbookMetadata(status="En revisión inicial"))

            def process_status_events(self, now: datetime, max_items: int = 20) -> int:
                return 0

            def control_records(self) -> list[ControlRecord]:
                return [
                    record(
                        state="En progreso",
                        permission_granted=True,
                        assignment_email_sent=True,
                        assignment_date=NOW - timedelta(days=2),
                    )
                ]

        backend = PollingBackend()
        summary = run_system2(backend, NOW, 20)
        self.assertEqual(1, summary["versioned"])
        self.assertEqual(["10"], backend.versions)
        merged = {
            key: value
            for _, patch in backend.patches
            for key, value in patch.items()
        }
        self.assertEqual("Actual", merged["EstadoGantt"])
        self.assertEqual("v1.0", merged["VersionActual"])
        self.assertTrue(
            any(
                patch.get("EstadoGantt") == "Entregar"
                for _, patch in backend.patches
            )
        )

    def test_isolated_run_only_processes_requested_control_item(self) -> None:
        class IsolatedBackend(FakeBackend):
            def process_status_events(self, now: datetime, max_items: int = 20) -> int:
                raise AssertionError("isolated run must not process global status events")

            def control_records(self) -> list[ControlRecord]:
                return [
                    record(item_id="10", state="Aprobado / Versionado"),
                    record(item_id="11", state="Aprobado / Versionado"),
                ]

        summary = run_system2(IsolatedBackend(), NOW, 20, "11")
        self.assertEqual(1, summary["control_items"])
        self.assertEqual(0, summary["status_events"])

    def test_isolated_run_fails_clearly_for_unknown_control_item(self) -> None:
        class IsolatedBackend(FakeBackend):
            def control_records(self) -> list[ControlRecord]:
                return [record(item_id="10")]

        with self.assertRaisesRegex(RuntimeError, "No existe el item 99"):
            run_system2(IsolatedBackend(), NOW, 20, "99")

    def test_force_version_recheck_reopens_only_requested_review_workbook(self) -> None:
        class RecheckBackend(FakeBackend):
            def __init__(self) -> None:
                super().__init__(
                    WorkbookMetadata(status="En revisión inicial")
                )

            def control_records(self) -> list[ControlRecord]:
                return [
                    record(
                        item_id="11",
                        state="Aprobado / Versionado",
                        current_version="v1.0",
                    )
                ]

        backend = RecheckBackend()
        summary = run_system2(
            backend,
            NOW,
            20,
            "11",
            force_version_recheck=True,
        )
        self.assertEqual(1, summary["versioned"])
        self.assertTrue(
            any(
                patch.get("EstadoGantt") == "Entregar"
                and patch.get("SolicitarVersionado") is True
                for _, patch in backend.patches
            )
        )

    def test_test_delivery_mode_redirects_and_removes_real_cc(self) -> None:
        notification = Notification(
            "Advertencia2",
            "ingeniero@example.com",
            "supervisor@example.com",
            "Aviso",
            "Contenido",
        )
        with patch.dict(
            "os.environ",
            {
                "NOTIFICATION_DELIVERY_MODE": "test",
                "NOTIFICATION_TEST_RECIPIENT": "auto.sys@prowallpanama.com",
            },
        ):
            redirected = apply_notification_delivery_mode(notification)
        self.assertEqual("auto.sys@prowallpanama.com", redirected.to)
        self.assertEqual("", redirected.cc)
        self.assertEqual("[PRUEBA] Aviso", redirected.subject)
        self.assertIn("ingeniero@example.com", redirected.body)
        self.assertIn("supervisor@example.com", redirected.body)

    def test_live_delivery_mode_keeps_original_recipients(self) -> None:
        notification = Notification("Advertencia1", "ingeniero@example.com", "", "Aviso", "Body")
        with patch.dict("os.environ", {"NOTIFICATION_DELIVERY_MODE": "live"}):
            delivered = apply_notification_delivery_mode(notification)
        self.assertEqual(notification, delivered)

    def test_engineer_guide_url_is_inserted_in_assignment(self) -> None:
        with patch.dict(
            "os.environ",
            {"ENGINEER_GUIDE_URL": "https://contoso.example/guia.pdf"},
        ):
            backend = FakeBackend()
            AutomationService(backend, NOW).assign(record())
        self.assertIn(
            "https://contoso.example/guia.pdf",
            backend.notifications[0][1].body,
        )


def json_keys(patches: list[tuple[str, dict]]) -> str:
    return " ".join(key for _, patch in patches for key in patch)


if __name__ == "__main__":
    unittest.main()
