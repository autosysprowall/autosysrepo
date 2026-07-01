from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from budget_extraction import (  # noqa: E402
    BudgetExtractionError,
    extract_budget_structure,
    normalize_text,
)
from sistema1_gantt_builder import (  # noqa: E402
    DATA_START_ROW,
    GanttReviewRequiredError,
    HEADER_ROW,
    LlmOptions,
    build_gantt_workbook,
    derive_project_identity,
    read_project_window,
    select_budget_sheet,
)


def add_datos(
    wb: Workbook,
    *,
    start: date = date(2026, 7, 1),
    end: date = date(2026, 7, 12),
    duration: str = "2 meses",
) -> None:
    ws = wb.create_sheet("Datos")
    ws.append(["Ingeniero residente", "ingeniero@example.com"])
    ws.append(["Supervisores", "supervisor@example.com"])
    ws.append(["Fecha de Inicio", start])
    ws.append(["Fecha Final", end])
    ws.append(["Duracion", duration])


def make_budget(
    path: Path,
    *,
    headers: list[str],
    rows: list[list[object]],
    preamble_rows: int = 0,
    with_datos: bool = True,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws["A1"] = "PRESUPUESTO GENERAL"
    ws["A1"].fill = PatternFill("solid", fgColor="C00000")
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    for _ in range(preamble_rows):
        ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws["L2"] = "=SUM(1,2)"
    if with_datos:
        add_datos(wb)
    wb.save(path)
    wb.close()


class BudgetExtractionTests(unittest.TestCase):
    def test_presupuesto_flexio_is_the_official_preferred_sheet(self) -> None:
        wb = Workbook()
        wb.active.title = "PRESUPUESTO"
        wb.create_sheet("PRESUPUESTO FLEXIO")
        self.assertEqual("PRESUPUESTO FLEXIO", select_budget_sheet(wb))

    def test_normalizes_accents_punctuation_and_line_breaks(self) -> None:
        self.assertEqual("precio unitario", normalize_text("  PRÉCIO\nUnitario  "))
        self.assertEqual("p u", normalize_text("P.U."))

    def test_normal_columns_preserve_costs_and_cc(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Presupuesto"
        ws.append(["Código", "Descripción", "Unidad", "Cantidad", "Costo Unitario", "Costo Total"])
        ws.append(["1.1", "Concreto", "m3", 5, 100, 500])
        result = extract_budget_structure(
            ws,
            source_file="2026-001 - Prueba.xlsx",
            source_sheet=ws.title,
        )
        self.assertEqual(1, result.header_row)
        self.assertEqual("Concreto", result.rows[0].actividad)
        self.assertEqual("1.1", result.rows[0].cc)
        self.assertEqual(100, result.rows[0].costo_unitario)
        self.assertEqual(500, result.rows[0].costo_total)
        reusable = result.to_dict()
        self.assertEqual("Presupuesto", reusable["source_sheet"])
        self.assertEqual(2, reusable["rows"][0]["row_number"])

    def test_reordered_alias_columns_and_preamble_are_detected(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Cliente", "Prowall"])
        ws.append([])
        ws.append(["Total", "Cant", "Partida", "P.U.", "Und"])
        ws.append([750, 3, "Formaleta", 250, "m2"])
        result = extract_budget_structure(
            ws,
            source_file="alias.xlsx",
            source_sheet=ws.title,
        )
        row = result.rows[0]
        self.assertEqual(3, result.header_row)
        self.assertEqual("Formaleta", row.actividad)
        self.assertEqual("m2", row.unidad)
        self.assertEqual(3, row.cantidad)
        self.assertEqual(250, row.costo_unitario)
        self.assertEqual(750, row.costo_total)

    def test_empty_totals_and_sections_keep_source_order(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Código", "Descripción", "Und", "Cant", "P.U.", "Total"])
        ws.append(["1", "OBRA CIVIL", None, None, None, None])
        ws.append([None, None, None, None, None, None])
        ws.append(["1.1", "Excavación", "m3", 2, 30, 60])
        ws.append([None, "Subtotal", None, None, None, 60])
        result = extract_budget_structure(
            ws,
            source_file="sections.xlsx",
            source_sheet=ws.title,
        )
        self.assertEqual(
            ["section", "spacer", "activity"],
            [row.row_type for row in result.rows],
        )
        self.assertEqual(
            ["OBRA CIVIL", "", "Excavación"],
            [row.actividad for row in result.rows],
        )
        reasons = {item["reason"] for item in result.assessment.ignored_rows}
        self.assertIn("empty", reasons)
        self.assertIn("total", reasons)

    def test_repeated_header_can_change_column_order_mid_sheet(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Código", "Descripción", "Und", "Cant", "P.U.", "Total"])
        ws.append(["1.1", "Concreto", "m3", 2, 50, 100])
        ws.append(["Descripción", "Und", "Código", "Total", "Cant", "P.U."])
        ws.append(["Formaleta", "m2", "1.2", 120, 3, 40])
        result = extract_budget_structure(
            ws,
            source_file="multiple-blocks.xlsx",
            source_sheet=ws.title,
        )
        self.assertEqual(["Concreto", "Formaleta"], [row.actividad for row in result.rows])
        self.assertEqual("1.2", result.rows[1].cc)
        self.assertEqual("m2", result.rows[1].unidad)
        self.assertEqual(3, result.rows[1].cantidad)
        self.assertEqual(40, result.rows[1].costo_unitario)
        self.assertEqual(120, result.rows[1].costo_total)

    def test_missing_total_is_calculated_without_overwriting_source_total(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Descripción", "Unidad", "Cantidad", "Costo Unitario", "Costo Total"])
        ws.append(["Acero", "kg", 4, 25, None])
        ws.append(["Concreto", "m3", 2, 10, 100])
        result = extract_budget_structure(
            ws,
            source_file="costs.xlsx",
            source_sheet=ws.title,
        )
        self.assertEqual(100, result.rows[0].costo_total)
        self.assertEqual([2], result.assessment.calculated_cost_rows)
        self.assertEqual(100, result.rows[1].costo_total)
        self.assertTrue(result.rows[1].warnings)

    def test_llm_mapper_hook_only_fills_existing_columns(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Trabajo comercial", "Medición", "Vol", "Tarifa", "Extensión"])
        ws.append(["Instalación", "m2", 2, 20, 40])
        calls: list[tuple[str, ...]] = []

        def mapper(sheet_name, headers, missing):
            calls.append(missing)
            return {
                "actividad": 1,
                "unidad": 2,
                "cantidad": 3,
                "costo_unitario": 4,
                "costo_total": 5,
                "cc": 99,
            }

        result = extract_budget_structure(
            ws,
            source_file="llm.xlsx",
            source_sheet=ws.title,
            column_mapper=mapper,
        )
        self.assertTrue(calls)
        self.assertTrue(result.assessment.llm_mapping_used)
        self.assertIsNone(result.column_mapping["cc"])
        self.assertEqual("Instalación", result.rows[0].actividad)

    def test_unreliable_header_raises_review_error(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Cliente", "Prowall"])
        ws.append(["Comentario", "Sin tabla de costos"])
        with self.assertRaises(BudgetExtractionError):
            extract_budget_structure(
                ws,
                source_file="invalid.xlsx",
                source_sheet=ws.title,
            )

    def test_quality_gate_rejects_missing_general_budget(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Actividad", "Unidad", "Cantidad", "Costo Unitario", "Costo Total"])
        ws.append(["Concreto", "m3", 2, 100, 200])
        with self.assertRaisesRegex(
            BudgetExtractionError,
            "SIN_PRESUPUESTO_GENERAL",
        ):
            extract_budget_structure(
                ws,
                source_file="dividido.xlsx",
                source_sheet=ws.title,
                enforce_quality_gate=True,
            )

    def test_quality_gate_rejects_semantically_swapped_columns(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["PRESUPUESTO GENERAL"])
        ws.append(["Actividad", "Unidad", "Cantidad", "Costo Unitario", "Costo Total"])
        ws.append(["Concreto", 25, "mucho", 100, 200])
        ws.append(["Formaleta", 30, "varios", 50, 150])
        with self.assertRaisesRegex(
            BudgetExtractionError,
            "CONTENIDO_DE_COLUMNAS_INVALIDO",
        ):
            extract_budget_structure(
                ws,
                source_file="encabezados-rotos.xlsx",
                source_sheet=ws.title,
                enforce_quality_gate=True,
            )

    def test_quality_gate_rejects_repeated_budget_headers(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["PRESUPUESTO GENERAL"])
        ws.append(["Código", "Descripción", "Und", "Cant", "P.U.", "Total"])
        ws.append(["1.1", "Concreto", "m3", 2, 50, 100])
        ws.append(["Descripción", "Und", "Código", "Total", "Cant", "P.U."])
        ws.append(["Formaleta", "m2", "1.2", 120, 3, 40])
        with self.assertRaisesRegex(
            BudgetExtractionError,
            "ENCABEZADOS_REPETIDOS",
        ):
            extract_budget_structure(
                ws,
                source_file="multiple-blocks.xlsx",
                source_sheet=ws.title,
                enforce_quality_gate=True,
            )

    def test_quality_gate_requires_a_financial_column(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["PRESUPUESTO GENERAL"])
        ws.append(["Actividad", "Unidad", "Cantidad"])
        ws.append(["Concreto", "m3", 2])
        ws.append(["Formaleta", "m2", 3])
        with self.assertRaisesRegex(
            BudgetExtractionError,
            "COLUMNAS_DE_COSTO_AUSENTES",
        ):
            extract_budget_structure(
                ws,
                source_file="sin-costos.xlsx",
                source_sheet=ws.title,
                enforce_quality_gate=True,
            )

    def test_item_material_and_multiple_financial_groups_use_content_validation(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Presupuesto FLEXIO"
        ws["C17"] = "PRESUPUESTO GENERAL (1 CASA)"
        ws.merge_cells("C17:I18")
        ws["J17"] = "GLOBAL (2 CASAS)"
        ws.merge_cells("J17:M18")
        ws["N17"] = "ACUMULADO (2 CASAS)"
        ws.merge_cells("N17:Q18")
        headers = [
            "PROWALL",
            "CC",
            "Ítems",
            "Material",
            "Cantidad",
            "Unidad",
            "Costo unitario",
            "Costo total (1 casa)",
            "Precio Unitario (1 casa)",
            "Costo Total (2 casas)",
            "Precio Total (2 casas)",
            "Utilidad Total",
            "Margen",
            "Costo Total (2 casas)",
            "Precio Total (2 casas)",
            "Utilidad Total",
            "Margen",
        ]
        for column, value in enumerate(headers, start=1):
            ws.cell(19, column, value)
        ws.append([])
        ws["J20"] = 100
        rows = [
            [
                "PROWALL/CAMPO",
                None,
                1.1,
                "Vaciado de Piso de Fundación",
                None,
                None,
                None,
                3245.73,
                3836.56,
                6491.46,
                7673.12,
                1181.66,
                0.154,
                7756.46,
                9168.39,
                1411.93,
                0.154,
            ],
            [
                "PROWALL/CAMPO",
                "CC.03.03.01.02",
                "1.1.1",
                "Concreto de 3000 psi",
                10,
                "m3",
                120,
                1200,
                None,
                2400,
                2836.88,
                436.88,
                0.154,
                2841.6,
                3358.87,
                517.27,
                0.154,
            ],
        ]
        for row_number, values in enumerate(rows, start=21):
            for column, value in enumerate(values, start=1):
                ws.cell(row_number, column, value)
        ws.auto_filter.ref = "A19:A22"
        calls: list[tuple[str, ...]] = []

        def mapper(sheet_name, candidate_headers, requested_fields):
            calls.append(requested_fields)
            return {"item": 4, "actividad": 3}

        result = extract_budget_structure(
            ws,
            source_file="multi-cost.xlsx",
            source_sheet=ws.title,
            column_mapper=mapper,
        )
        self.assertTrue(calls)
        self.assertEqual(3, result.column_mapping["item"])
        self.assertEqual(4, result.column_mapping["actividad"])
        self.assertEqual("1.1.1", str(result.rows[1].item))
        self.assertEqual("Concreto de 3000 psi", result.rows[1].actividad)
        self.assertEqual(11, len(result.cost_columns))
        self.assertEqual(
            ["PRESUPUESTO GENERAL (1 CASA)", "GLOBAL (2 CASAS)", "ACUMULADO (2 CASAS)"],
            list(dict.fromkeys(column.group for column in result.cost_columns)),
        )
        self.assertEqual("complex_validation", result.assessment.llm_validation["mode"])
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2025-111 - Multi costo.xlsx"
            output = root / "2025-111_gantt_WORKING.xlsx"
            add_datos(wb)
            wb.save(source)
            build_gantt_workbook(
                source,
                output,
                source.name,
                LlmOptions(enabled=False),
            )
            generated = load_workbook(output, data_only=False)
            try:
                gantt = generated["Gantt"]
                self.assertEqual("Ítem", gantt.cell(HEADER_ROW, 1).value)
                self.assertEqual("Actividad", gantt.cell(HEADER_ROW, 2).value)
                self.assertEqual(1.1, gantt.cell(DATA_START_ROW, 1).value)
                self.assertEqual(
                    "Vaciado de Piso de Fundación",
                    gantt.cell(DATA_START_ROW, 2).value,
                )
                group_labels = {
                    gantt.cell(HEADER_ROW - 1, column).value
                    for column in range(1, gantt.max_column + 1)
                    if isinstance(gantt.cell(HEADER_ROW - 1, column).value, str)
                }
                self.assertIn("GLOBAL (2 CASAS)", group_labels)
                self.assertIn("ACUMULADO (2 CASAS)", group_labels)
                header_indexes = {
                    gantt.cell(HEADER_ROW, column).value: column
                    for column in range(1, gantt.max_column + 1)
                    if isinstance(gantt.cell(HEADER_ROW, column).value, str)
                }
                global_cost_header = next(
                    header
                    for header in header_indexes
                    if "Costo Total (2 casas)" in header
                    and "GLOBAL (2 CASAS)" in header
                )
                self.assertEqual(
                    6491.46,
                    gantt.cell(
                        DATA_START_ROW,
                        header_indexes[global_cost_header],
                    ).value,
                )
            finally:
                generated.close()


class GanttWorkbookTests(unittest.TestCase):
    def test_project_calendar_understands_duration_units_and_prioritizes_end(self) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        datos = wb.create_sheet("Datos")
        datos.append(["Fecha de Inicio", date(2026, 7, 1)])
        datos.append(["Duración", "2 semanas"])
        start, end, _ = read_project_window(wb)
        self.assertEqual(date(2026, 7, 1), start)
        self.assertEqual(date(2026, 7, 14), end)

        datos.append(["Fecha Final", date(2026, 8, 20)])
        _, prioritized_end, notes = read_project_window(wb)
        self.assertEqual(date(2026, 8, 20), prioritized_end)
        self.assertTrue(any("prioridad" in note for note in notes))

    def test_invalid_budget_is_stopped_for_manual_review(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2026-099 - Inválido.xlsx"
            output = root / "2026-099_gantt_WORKING.xlsx"
            make_budget(
                source,
                headers=["Cliente", "Comentario"],
                rows=[["Prowall", "Sin tabla comercial"]],
                with_datos=False,
            )
            with self.assertRaises(GanttReviewRequiredError):
                build_gantt_workbook(
                    source,
                    output,
                    source.name,
                    LlmOptions(enabled=False),
                )
            self.assertFalse(output.exists())

    def test_unreadable_excel_is_rejected_with_unknown_cause(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2026-098 - Roto.xlsx"
            output = root / "2026-098_gantt_WORKING.xlsx"
            source.write_bytes(b"not-an-excel-workbook")
            with self.assertRaisesRegex(
                GanttReviewRequiredError,
                "ARCHIVO_ROTO_CAUSA_DESCONOCIDA",
            ):
                build_gantt_workbook(
                    source,
                    output,
                    source.name,
                    LlmOptions(enabled=False),
                )
            self.assertFalse(output.exists())

    def test_generated_workbook_preserves_budget_and_builds_functional_gantt(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2026-100 - Proyecto Costos.xlsx"
            output = root / "2026-100_gantt_WORKING.xlsx"
            make_budget(
                source,
                headers=["Código", "Descripción", "Und", "Cant", "P.U.", "Total"],
                rows=[
                    ["1", "PRELIMINARES", None, None, None, None],
                    ["1.1", "Movilización", "gl", 1, 250, 250],
                    ["1.2", "Concreto", "m3", 5, 100, 500],
                    [None, "Gran Total", None, None, None, 750],
                ],
                preamble_rows=4,
            )
            result = build_gantt_workbook(
                source,
                output,
                source.name,
                llm_options=LlmOptions(enabled=False),
            )
            self.assertEqual("ok", result.validation["status"])
            self.assertEqual(2, result.rows_written)

            wb = load_workbook(output, data_only=False)
            try:
                self.assertIn("Presupuesto", wb.sheetnames)
                self.assertIn("Gantt", wb.sheetnames)
                self.assertIn("Datos", wb.sheetnames)
                self.assertIn("Assessment", wb.sheetnames)
                self.assertEqual("00C00000", wb["Presupuesto"]["A1"].fill.fgColor.rgb)
                self.assertEqual("=SUM(1,2)", wb["Presupuesto"]["L2"].value)

                ws = wb["Gantt"]
                headers = [
                    ws.cell(HEADER_ROW, column).value
                    for column in range(1, 10)
                ]
                self.assertEqual(
                    [
                        "Actividad",
                        "CC",
                        "Fecha de Inicio",
                        "Fecha de Fin",
                        "Estatus",
                        "Unidad",
                        "Cantidad",
                        "Costo Unitario",
                        "Costo Total",
                    ],
                    headers,
                )
                self.assertIsNone(ws.freeze_panes)
                activity_rows = [DATA_START_ROW + 1, DATA_START_ROW + 2]
                for row_number in activity_rows:
                    self.assertIsNone(ws.cell(row_number, 3).value)
                    self.assertIsNone(ws.cell(row_number, 4).value)
                self.assertEqual(250, ws.cell(activity_rows[0], 9).value)
                self.assertEqual(500, ws.cell(activity_rows[1], 9).value)
                self.assertGreater(len(ws.data_validations.dataValidation), 0)
                self.assertGreater(len(ws.conditional_formatting), 0)
                section_range = next(
                    merged
                    for merged in ws.merged_cells.ranges
                    if merged.min_row == DATA_START_ROW
                )
                self.assertGreater(section_range.max_col, 9)
                self.assertEqual("right", ws.cell(DATA_START_ROW, 1).alignment.horizontal)
                calendar_start = 10
                self.assertEqual(date(2026, 7, 1), ws.cell(HEADER_ROW, calendar_start).value.date())
                self.assertEqual(date(2026, 7, 12), ws.cell(HEADER_ROW, calendar_start + 11).value.date())
            finally:
                wb.close()

    def test_cc_column_is_omitted_when_source_has_no_independent_cc(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2026-101 - Sin CC.xlsx"
            output = root / "2026-101_gantt_WORKING.xlsx"
            make_budget(
                source,
                headers=["Actividad", "Unidad", "Cantidad", "Costo Unitario", "Costo Total"],
                rows=[["Instalación", "m2", 2, 25, 50]],
            )
            build_gantt_workbook(source, output, source.name, LlmOptions(enabled=False))
            wb = load_workbook(output)
            try:
                headers = [
                    wb["Gantt"].cell(HEADER_ROW, column).value
                    for column in range(1, 9)
                ]
                self.assertNotIn("CC", headers)
                self.assertEqual("Fecha de Inicio", headers[1])
            finally:
                wb.close()

    def test_source_cost_formulas_are_linked_instead_of_lost(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2026-102 - Fórmulas.xlsx"
            output = root / "2026-102_gantt_WORKING.xlsx"
            make_budget(
                source,
                headers=["Código", "Descripción", "Und", "Cant", "P.U.", "Total"],
                rows=[["1.1", "Concreto", "m3", "=2+3", "=200/2", "=D3*E3"]],
            )
            build_gantt_workbook(source, output, source.name, LlmOptions(enabled=False))
            wb = load_workbook(output, data_only=False)
            try:
                ws = wb["Gantt"]
                self.assertEqual("='Presupuesto'!D3", ws.cell(DATA_START_ROW, 7).value)
                self.assertEqual("='Presupuesto'!E3", ws.cell(DATA_START_ROW, 8).value)
                self.assertEqual("='Presupuesto'!F3", ws.cell(DATA_START_ROW, 9).value)
            finally:
                wb.close()

    def test_flexio_costs_styles_spacers_and_table_boundary_are_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "2025-135 Escuela Santa Cecilia.xlsx"
            output = root / "2025-135_gantt_WORKING.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "PRESUPUESTO FLEXIO"
            ws.append(
                [
                    "Actividad",
                    "Unidad",
                    "Cantidad",
                    "C. Unitario",
                    "C. Total",
                    "P.U.",
                    "COSTO TOTAL",
                ]
            )
            ws["J1"] = "PRESUPUESTO GENERAL"
            ws.append(
                [
                    "Administracion",
                    "m2",
                    4436.5,
                    32.72,
                    145184.23,
                    39.65,
                    175934.67,
                ]
            )
            ws.append([])
            ws.append(
                [
                    "Suministro",
                    "m2",
                    3793.24,
                    32,
                    121383.68,
                    None,
                    121383.68,
                ]
            )
            for _ in range(5):
                ws.append([])
            ws.append(["PROYECTO", "Area", 2, 74.36, 148.72, None, None])
            ws.auto_filter.ref = "A1:A4"
            ws.row_dimensions[2].height = 23.25
            gray = PatternFill("solid", fgColor="D0D0D0")
            pink = PatternFill("solid", fgColor="F2CEEF")
            ws["A2"].fill = gray
            ws["A2"].font = Font(bold=True)
            ws["D2"].fill = pink
            ws["D2"].font = Font(bold=True)
            ws["E2"].fill = gray
            ws["E2"].font = Font(bold=True)
            currency_format = (
                '" "[$B/.-180A]* #,##0.00" ";"-"[$B/.-180A]* #,##0.00'
            )
            ws["D2"].number_format = currency_format
            ws["E2"].number_format = currency_format
            add_datos(wb)
            wb.save(source)
            wb.close()

            build_gantt_workbook(source, output, source.name, LlmOptions(enabled=False))
            generated = load_workbook(output, data_only=False)
            try:
                gantt = generated["Gantt"]
                headers = {
                    gantt.cell(HEADER_ROW, column).value: column
                    for column in range(1, 9)
                }
                self.assertEqual(
                    32.72,
                    gantt.cell(DATA_START_ROW, headers["Costo Unitario"]).value,
                )
                self.assertEqual(
                    145184.23,
                    gantt.cell(DATA_START_ROW, headers["Costo Total"]).value,
                )
                self.assertEqual(
                    "00F2CEEF",
                    gantt.cell(
                        DATA_START_ROW,
                        headers["Costo Unitario"],
                    ).fill.fgColor.rgb,
                )
                self.assertEqual(
                    "00D0D0D0",
                    gantt.cell(DATA_START_ROW, headers["Actividad"]).fill.fgColor.rgb,
                )
                self.assertEqual(
                    currency_format,
                    gantt.cell(
                        DATA_START_ROW,
                        headers["Costo Unitario"],
                    ).number_format,
                )
                self.assertEqual(23.25, gantt.row_dimensions[DATA_START_ROW].height)
                self.assertIsNone(gantt.cell(DATA_START_ROW + 1, 1).value)
                self.assertEqual(
                    "Suministro",
                    gantt.cell(DATA_START_ROW + 2, 1).value,
                )
                activities = [
                    gantt.cell(row, 1).value
                    for row in range(DATA_START_ROW, gantt.max_row + 1)
                ]
                self.assertNotIn("PROYECTO", activities)
                calendar_column = next(
                    column
                    for column in range(1, gantt.max_column + 1)
                    if isinstance(gantt.cell(HEADER_ROW, column).value, date)
                )
                self.assertEqual(
                    "00D0D0D0",
                    gantt.cell(DATA_START_ROW, calendar_column).fill.fgColor.rgb,
                )
            finally:
                generated.close()

    def test_output_identity_uses_required_working_name(self) -> None:
        identity = derive_project_identity("2026-123 - Proyecto.xlsx")
        self.assertEqual("2026-123_gantt_WORKING.xlsx", identity.gantt_file_name)


if __name__ == "__main__":
    unittest.main()
