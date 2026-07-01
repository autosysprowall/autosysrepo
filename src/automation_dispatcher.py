from __future__ import annotations

import argparse
import base64
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, replace
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Protocol

import requests
from openpyxl import load_workbook

from src.gantt_versioning import decide_version


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from sistema1_poll_queue import (  # noqa: E402
    GRAPH_BASE,
    acquire_token,
    build_field_map,
    drive_path_from_queue_fields,
    env_value,
    graph_get,
    graph_patch,
    graph_post,
    graph_put_bytes,
    encoded_drive_path,
    ensure_drive_folder,
    get_drive_item_by_path,
    list_columns,
    load_settings,
    path_exists,
    pick_field,
    print_token_diagnostics,
    resolve_list,
    resolve_site,
)


CONTROL_LIST_NAME = "Control_Gantt_Asignaciones"
QUEUE_LIST_NAME = "Cola_Automatizacion_Proyectos"
NOTIFICATION_LIST_NAME = "Cola_Notificaciones_Gantt"
FORBIDDEN_PATH = "proyectos terminados"
ACTIVE_TRACKING_STATES = {"asignado", "en progreso"}
CLOSED_STATES = {"aprobado / versionado", "vencido"}
REVIEW_STATE = "En revisión inicial"
EMAIL_RE = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")


CONTROL_COLUMNS: dict[str, dict[str, Any]] = {
    "GanttWorkingIdentifier": {"text": {}},
    "PresupuestoIdentifier": {"text": {}},
    "SupervisoresEmail": {"text": {}},
    "PermisoGanttOtorgado": {"boolean": {}},
    "FechaPermisoOtorgado": {"dateTime": {"format": "dateTime"}},
    "CorreoAsignacionEnviado": {"boolean": {}},
    "FechaCorreoAsignacion": {"dateTime": {"format": "dateTime"}},
    "FechaAdvertencia1": {"dateTime": {"format": "dateTime"}},
    "FechaAdvertencia2": {"dateTime": {"format": "dateTime"}},
    "Advertencia1Enviada": {"boolean": {}},
    "Advertencia2Enviada": {"boolean": {}},
    "VencimientoNotificado": {"boolean": {}},
    "FechaVencimientoNotificado": {"dateTime": {"format": "dateTime"}},
    "UltimoTrackingRun": {"dateTime": {"format": "dateTime"}},
    "TrackingIntentos": {"number": {}},
    "UltimoErrorTracking": {"text": {"allowMultipleLines": True}},
    "StatusExcel": {"text": {}},
    "FechaLecturaStatusExcel": {"dateTime": {"format": "dateTime"}},
    "SolicitarVersionado": {"boolean": {}},
    "VersionActual": {"text": {}},
    "GanttVersionLink": {"text": {}},
    "GanttVersionIdentifier": {"text": {}},
    "FechaUltimoVersionado": {"dateTime": {"format": "dateTime"}},
    "VersionadoIntentos": {"number": {}},
    "UltimoErrorVersionado": {"text": {"allowMultipleLines": True}},
    "MotivoUltimoVersionado": {"text": {"allowMultipleLines": True}},
}

NOTIFICATION_COLUMNS: dict[str, dict[str, Any]] = {
    "ProyectoID": {"text": {}},
    "TipoNotificacion": {
        "choice": {
            "allowTextEntry": False,
            "choices": ["AsignacionGantt", "Advertencia1", "Advertencia2", "Vencimiento"],
            "displayAs": "dropDownMenu",
        }
    },
    "EstadoNotificacion": {
        "choice": {
            "allowTextEntry": False,
            "choices": ["Pendiente", "Enviado", "Error"],
            "displayAs": "dropDownMenu",
        }
    },
    "To": {"text": {}},
    "Cc": {"text": {}},
    "Subject": {"text": {}},
    "Body": {"text": {"allowMultipleLines": True}},
    "RelatedControlItemID": {"text": {}},
    "Intentos": {"number": {}},
    "UltimoError": {"text": {"allowMultipleLines": True}},
    "FechaCreacion": {"dateTime": {"format": "dateTime"}},
    "FechaEnvio": {"dateTime": {"format": "dateTime"}},
    "Notas": {"text": {"allowMultipleLines": True}},
}


def normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().casefold()


def normalize_emails(value: Any) -> str:
    seen: set[str] = set()
    result: list[str] = []
    for match in EMAIL_RE.findall(str(value or "").replace(";", " ").replace(",", " ")):
        email = match.casefold()
        if email not in seen:
            seen.add(email)
            result.append(email)
    return ";".join(result)


def is_valid_email(value: str) -> bool:
    emails = normalize_emails(value).split(";") if value else []
    return len(emails) == 1 and emails[0] == value.strip().casefold()


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return normalized(value) in {"true", "1", "si", "yes"}


def parse_int(value: Any) -> int:
    try:
        return int(float(str(value or "0").strip() or "0"))
    except ValueError:
        return 0


def sanitize_error(exc: Exception) -> str:
    message = str(exc)
    for name in (
        "MS_TENANT_ID",
        "MS_CLIENT_ID",
        "MS_CLIENT_SECRET",
        "MS_GRAPH_CLIENT_SECRET",
        "OPENAI_API_KEY",
    ):
        secret = os.getenv(name, "")
        if secret:
            message = message.replace(secret, "***")
    return message[:500]


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def iso_utc(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def link_text(value: Any) -> str:
    if isinstance(value, dict):
        return str(value.get("Url") or value.get("url") or value.get("Description") or "")
    return str(value or "").strip()


def contains_forbidden_path(*values: Any) -> bool:
    combined = " ".join(link_text(value) for value in values)
    decoded = requests.utils.unquote(combined).replace("_", " ").replace("-", " ")
    return FORBIDDEN_PATH in normalized(decoded)


@dataclass(frozen=True)
class WorkbookMetadata:
    engineer_email: str = ""
    supervisors_email: str = ""
    status: str = ""
    project_id: str = ""


def _neighbor_values(ws: Any, row: int, column: int) -> list[Any]:
    values = [ws.cell(row, column).value]
    for offset in range(1, 7):
        values.append(ws.cell(row, column + offset).value)
    return values


def extract_workbook_metadata(content: bytes) -> WorkbookMetadata:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    try:
        if "Datos" not in workbook.sheetnames:
            return WorkbookMetadata()
        ws = workbook["Datos"]
        engineer = ""
        supervisors = ""
        status = ""
        project_id = ""
        for row in ws.iter_rows():
            for cell in row:
                label = normalized(cell.value)
                if not label:
                    continue
                values = _neighbor_values(ws, cell.row, cell.column)
                joined = " ".join(str(value or "") for value in values)
                if not engineer and ("ingeniero residente" in label or label == "ingenieroemail"):
                    engineer = normalize_emails(joined)
                elif not supervisors and ("supervisores" in label or label == "supervisoresemail"):
                    supervisors = normalize_emails(joined)
                elif not status and label in {"estadogantt", "estado gantt", "statusgantt", "status gantt", "status"}:
                    for value in values[1:]:
                        candidate = str(value or "").strip()
                        if candidate:
                            status = candidate
                            break
                elif not project_id and label in {"proyectoid", "proyecto id"}:
                    for value in values[1:]:
                        candidate = str(value or "").strip()
                        if candidate:
                            project_id = candidate
                            break
        return WorkbookMetadata(
            engineer_email=engineer.split(";")[0] if engineer else "",
            supervisors_email=supervisors,
            status=status,
            project_id=project_id,
        )
    finally:
        workbook.close()


@dataclass(frozen=True)
class ControlRecord:
    item_id: str
    title: str = ""
    project_id: str = ""
    project_name: str = ""
    budget_link: str = ""
    gantt_link: str = ""
    budget_identifier: str = ""
    gantt_identifier: str = ""
    engineer_email: str = ""
    supervisors_email: str = ""
    state: str = ""
    assignment_date: datetime | None = None
    deadline: datetime | None = None
    review_date: datetime | None = None
    permission_granted: bool = False
    assignment_email_sent: bool = False
    warning1_sent: bool = False
    warning2_sent: bool = False
    expiration_notified: bool = False
    tracking_attempts: int = 0
    version_requested: bool = False
    current_version: str = ""
    version_link: str = ""
    version_identifier: str = ""
    versioned_at: datetime | None = None
    version_attempts: int = 0


@dataclass(frozen=True)
class VersionArtifact:
    identifier: str
    web_url: str
    file_name: str
    created: bool
    version: str = "v1.0"
    reasons: tuple[str, ...] = ()


@dataclass(frozen=True)
class Notification:
    kind: str
    to: str
    cc: str
    subject: str
    body: str


class AutomationBackend(Protocol):
    def patch_control(self, item_id: str, updates: dict[str, Any]) -> None: ...

    def load_metadata(self, record: ControlRecord, *, prefer_gantt: bool = True) -> WorkbookMetadata: ...

    def grant_edit_access(self, record: ControlRecord, email: str) -> None: ...

    def notification_exists(self, item_id: str, kind: str) -> bool: ...

    def queue_notification(self, record: ControlRecord, notification: Notification) -> None: ...

    def create_initial_version(self, record: ControlRecord) -> VersionArtifact: ...


def assignment_notification(record: ControlRecord, assigned: datetime, deadline: datetime) -> Notification:
    subject = f"Asignación de Gantt - {record.project_id} {record.project_name}".strip()
    body = (
        "Hola,\n\n"
        f"Se te ha asignado el Gantt del proyecto {record.project_name or record.project_id}.\n\n"
        f"Archivo de trabajo:\n{record.gantt_link}\n\n"
        f"Fecha de asignación: {assigned.date().isoformat()}\n"
        f"Fecha límite: {deadline.date().isoformat()}\n\n"
        "Por favor trabaja siempre sobre este mismo archivo. No crees copias manuales, "
        "no cambies el nombre del archivo y no lo muevas de carpeta.\n\n"
        "Cuando termines, cambia el status del Gantt a:\n\n"
        "En revisión inicial\n\nSaludos."
    )
    return Notification("AsignacionGantt", record.engineer_email, "", subject, body)


def tracking_notification(record: ControlRecord, kind: str, days: int) -> Notification:
    base = f"{record.project_id} {record.project_name}".strip()
    if kind == "Advertencia1":
        subject = f"Advertencia 1 - Gantt pendiente - {base}"
        body = (
            f"Han transcurrido {days} días desde la asignación del Gantt de {base}.\n\n"
            f"Archivo de trabajo:\n{record.gantt_link}\n\n"
            "Esta es la primera de dos advertencias."
        )
        return Notification(kind, record.engineer_email, "", subject, body)
    if kind == "Advertencia2":
        subject = f"Advertencia 2 - Gantt pendiente - {base}"
        body = (
            f"Han transcurrido {days} días desde la asignación del Gantt de {base}.\n\n"
            f"Archivo de trabajo:\n{record.gantt_link}\n\n"
            "Esta es la segunda y última advertencia antes del vencimiento."
        )
        return Notification(kind, record.engineer_email, record.supervisors_email, subject, body)
    recipients = record.supervisors_email or record.engineer_email
    cc = record.engineer_email if record.supervisors_email else ""
    subject = f"Gantt vencido - {base}"
    body = (
        f"El Gantt de {base} llegó al día {days} sin pasar a revisión inicial.\n\n"
        f"Archivo de trabajo:\n{record.gantt_link}\n\n"
        "El registro fue marcado como Vencido."
    )
    return Notification("Vencimiento", recipients, cc, subject, body)


def apply_notification_delivery_mode(notification: Notification) -> Notification:
    mode = normalized(os.getenv("NOTIFICATION_DELIVERY_MODE", "test"))
    if mode == "live":
        return notification
    test_recipient = normalize_emails(
        os.getenv("NOTIFICATION_TEST_RECIPIENT", "auto.sys@prowallpanama.com")
    )
    if not is_valid_email(test_recipient):
        raise RuntimeError(
            "NOTIFICATION_TEST_RECIPIENT debe contener un único correo válido en modo test."
        )
    audit = (
        "\n\n--- MODO PRUEBA ---\n"
        f"Destinatario real previsto: {notification.to or '(vacío)'}\n"
        f"CC real previsto: {notification.cc or '(vacío)'}\n"
        "Este correo fue redirigido y no se envió a destinatarios reales."
    )
    return replace(
        notification,
        to=test_recipient,
        cc="",
        subject=f"[PRUEBA] {notification.subject}",
        body=notification.body + audit,
    )


def days_since(assigned: datetime, now: datetime) -> int:
    return max(0, (now.astimezone(timezone.utc).date() - assigned.astimezone(timezone.utc).date()).days)


def due_tracking_kind(record: ControlRecord, now: datetime) -> str:
    if not record.assignment_date or normalized(record.state) not in ACTIVE_TRACKING_STATES:
        return ""
    elapsed = days_since(record.assignment_date, now)
    if elapsed >= 9 and not record.expiration_notified:
        return "Vencimiento"
    if elapsed >= 6 and not record.warning2_sent:
        return "Advertencia2"
    if elapsed >= 3 and not record.warning1_sent:
        return "Advertencia1"
    return ""


class AutomationService:
    def __init__(self, backend: AutomationBackend, now: datetime | None = None) -> None:
        self.backend = backend
        self.now = now or datetime.now(timezone.utc)

    def sync_excel_status(self, record: ControlRecord, metadata: WorkbookMetadata) -> ControlRecord:
        updates: dict[str, Any] = {
            "StatusExcel": metadata.status,
            "FechaLecturaStatusExcel": iso_utc(self.now),
            "UltimoTrackingRun": iso_utc(self.now),
        }
        if normalized(metadata.status) == normalized(REVIEW_STATE):
            review_date = record.review_date or self.now
            updates.update(
                {
                    "EstadoGantt": REVIEW_STATE,
                    "FechaEnvioRevision": iso_utc(review_date),
                }
            )
            if record.assignment_date:
                updates["DiasParaCompletar"] = days_since(record.assignment_date, review_date)
            record = replace(record, state=REVIEW_STATE, review_date=review_date)
        self.backend.patch_control(record.item_id, updates)
        return record

    def version(self, record: ControlRecord) -> ControlRecord:
        if not record.version_requested:
            return record
        if contains_forbidden_path(
            record.gantt_link,
            record.version_link,
            record.budget_link,
        ):
            self.backend.patch_control(
                record.item_id,
                {
                    "UltimoErrorVersionado": (
                        "Ruta rechazada: Proyectos Terminados está fuera del procesamiento."
                    ),
                    "VersionadoIntentos": record.version_attempts + 1,
                },
            )
            return record

        if normalized(record.state) != normalized(REVIEW_STATE):
            self.backend.patch_control(
                record.item_id,
                {
                    "UltimoErrorVersionado": (
                        "SolicitarVersionado requiere EstadoGantt = En revisión inicial."
                    ),
                    "VersionadoIntentos": record.version_attempts + 1,
                },
            )
            return record

        try:
            artifact = self.backend.create_initial_version(record)
        except Exception as exc:
            self.backend.patch_control(
                record.item_id,
                {
                    "UltimoErrorVersionado": sanitize_error(exc)[:500],
                    "VersionadoIntentos": record.version_attempts + 1,
                },
            )
            return record

        updates = {
            "SolicitarVersionado": False,
            "VersionActual": artifact.version,
            "GanttVersionLink": artifact.web_url,
            "GanttVersionIdentifier": artifact.identifier,
            "FechaUltimoVersionado": iso_utc(self.now),
            "FechaAprobacion": iso_utc(self.now),
            "EstadoGantt": "Aprobado / Versionado",
            "UltimoErrorVersionado": "",
            "MotivoUltimoVersionado": " | ".join(artifact.reasons),
        }
        self.backend.patch_control(record.item_id, updates)
        print(
            f"Versioned control={record.item_id} version={artifact.version} "
            f"file={artifact.file_name} created={artifact.created}"
        )
        return replace(
            record,
            version_requested=False,
            current_version=artifact.version,
            version_link=artifact.web_url,
            version_identifier=artifact.identifier,
            versioned_at=self.now,
            state="Aprobado / Versionado",
        )

    def assign(self, record: ControlRecord) -> ControlRecord:
        state = normalized(record.state)
        if state in CLOSED_STATES or state == normalized(REVIEW_STATE):
            return record
        if contains_forbidden_path(record.gantt_link, record.budget_link):
            self.backend.patch_control(
                record.item_id,
                {
                    "EstadoGantt": "Requiere revisión manual",
                    "UltimoErrorTracking": "Ruta rechazada: Proyectos Terminados está fuera del procesamiento.",
                    "UltimoTrackingRun": iso_utc(self.now),
                },
            )
            return replace(record, state="Requiere revisión manual")

        engineer = normalize_emails(record.engineer_email).split(";")[0] if record.engineer_email else ""
        supervisors = normalize_emails(record.supervisors_email)
        metadata_error = ""
        if not engineer or not supervisors:
            try:
                metadata = self.backend.load_metadata(record, prefer_gantt=False)
                engineer = engineer or metadata.engineer_email
                supervisors = supervisors or metadata.supervisors_email
            except Exception as exc:
                metadata_error = sanitize_error(exc)
            contact_updates: dict[str, Any] = {"UltimoTrackingRun": iso_utc(self.now)}
            if engineer:
                contact_updates["IngenieroEmail"] = engineer
            if supervisors:
                contact_updates["SupervisoresEmail"] = supervisors
            self.backend.patch_control(record.item_id, contact_updates)
            record = replace(record, engineer_email=engineer, supervisors_email=supervisors)

        if not is_valid_email(engineer):
            self.backend.patch_control(
                record.item_id,
                {
                    "EstadoGantt": "Pendiente de asignación",
                    "UltimoErrorTracking": (
                        "IngenieroEmail vacío o inválido; no se otorgó permiso ni se encoló correo."
                        + (f" Detalle de Datos: {metadata_error}" if metadata_error else "")
                    )[:500],
                    "TrackingIntentos": record.tracking_attempts + 1,
                    "UltimoTrackingRun": iso_utc(self.now),
                },
            )
            return replace(record, state="Pendiente de asignación")
        if not record.gantt_link and not record.gantt_identifier:
            self.backend.patch_control(
                record.item_id,
                {
                    "EstadoGantt": "Requiere revisión manual",
                    "UltimoErrorTracking": "GanttWorkingLink/GanttWorkingIdentifier vacío.",
                    "TrackingIntentos": record.tracking_attempts + 1,
                    "UltimoTrackingRun": iso_utc(self.now),
                },
            )
            return replace(record, state="Requiere revisión manual")

        assignment_queued = self.backend.notification_exists(record.item_id, "AsignacionGantt")
        if assignment_queued and not record.permission_granted:
            # El dispatcher siempre concede permiso antes de crear esta notificación.
            # Esto mantiene idempotencia cuando la lista antigua aún no tiene
            # PermisoGanttOtorgado y Graph no permite ampliar su esquema.
            record = replace(record, permission_granted=True)

        if not record.permission_granted:
            try:
                self.backend.grant_edit_access(record, engineer)
            except Exception as exc:
                error = sanitize_error(exc)
                self.backend.patch_control(
                    record.item_id,
                    {
                        "UltimoErrorTracking": f"No se pudo compartir el Gantt: {error}",
                        "TrackingIntentos": record.tracking_attempts + 1,
                        "UltimoTrackingRun": iso_utc(self.now),
                    },
                )
                return record
            self.backend.patch_control(
                record.item_id,
                {
                    "PermisoGanttOtorgado": True,
                    "FechaPermisoOtorgado": iso_utc(self.now),
                    "UltimoErrorTracking": "",
                },
            )
            record = replace(record, permission_granted=True)

        assigned = record.assignment_date or self.now
        deadline = record.deadline or (assigned + timedelta(days=9))
        target_state = record.state if normalized(record.state) in ACTIVE_TRACKING_STATES else "Asignado"
        self.backend.patch_control(
            record.item_id,
            {
                "EstadoGantt": target_state,
                "FechaAsignacion": iso_utc(assigned),
                "FechaLimite": iso_utc(deadline),
                "UltimoTrackingRun": iso_utc(self.now),
                "UltimoErrorTracking": "",
            },
        )
        record = replace(record, state=target_state, assignment_date=assigned, deadline=deadline)
        if not record.assignment_email_sent and not assignment_queued:
            self.backend.queue_notification(record, assignment_notification(record, assigned, deadline))
        return record

    def track(self, record: ControlRecord) -> ControlRecord:
        kind = due_tracking_kind(record, self.now)
        self.backend.patch_control(record.item_id, {"UltimoTrackingRun": iso_utc(self.now)})
        if not kind:
            return record
        if self.backend.notification_exists(record.item_id, kind):
            if kind == "Vencimiento":
                self.backend.patch_control(record.item_id, {"EstadoGantt": "Vencido"})
                return replace(record, state="Vencido")
            return record
        elapsed = days_since(record.assignment_date, self.now) if record.assignment_date else 0
        if kind == "Vencimiento":
            self.backend.queue_notification(record, tracking_notification(record, kind, elapsed))
            self.backend.patch_control(record.item_id, {"EstadoGantt": "Vencido"})
            return replace(record, state="Vencido")
        self.backend.queue_notification(record, tracking_notification(record, kind, elapsed))
        return record


class SharePointBackend:
    def __init__(self, token: str, site: dict[str, Any], *, ensure_schema: bool = True) -> None:
        self.token = token
        self.site = site
        self.site_id = str(site["id"])
        self.control_list = resolve_list(token, self.site_id, CONTROL_LIST_NAME)
        self.queue_list = resolve_list(token, self.site_id, QUEUE_LIST_NAME)
        if ensure_schema:
            self._ensure_columns(self.control_list, CONTROL_COLUMNS)
        self.notification_list = self._resolve_or_create_notification_list(ensure_schema)
        self.control_fields = build_field_map(list_columns(token, self.site_id, self.control_list["id"]))
        self.queue_fields = build_field_map(list_columns(token, self.site_id, self.queue_list["id"]))
        self.notification_fields = build_field_map(
            list_columns(token, self.site_id, self.notification_list["id"])
        )
        self._notification_keys: set[tuple[str, str]] | None = None

    def _ensure_columns(self, target_list: dict[str, Any], definitions: dict[str, dict[str, Any]]) -> None:
        existing = build_field_map(list_columns(self.token, self.site_id, target_list["id"]))
        missing = [name for name in definitions if not pick_field(existing, (name,))]
        for index, name in enumerate(missing):
            column_type = definitions[name]
            try:
                graph_post(
                    self.token,
                    f"{GRAPH_BASE}/sites/{self.site_id}/lists/{target_list['id']}/columns",
                    {"name": name, "displayName": name, **column_type},
                )
                print(f"Created SharePoint column {target_list['displayName']}.{name}")
            except RuntimeError as exc:
                if "403" not in str(exc) and "accessDenied" not in str(exc):
                    raise
                print(
                    "WARNING: no se pudieron crear columnas en "
                    f"{target_list['displayName']}. Faltan: {', '.join(missing[index:])}. "
                    "Conceder Sites.Manage.All o crearlas manualmente."
                )
                break

    def _resolve_or_create_notification_list(self, ensure_schema: bool) -> dict[str, Any]:
        created = False
        try:
            target = resolve_list(self.token, self.site_id, NOTIFICATION_LIST_NAME)
        except RuntimeError:
            if not ensure_schema:
                raise
            try:
                target = graph_post(
                    self.token,
                    f"{GRAPH_BASE}/sites/{self.site_id}/lists",
                    {
                        "displayName": NOTIFICATION_LIST_NAME,
                        "columns": [
                            {"name": name, "displayName": name, **column_type}
                            for name, column_type in NOTIFICATION_COLUMNS.items()
                        ],
                        "list": {"template": "genericList"},
                    },
                )
            except RuntimeError as exc:
                if "403" not in str(exc) and "accessDenied" not in str(exc):
                    raise
                raise RuntimeError(
                    f"No existe {NOTIFICATION_LIST_NAME} y Microsoft Graph no tiene permiso "
                    "para crearla. Crear la lista según docs/power_automate_tracking_flows.md "
                    "o conceder Sites.Manage.All con consentimiento de administrador."
                ) from exc
            created = True
            print(f"Created SharePoint list {NOTIFICATION_LIST_NAME}")
        if ensure_schema and not created:
            self._ensure_columns(target, NOTIFICATION_COLUMNS)
        return target

    def _get(self, fields: dict[str, Any], field_map: dict[str, str], *names: str) -> Any:
        field_name = pick_field(field_map, tuple(names))
        return fields.get(field_name) if field_name else None

    def _map_updates(self, field_map: dict[str, str], updates: dict[str, Any]) -> dict[str, Any]:
        mapped: dict[str, Any] = {}
        aliases = {
            "EstadoGantt": ("EstadoGantt", "EstadoGannt"),
            "GanttWorkingLink": ("GanttWorkingLink", "GanntWorkingLink"),
            "FechaGanttGenerado": ("FechaGanttGenerado", "FechaGanntGenerado"),
            "VersionActual": ("VersionActual", "VersionadoActual"),
            "GanttVersionLink": ("GanttVersionLink", "GanntVersionLink"),
            "GanttVersionIdentifier": (
                "GanttVersionIdentifier",
                "GanntVersionIdentifier",
            ),
        }
        for canonical, value in updates.items():
            field_name = pick_field(field_map, aliases.get(canonical, (canonical,)))
            if field_name:
                mapped[field_name] = value
        return mapped

    def list_items(self, target_list: dict[str, Any], top: int = 500) -> list[dict[str, Any]]:
        url = (
            f"{GRAPH_BASE}/sites/{self.site_id}/lists/{target_list['id']}/items"
            f"?$top={top}&$orderby=createdDateTime asc&expand=fields"
        )
        result: list[dict[str, Any]] = []
        while url:
            data = graph_get(self.token, url)
            result.extend(data.get("value") or [])
            url = str(data.get("@odata.nextLink") or "")
        return result

    def control_records(self) -> list[ControlRecord]:
        records: list[ControlRecord] = []
        for item in self.list_items(self.control_list):
            fields = item.get("fields") or {}
            get = lambda *names: self._get(fields, self.control_fields, *names)
            records.append(
                ControlRecord(
                    item_id=str(item.get("id") or ""),
                    title=str(get("Title") or ""),
                    project_id=str(get("ProyectoID") or ""),
                    project_name=str(get("NombreProyecto") or ""),
                    budget_link=link_text(get("PresupuestoLink")),
                    gantt_link=link_text(get("GanttWorkingLink", "GanntWorkingLink")),
                    budget_identifier=str(get("PresupuestoIdentifier") or ""),
                    gantt_identifier=str(get("GanttWorkingIdentifier") or ""),
                    engineer_email=str(get("IngenieroEmail") or ""),
                    supervisors_email=str(get("SupervisoresEmail") or ""),
                    state=str(get("EstadoGantt", "EstadoGannt") or ""),
                    assignment_date=parse_datetime(get("FechaAsignacion")),
                    deadline=parse_datetime(get("FechaLimite")),
                    review_date=parse_datetime(get("FechaEnvioRevision")),
                    permission_granted=parse_bool(get("PermisoGanttOtorgado")),
                    assignment_email_sent=parse_bool(get("CorreoAsignacionEnviado")),
                    warning1_sent=parse_bool(get("Advertencia1Enviada")),
                    warning2_sent=parse_bool(get("Advertencia2Enviada")),
                    expiration_notified=parse_bool(get("VencimientoNotificado")),
                    tracking_attempts=parse_int(get("TrackingIntentos")),
                    version_requested=parse_bool(get("SolicitarVersionado")),
                    current_version=str(
                        get("VersionActual", "VersionadoActual") or ""
                    ),
                    version_link=link_text(
                        get("GanttVersionLink", "GanntVersionLink")
                    ),
                    version_identifier=str(
                        get("GanttVersionIdentifier", "GanntVersionIdentifier")
                        or ""
                    ),
                    versioned_at=parse_datetime(get("FechaUltimoVersionado")),
                    version_attempts=parse_int(get("VersionadoIntentos")),
                )
            )
        return records

    def patch_control(self, item_id: str, updates: dict[str, Any]) -> None:
        mapped = self._map_updates(self.control_fields, updates)
        if mapped:
            graph_patch(
                self.token,
                f"{GRAPH_BASE}/sites/{self.site_id}/lists/{self.control_list['id']}/items/{item_id}/fields",
                mapped,
            )

    @staticmethod
    def _share_id(url: str) -> str:
        encoded = base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii").rstrip("=")
        return "u!" + encoded

    def resolve_drive_item(self, link: str, identifier: str) -> dict[str, Any]:
        if identifier:
            try:
                return graph_get(
                    self.token,
                    f"{GRAPH_BASE}/sites/{self.site_id}/drive/items/{identifier}",
                )
            except RuntimeError:
                pass
        if link:
            return graph_get(self.token, f"{GRAPH_BASE}/shares/{self._share_id(link)}/driveItem")
        raise RuntimeError("No hay link o identificador para resolver el archivo.")

    def _download_drive_item(self, item: dict[str, Any]) -> bytes:
        drive_id = str((item.get("parentReference") or {}).get("driveId") or "")
        item_id = str(item.get("id") or "")
        if not drive_id or not item_id:
            raise RuntimeError("El archivo no devolvió driveId/itemId.")
        response = requests.get(
            f"{GRAPH_BASE}/drives/{drive_id}/items/{item_id}/content",
            headers={"Authorization": f"Bearer {self.token}"},
            timeout=120,
        )
        if response.status_code >= 400:
            raise RuntimeError(f"Graph download {response.status_code}: {response.text[:1000]}")
        return response.content

    def load_metadata(self, record: ControlRecord, *, prefer_gantt: bool = True) -> WorkbookMetadata:
        candidates = (
            [
                (record.gantt_link, record.gantt_identifier),
                (record.budget_link, record.budget_identifier),
            ]
            if prefer_gantt
            else [
                (record.budget_link, record.budget_identifier),
                (record.gantt_link, record.gantt_identifier),
            ]
        )
        errors: list[str] = []
        combined = WorkbookMetadata()
        for link, identifier in candidates:
            if not link and not identifier:
                continue
            if contains_forbidden_path(link):
                raise RuntimeError("Ruta rechazada: Proyectos Terminados.")
            try:
                metadata = extract_workbook_metadata(
                    self._download_drive_item(self.resolve_drive_item(link, identifier))
                )
                combined = WorkbookMetadata(
                    engineer_email=combined.engineer_email or metadata.engineer_email,
                    supervisors_email=combined.supervisors_email or metadata.supervisors_email,
                    status=combined.status or metadata.status,
                    project_id=combined.project_id or metadata.project_id,
                )
                if combined.engineer_email and (combined.status or not prefer_gantt):
                    return combined
            except Exception as exc:
                errors.append(str(exc))
        if not any((combined.engineer_email, combined.supervisors_email, combined.status, combined.project_id)):
            raise RuntimeError("No se pudo leer la hoja Datos: " + " | ".join(errors[:2]))
        return combined

    def grant_edit_access(self, record: ControlRecord, email: str) -> None:
        item = self.resolve_drive_item(record.gantt_link, record.gantt_identifier)
        if contains_forbidden_path(item.get("webUrl"), record.gantt_link):
            raise RuntimeError("Ruta rechazada: Proyectos Terminados.")
        drive_id = str((item.get("parentReference") or {}).get("driveId") or "")
        item_id = str(item.get("id") or "")
        if not drive_id or not item_id:
            raise RuntimeError("No se pudo resolver driveId/itemId para compartir.")
        graph_post(
            self.token,
            f"{GRAPH_BASE}/drives/{drive_id}/items/{item_id}/invite",
            {
                "recipients": [{"email": email}],
                "roles": ["write"],
                "requireSignIn": True,
                "sendInvitation": False,
            },
        )

    def create_initial_version(self, record: ControlRecord) -> VersionArtifact:
        source = self.resolve_drive_item(record.gantt_link, record.gantt_identifier)
        source_url = str(source.get("webUrl") or "")
        parent = source.get("parentReference") or {}
        parent_path_raw = str(parent.get("path") or "")
        if contains_forbidden_path(source_url, parent_path_raw):
            raise RuntimeError("Ruta rechazada: Proyectos Terminados.")
        if "root:" not in parent_path_raw:
            raise RuntimeError("No se pudo determinar la carpeta del Gantt WORKING.")
        working_folder = parent_path_raw.split("root:", 1)[1].strip("/")
        normalized_folder = normalized(
            requests.utils.unquote(working_folder).replace("_", " ").replace("-", " ")
        )
        if "proyectos/proyectos activos" not in working_folder.casefold():
            raise RuntimeError("El Gantt WORKING no está dentro de Proyectos Activos.")
        if not normalized_folder.endswith("gantts/working"):
            raise RuntimeError("El archivo no está dentro de /gantts/working/.")

        project_id = re.sub(r'[<>:"/\\|?*]+', "_", record.project_id.strip())
        if not project_id:
            raise RuntimeError("ProyectoID vacío; no se puede nombrar la versión.")
        gantts_folder = working_folder.rsplit("/", 1)[0]
        version_folder = f"{gantts_folder}/versionados"
        v1_name = f"{project_id}_gantt_v1.0.xlsx"
        v1_path = f"{version_folder}/{v1_name}"
        working_content = self._download_drive_item(source)
        baseline_content = None
        if path_exists(self.token, self.site_id, v1_path):
            baseline_item = get_drive_item_by_path(
                self.token,
                self.site_id,
                v1_path,
            )
            baseline_content = self._download_drive_item(baseline_item)

        decision = decide_version(working_content, baseline_content)
        version_name = f"{project_id}_gantt_{decision.version}.xlsx"
        version_path = f"{version_folder}/{version_name}"
        created = not path_exists(self.token, self.site_id, version_path)
        ensure_drive_folder(self.token, self.site_id, version_folder)
        target = graph_put_bytes(
            self.token,
            (
                f"{GRAPH_BASE}/sites/{self.site_id}/drive/root:/"
                f"{encoded_drive_path(version_path)}:/content"
            ),
            working_content,
        )
        return VersionArtifact(
            identifier=str(target.get("id") or ""),
            web_url=str(target.get("webUrl") or ""),
            file_name=str(target.get("name") or version_name),
            created=created,
            version=decision.version,
            reasons=decision.reasons,
        )

    def _load_notification_keys(self) -> set[tuple[str, str]]:
        keys: set[tuple[str, str]] = set()
        for item in self.list_items(self.notification_list):
            fields = item.get("fields") or {}
            control_id = str(
                self._get(fields, self.notification_fields, "RelatedControlItemID") or ""
            )
            kind = str(self._get(fields, self.notification_fields, "TipoNotificacion") or "")
            if control_id and kind:
                keys.add((control_id, kind.casefold()))
        return keys

    def notification_exists(self, item_id: str, kind: str) -> bool:
        if self._notification_keys is None:
            self._notification_keys = self._load_notification_keys()
        return (item_id, kind.casefold()) in self._notification_keys

    def queue_notification(self, record: ControlRecord, notification: Notification) -> None:
        notification = apply_notification_delivery_mode(notification)
        fields = self._map_updates(
            self.notification_fields,
            {
                "Title": f"{record.item_id}:{notification.kind}",
                "ProyectoID": record.project_id,
                "TipoNotificacion": notification.kind,
                "EstadoNotificacion": "Pendiente",
                "To": notification.to,
                "Cc": notification.cc,
                "Subject": notification.subject,
                "Body": notification.body,
                "RelatedControlItemID": record.item_id,
                "Intentos": 0,
                "FechaCreacion": iso_utc(datetime.now(timezone.utc)),
                "Notas": "Creado por automation_dispatcher.py",
            },
        )
        if not fields.get(pick_field(self.notification_fields, ("To",)), ""):
            raise RuntimeError(f"Notificación {notification.kind} sin destinatario.")
        graph_post(
            self.token,
            f"{GRAPH_BASE}/sites/{self.site_id}/lists/{self.notification_list['id']}/items",
            {"fields": fields},
        )
        if self._notification_keys is None:
            self._notification_keys = set()
        self._notification_keys.add((record.item_id, notification.kind.casefold()))
        print(f"Queued notification control={record.item_id} type={notification.kind}")

    def _find_control_for_metadata(self, metadata: WorkbookMetadata, queue_fields: dict[str, Any]) -> ControlRecord | None:
        identifier = str(
            self._get(queue_fields, self.queue_fields, "FileIdentifier", "FileID") or ""
        )
        file_link = link_text(self._get(queue_fields, self.queue_fields, "FileLink"))
        for record in self.control_records():
            if identifier and record.gantt_identifier and identifier == record.gantt_identifier:
                return record
            if file_link and record.gantt_link and file_link.casefold() == record.gantt_link.casefold():
                return record
            if metadata.project_id and record.project_id.casefold() == metadata.project_id.casefold():
                return record
        return None

    def process_status_events(self, now: datetime, max_items: int = 20) -> int:
        processed = 0
        for item in self.list_items(self.queue_list):
            if processed >= max_items:
                break
            fields = item.get("fields") or {}
            event_type = str(self._get(fields, self.queue_fields, "EventType") or "")
            state = str(self._get(fields, self.queue_fields, "Estado") or "")
            if normalized(event_type) != "gantt_working_modificado" or normalized(state) != "pendiente":
                continue
            item_id = str(item.get("id") or "")
            folder_path = str(self._get(fields, self.queue_fields, "FolderPath") or "")
            file_name = str(
                self._get(fields, self.queue_fields, "FileName", "Filename", "Title") or ""
            )
            if contains_forbidden_path(folder_path, file_name):
                self._patch_queue(
                    item_id,
                    {
                        "Estado": "RequiereRevision",
                        "UltimoError": "Ruta rechazada: Proyectos Terminados.",
                    },
                )
                continue
            self._patch_queue(item_id, {"Estado": "Procesando", "UltimoError": ""})
            try:
                identifier = str(
                    self._get(fields, self.queue_fields, "FileIdentifier", "FileID") or ""
                )
                file_link = link_text(self._get(fields, self.queue_fields, "FileLink"))
                if identifier or file_link:
                    drive_item = self.resolve_drive_item(file_link, identifier)
                    content = self._download_drive_item(drive_item)
                else:
                    drive_path = drive_path_from_queue_fields(
                        {
                            "Filename": file_name,
                            "FolderPath": folder_path,
                        }
                    )
                    response = requests.get(
                        f"{GRAPH_BASE}/sites/{self.site_id}/drive/root:/{requests.utils.quote(drive_path, safe='/')}:/content",
                        headers={"Authorization": f"Bearer {self.token}"},
                        timeout=120,
                    )
                    if response.status_code >= 400:
                        raise RuntimeError(f"Graph download {response.status_code}: {response.text[:1000]}")
                    content = response.content
                metadata = extract_workbook_metadata(content)
                record = self._find_control_for_metadata(metadata, fields)
                if not record:
                    raise RuntimeError("No se encontró Control_Gantt_Asignaciones para el Gantt modificado.")
                AutomationService(self, now).sync_excel_status(record, metadata)
                self._patch_queue(
                    item_id,
                    {
                        "Estado": "Procesado",
                        "FechaProcesado": iso_utc(now),
                        "ProyectoID": record.project_id,
                        "Notas": (
                            f"StatusExcel={metadata.status or 'sin campo definido'}; "
                            "no se creó ninguna versión."
                        ),
                    },
                )
                processed += 1
            except Exception as exc:
                error = sanitize_error(exc)
                self._patch_queue(
                    item_id,
                    {
                        "Estado": "Error",
                        "UltimoError": error[:240],
                        "Notas": "Error leyendo status del Gantt WORKING.",
                    },
                )
        return processed

    def _patch_queue(self, item_id: str, updates: dict[str, Any]) -> None:
        mapped = self._map_updates(self.queue_fields, updates)
        if mapped:
            graph_patch(
                self.token,
                f"{GRAPH_BASE}/sites/{self.site_id}/lists/{self.queue_list['id']}/items/{item_id}/fields",
                mapped,
            )


def run_system1(top: int, max_items: int, item_id: str = "") -> None:
    command = [
        sys.executable,
        str(SCRIPTS / "sistema1_poll_queue.py"),
        "--top",
        str(top),
        "--max-items",
        str(max_items),
        "--process",
    ]
    if item_id:
        command.extend(["--item-id", item_id])
    result = subprocess.run(command, cwd=ROOT, check=False)
    if result.returncode:
        raise RuntimeError(f"Sistema 1 terminó con código {result.returncode}.")


def run_system2(
    backend: SharePointBackend,
    now: datetime,
    max_status_items: int,
    control_item_id: str = "",
) -> dict[str, int]:
    isolated_run = control_item_id.strip()
    status_events = 0 if isolated_run else backend.process_status_events(now, max_status_items)
    service = AutomationService(backend, now)
    assigned = 0
    tracked = 0
    versioned = 0
    version_errors = 0
    errors = 0
    records = backend.control_records()
    if isolated_run:
        records = [record for record in records if record.item_id == isolated_run]
        if not records:
            raise RuntimeError(
                f"No existe el item {isolated_run} en Control_Gantt_Asignaciones."
            )
        print(f"System 2 isolated test: control item {isolated_run}")
    for record in records:
        if record.version_requested:
            updated = service.version(record)
            if (
                normalized(updated.current_version) in {"v1.0", "v2.0"}
                and normalized(updated.state) == "aprobado / versionado"
            ):
                versioned += 1
            else:
                version_errors += 1
            continue
        if (
            normalized(record.state) in CLOSED_STATES
            or normalized(record.state) == normalized(REVIEW_STATE)
            or normalized(record.state) == "requiere revision manual"
        ):
            continue
        try:
            updated = service.assign(record)
            if normalized(updated.state) in ACTIVE_TRACKING_STATES:
                assigned += 1
                service.track(updated)
                tracked += 1
        except Exception as exc:
            errors += 1
            error = sanitize_error(exc)
            backend.patch_control(
                record.item_id,
                {
                    "UltimoErrorTracking": error[:240],
                    "TrackingIntentos": record.tracking_attempts + 1,
                    "UltimoTrackingRun": iso_utc(now),
                },
            )
            print(f"ERROR control item {record.item_id}: {error}", file=sys.stderr)
    return {
        "control_items": len(records),
        "assigned_or_active": assigned,
        "tracked": tracked,
        "versioned": versioned,
        "version_errors": version_errors,
        "status_events": status_events,
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Dispatcher único de automatización SharePoint.")
    parser.add_argument("--top", type=int, default=50)
    parser.add_argument("--max-items", type=int, default=5)
    parser.add_argument("--max-status-items", type=int, default=20)
    parser.add_argument("--item-id", default="")
    parser.add_argument(
        "--control-item-id",
        default="",
        help="Limita Sistema 2 a un item de Control_Gantt_Asignaciones.",
    )
    parser.add_argument("--skip-system1", action="store_true")
    parser.add_argument(
        "--skip-system2",
        action="store_true",
        help="Ejecuta solo Sistema 1; no comparte archivos, encola correos ni procesa tracking.",
    )
    parser.add_argument("--skip-schema", action="store_true")
    args = parser.parse_args()

    system1_error = ""
    if not args.skip_system1:
        try:
            run_system1(max(1, args.top), max(1, args.max_items), args.item_id.strip())
        except Exception as exc:
            system1_error = sanitize_error(exc)
            print(
                "Sistema 1 terminó con error; Sistema 2 continuará para no bloquear "
                f"asignaciones y tracking: {system1_error}"
            )

    if args.skip_system2:
        print("Sistema 2 omitido: no se procesaron asignaciones, correos, permisos ni tracking.")
        return 1 if system1_error else 0

    settings = load_settings()
    token = acquire_token(settings)
    print_token_diagnostics(token)
    site = resolve_site(token, settings)
    backend = SharePointBackend(token, site, ensure_schema=not args.skip_schema)
    summary = run_system2(
        backend,
        datetime.now(timezone.utc),
        max(1, args.max_status_items),
        args.control_item_id.strip(),
    )
    if system1_error:
        summary["system1_errors"] = 1
    print("Autosys automation dispatcher summary:")
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    if summary["errors"] or summary["version_errors"] or system1_error:
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {sanitize_error(exc)}", file=sys.stderr)
        raise SystemExit(1)
