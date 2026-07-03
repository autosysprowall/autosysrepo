from __future__ import annotations

import io
import hashlib
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


HEADER_SCAN_LIMIT = 30
COST_TOLERANCE = 0.01
VERSION_PATTERN = re.compile(r"^v?(\d+)\.(\d+)$", re.IGNORECASE)
PROJECT_STATUS_LABELS = {
    "estado general del gantt",
    "estado general gantt",
}
PLANNING_HEADERS = (
    "actividad",
    "item",
    "cc",
    "fecha de inicio",
    "fecha de fin",
    "unidad",
    "cantidad",
    "costo unitario",
    "costo total",
)


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().casefold()


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    return None


def _find_gantt_header(ws) -> tuple[int, dict[str, int]]:
    best: tuple[int, dict[str, int]] | None = None
    for row_number in range(1, min(ws.max_row, HEADER_SCAN_LIMIT) + 1):
        headers = {
            _normalized(ws.cell(row_number, column).value): column
            for column in range(1, ws.max_column + 1)
            if _normalized(ws.cell(row_number, column).value)
        }
        if "actividad" in headers and "fecha de fin" in headers:
            if best is None or len(headers) > len(best[1]):
                best = (row_number, headers)
    if best is None:
        raise ValueError("No se pudo detectar el encabezado de la hoja Gantt.")
    return best


def _contractual_end(workbook) -> date | None:
    if "AutosysVersionBaseline" in workbook.sheetnames:
        baseline_ws = workbook["AutosysVersionBaseline"]
        for row in baseline_ws.iter_rows():
            for cell in row:
                label = _normalized(cell.value)
                if "fecha final contractual" not in label:
                    continue
                for offset in range(1, 5):
                    parsed = _as_date(
                        baseline_ws.cell(
                            cell.row,
                            cell.column + offset,
                        ).value
                    )
                    if parsed:
                        return parsed
    if "Datos" not in workbook.sheetnames:
        return None
    ws = workbook["Datos"]
    for row in ws.iter_rows():
        for cell in row:
            label = _normalized(cell.value)
            if "fecha" not in label or not any(
                term in label
                for term in (
                    "fecha final",
                    "fecha de fin",
                    "fecha fin",
                    "fin contractual",
                )
            ):
                continue
            for offset in range(1, 5):
                parsed = _as_date(ws.cell(cell.row, cell.column + offset).value)
                if parsed:
                    return parsed
            parsed = _as_date(ws.cell(cell.row + 1, cell.column).value)
            if parsed:
                return parsed
    return None


def _cost_columns(headers: dict[str, int]) -> dict[str, int]:
    exact = [
        (header, column)
        for header, column in headers.items()
        if header == "costo total"
    ]
    candidates = exact or [
        (header, column)
        for header, column in headers.items()
        if header.startswith("costo total")
        and not any(term in header for term in ("utilidad", "margen"))
    ]
    if not candidates:
        return {}
    header, column = min(candidates, key=lambda item: item[1])
    return {header: column}


def _numeric(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


@dataclass(frozen=True)
class GanttSnapshot:
    costs: dict[str, float]
    baseline_costs: dict[str, float]
    contractual_end: date | None
    latest_activity_end: date | None
    planning_fingerprint: str

    @property
    def total_cost(self) -> float:
        return sum(self.costs.values())

    @property
    def baseline_total_cost(self) -> float | None:
        if not self.baseline_costs:
            return None
        return sum(self.baseline_costs.values())

    @property
    def schedule_overrun(self) -> bool:
        return bool(
            self.contractual_end
            and self.latest_activity_end
            and self.latest_activity_end > self.contractual_end
        )


@dataclass(frozen=True)
class VersionDecision:
    version: str
    cost_increase: bool
    schedule_overrun: bool
    reasons: tuple[str, ...]
    previous_total_cost: float
    working_total_cost: float
    chronology_limit: date | None
    planning_changed: bool


def parse_version(value: str) -> tuple[int, int] | None:
    match = VERSION_PATTERN.fullmatch(str(value or "").strip())
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def format_version(value: tuple[int, int]) -> str:
    return f"v{value[0]}.{value[1]}"


def next_version(current_version: str, major_change: bool) -> str:
    parsed = parse_version(current_version)
    if parsed is None:
        return "v2.0" if major_change else "v1.0"
    major, minor = parsed
    if major_change:
        return f"v{major + 1}.0"
    return f"v{major}.{minor + 1}"


def snapshot_gantt(content: bytes) -> GanttSnapshot:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    try:
        if "Gantt" not in workbook.sheetnames:
            raise ValueError("El archivo no contiene la hoja Gantt.")
        ws = workbook["Gantt"]
        header_row, headers = _find_gantt_header(ws)
        cost_columns = _cost_columns(headers)
        costs = {header: 0.0 for header in cost_columns}
        end_column = headers["fecha de fin"]
        latest_end: date | None = None
        planning_columns = [
            (header, headers[header])
            for header in PLANNING_HEADERS
            if header in headers
        ]
        planning_rows: list[str] = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            activity = ws.cell(row_number, headers["actividad"]).value
            if activity in (None, ""):
                continue
            serialized: list[str] = []
            for header, column in planning_columns:
                value = ws.cell(row_number, column).value
                parsed_date = _as_date(value)
                if parsed_date:
                    text = parsed_date.isoformat()
                elif isinstance(value, float):
                    text = f"{value:.12g}"
                else:
                    text = _normalized(value)
                serialized.append(f"{header}={text}")
            planning_rows.append("|".join(serialized))
            activity_end = _as_date(ws.cell(row_number, end_column).value)
            if activity_end and (latest_end is None or activity_end > latest_end):
                latest_end = activity_end
            for header, column in cost_columns.items():
                value = _numeric(ws.cell(row_number, column).value)
                if value is not None:
                    costs[header] += value
        baseline_costs: dict[str, float] = {}
        if "AutosysVersionBaseline" in workbook.sheetnames:
            baseline_ws = workbook["AutosysVersionBaseline"]
            baseline_candidates: list[tuple[str, float]] = []
            for row_number in range(3, baseline_ws.max_row + 1):
                header = _normalized(baseline_ws.cell(row_number, 1).value)
                value = _numeric(baseline_ws.cell(row_number, 2).value)
                if (
                    header.startswith("costo total")
                    and value is not None
                    and not any(
                        term in header for term in ("utilidad", "margen")
                    )
                ):
                    baseline_candidates.append((header, value))
            exact_baseline = [
                candidate
                for candidate in baseline_candidates
                if candidate[0] == "costo total"
            ]
            selected_baseline = exact_baseline or baseline_candidates[:1]
            baseline_costs.update(selected_baseline)
        return GanttSnapshot(
            costs=costs,
            baseline_costs=baseline_costs,
            contractual_end=_contractual_end(workbook),
            latest_activity_end=latest_end,
            planning_fingerprint=hashlib.sha256(
                "\n".join(planning_rows).encode("utf-8")
            ).hexdigest(),
        )
    finally:
        workbook.close()


def decide_version(
    working_content: bytes,
    previous_version_content: bytes | None,
    *,
    current_version: str = "",
    prior_version_contents: tuple[bytes, ...] = (),
) -> VersionDecision:
    working = snapshot_gantt(working_content)
    reasons: list[str] = []

    previous_snapshot = (
        snapshot_gantt(previous_version_content)
        if previous_version_content is not None
        else None
    )
    previous_total = (
        previous_snapshot.total_cost
        if previous_snapshot is not None
        else working.baseline_total_cost
    )
    planning_changed = bool(
        previous_snapshot is None
        or working.planning_fingerprint
        != previous_snapshot.planning_fingerprint
    )
    if previous_total is None:
        raise ValueError(
            "El Gantt no contiene AutosysVersionBaseline y tampoco existe una "
            "versión anterior. Regenere el WORKING antes de versionar."
        )
    if not working.costs:
        raise ValueError(
            "No existe una columna canónica Costo Total en el Gantt WORKING."
        )

    cost_increase = working.total_cost > previous_total + COST_TOLERANCE
    if cost_increase:
        reasons.append(
            "Costo Total aumentó de "
            f"{previous_total:.2f} a {working.total_cost:.2f}."
        )

    history_contents = prior_version_contents
    if not history_contents and previous_version_content is not None:
        history_contents = (previous_version_content,)
    chronology_candidates = [
        value
        for value in (
            working.contractual_end,
            *(
                snapshot_gantt(content).latest_activity_end
                for content in history_contents
            ),
        )
        if value is not None
    ]
    chronology_limit = (
        max(chronology_candidates) if chronology_candidates else None
    )
    schedule_overrun = bool(
        chronology_limit
        and working.latest_activity_end
        and working.latest_activity_end > chronology_limit
    )
    if schedule_overrun:
        reasons.append(
            "La fecha más tardía aumentó de "
            f"{chronology_limit.isoformat()} a "
            f"{working.latest_activity_end.isoformat()}."
        )

    major_change = cost_increase or schedule_overrun
    if not reasons:
        reasons.append(
            "Costo Total sin aumento y fecha final dentro del máximo ya aprobado."
        )
    return VersionDecision(
        version=next_version(current_version, major_change),
        cost_increase=cost_increase,
        schedule_overrun=schedule_overrun,
        reasons=tuple(reasons),
        previous_total_cost=previous_total,
        working_total_cost=working.total_cost,
        chronology_limit=chronology_limit,
        planning_changed=planning_changed,
    )


def workbook_with_status(content: bytes, status: str) -> bytes:
    workbook = load_workbook(
        io.BytesIO(content),
        data_only=False,
        read_only=False,
    )
    try:
        if "Gantt" not in workbook.sheetnames:
            raise ValueError("El archivo no contiene la hoja Gantt.")
        ws = workbook["Gantt"]
        status_cell = None
        for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row)):
            for cell in row:
                if _normalized(cell.value) in PROJECT_STATUS_LABELS:
                    status_cell = ws.cell(cell.row, cell.column + 1)
                    break
            if status_cell is not None:
                break
        if status_cell is None:
            raise ValueError(
                "No se encontró la celda Estado general del Gantt."
            )
        status_cell.value = status
        validation_found = False
        for validation in ws.data_validations.dataValidation:
            if status_cell.coordinate in validation:
                validation.formula1 = '"Actual,En Progreso,Entregar"'
                validation_found = True
        if not validation_found:
            validation = DataValidation(
                type="list",
                formula1='"Actual,En Progreso,Entregar"',
                allow_blank=False,
            )
            ws.add_data_validation(validation)
            validation.add(status_cell.coordinate)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
