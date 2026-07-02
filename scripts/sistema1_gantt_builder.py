from __future__ import annotations

import calendar
import io
import json
import re
import unicodedata
from copy import copy
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from hashlib import sha1
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from budget_extraction import (
    BudgetCostColumn,
    BudgetExtractionError,
    BudgetExtractionResult,
    detect_header_candidate,
    extract_budget_structure,
)
from sistema1_llm_planner import request_llm_column_mapping, request_llm_plan


HEADER_ROW = 10
WEEKDAY_ROW = 11
DATA_START_ROW = 12
PROJECT_STATUS_LABEL_CELL = "A6"
PROJECT_STATUS_VALUE_CELL = "B6"
PROJECT_STATUS_LABEL = "Estado general del Gantt"
PROJECT_STATUS_DEFAULT = "En Progreso"
PROJECT_STATUS_VALUES = ("Actual", "En Progreso", "Entregar")
START_DATE_COL = 5
END_DATE_COL = 6
DAILY_START_COL = 9
DEFAULT_WINDOW_DAYS = 90
CALENDAR_EXTENSION_DAYS = 365
VERSION_BASELINE_SHEET = "AutosysVersionBaseline"
CATEGORY_ORDER = ("PRELIMINARES", "FABRICA", "CAMPO", "ACABADOS")

TOTAL_TERMS = ("subtotal", "total", "gran total", "itbms", "impuesto", "utilidad", "margen", "precio total")
INDIRECT_TERMS = ("alquiler", "combustible", "baño", "banos", "grúa", "grua", "campamento", "administracion")
RESOURCE_TERMS = ("material", "materiales", "acero", "concreto", "clavo", "alambre", "plastico", "plástico")
WORK_TERMS = (
    "fabricacion",
    "fabricación",
    "instalacion",
    "instalación",
    "construccion",
    "construcción",
    "vaciado",
    "colocacion",
    "colocación",
    "montaje",
    "formaleta",
    "pintura",
    "pasteo",
    "suministro",
    "transporte",
)

MONTHS_ES = {
    1: "ENE",
    2: "FEB",
    3: "MAR",
    4: "ABR",
    5: "MAY",
    6: "JUN",
    7: "JUL",
    8: "AGO",
    9: "SEP",
    10: "OCT",
    11: "NOV",
    12: "DIC",
}
WEEKDAYS_ES = ["L", "M", "X", "J", "V", "S", "D"]


@dataclass(frozen=True)
class ProjectIdentity:
    project_id: str
    project_name: str
    folder_name: str
    gantt_file_name: str
    budget_copy_name: str


@dataclass(frozen=True)
class BudgetRow:
    source_row: int
    line_type: str
    item: Any
    code: Any
    description: str
    quantity: Any
    unit: Any
    amount: Any
    requires_review: bool
    review_reason: str


@dataclass(frozen=True)
class BuildResult:
    output_path: Path
    selected_sheet: str
    header_row: int
    rows_written: int
    review_rows: int
    window_start: date
    window_end: date
    notes: list[str]
    validation: dict[str, Any]


@dataclass(frozen=True)
class LlmOptions:
    enabled: bool = False
    model: str = ""


class GanttReviewRequiredError(RuntimeError):
    """El presupuesto o el archivo generado requiere revisión antes de subirlo."""


def normalize_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    replacements = {
        "á": "a",
        "é": "e",
        "í": "i",
        "ó": "o",
        "ú": "u",
        "ü": "u",
        "ñ": "n",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text)


def material_process_kind(description: str) -> str:
    text = normalize_text(description)
    if "mano de obra" in text or re.search(r"\bmo\b", text):
        return "mano_obra"
    if "material de produccion" in text or "materal de produccion" in text:
        return "material_produccion"
    if "formaleta" in text or "molde" in text:
        return "formaleta"
    if text == "concreto" or text.startswith("concreto "):
        return "concreto"
    if text == "acero" or text.startswith("acero "):
        return "acero"
    if text == "transporte" or text.startswith("transporte "):
        return "transporte"
    return ""


def detect_material_process_budget(rows: list[BudgetRow]) -> bool:
    kinds = {material_process_kind(row.description) for row in rows}
    kinds.discard("")
    return len(kinds) >= 3 and bool(kinds.intersection({"formaleta", "mano_obra"}))


def material_process_bucket(kind: str) -> str:
    return "CAMPO" if kind == "transporte" else "FABRICA"


def enforce_material_process_rows(
    planned_rows: list[BudgetRow],
    source_rows: list[BudgetRow],
) -> tuple[list[BudgetRow], list[str]]:
    source_by_row = {row.source_row: row for row in source_rows}
    protected = {
        row.source_row: material_process_kind(row.description)
        for row in source_rows
        if material_process_kind(row.description)
    }
    if not protected:
        return planned_rows, []

    preserved = 0
    result: list[BudgetRow] = []
    for planned in planned_rows:
        kind = protected.get(planned.source_row)
        source = source_by_row.get(planned.source_row)
        if not kind or not source:
            result.append(planned)
            continue
        preserved += 1
        result.append(
            replace(
                planned,
                line_type=material_process_bucket(kind),
                description=source.description,
                quantity=source.quantity,
                unit=source.unit,
                requires_review=source.requires_review,
                review_reason=(
                    source.review_reason
                    if source.requires_review
                    else "Proceso preservado por estructura de presupuesto basada en materiales."
                ),
            )
        )
    return result, [
        "Estructura detectada: presupuesto por procesos/materiales.",
        f"Procesos de materiales preservados en el Gantt: {preserved}.",
    ]


def display_text(value: Any, max_len: int = 200) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_len]


def is_resident_engineer_row(*values: Any) -> bool:
    text = normalize_text(" ".join(display_text(value, 120) for value in values))
    return any(term in text for term in ("ingeniero residente", "residente obra", "ingeniero residente obra"))


def is_month_unit(value: Any) -> bool:
    text = normalize_text(value).strip(". ")
    return text in {"mes", "meses", "mensual"}


def safe_path_name(value: str, max_len: int = 95) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1F]+', "_", value).strip(" ._")
    name = re.sub(r"\s+", " ", name)
    if not name:
        name = "Proyecto"
    return name[:max_len].strip(" ._")


def safe_file_stem(value: str, max_len: int = 120) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1F]+', "_", value).strip(" ._")
    name = re.sub(r"\s+", " ", name)
    if not name:
        name = "presupuesto"
    return name[:max_len].strip(" ._")


def derive_project_identity(file_name: str) -> ProjectIdentity:
    stem = Path(file_name).stem
    match = re.match(r"(?P<id>\d{4}-\d{3}(?:-[A-Za-z0-9]+)?)\s*[-_ ]\s*(?P<name>.+)", stem)
    if match:
        project_id = match.group("id").upper()
        project_name = safe_path_name(match.group("name"), 80)
    else:
        digest = int(sha1(stem.encode("utf-8")).hexdigest()[:6], 16) % 1000
        project_id = f"PR-{datetime.now().year}-{digest:03d}"
        project_name = safe_path_name(stem, 80)

    folder_name = safe_path_name(f"{project_id}_{project_name}", 120)
    gantt_file_name = f"{project_id}_gantt_WORKING.xlsx"
    budget_copy_name = f"{project_id}_presupuesto_aprobado{Path(file_name).suffix or '.xlsx'}"
    return ProjectIdentity(
        project_id=project_id,
        project_name=project_name,
        folder_name=folder_name,
        gantt_file_name=gantt_file_name,
        budget_copy_name=budget_copy_name,
    )


def parse_date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = display_text(value, 40)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_int_value(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"\d+", display_text(value, 40).replace(",", ""))
    return int(match.group(0)) if match else None


def find_value_near_label(ws, label_terms: tuple[str, ...]) -> Any:
    for row in ws.iter_rows():
        for cell in row:
            label = normalize_text(cell.value)
            if not label or not all(term in label for term in label_terms):
                continue
            for offset in range(1, 5):
                value = ws.cell(cell.row, cell.column + offset).value
                if display_text(value):
                    return value
            value = ws.cell(cell.row + 1, cell.column).value
            if display_text(value):
                return value
    return None


def find_label_and_value(ws, label_terms: tuple[str, ...]) -> tuple[str, Any]:
    for row in ws.iter_rows():
        for cell in row:
            label = normalize_text(cell.value)
            if not label or not all(term in label for term in label_terms):
                continue
            for offset in range(1, 5):
                value = ws.cell(cell.row, cell.column + offset).value
                if display_text(value):
                    return label, value
            value = ws.cell(cell.row + 1, cell.column).value
            if display_text(value):
                return label, value
    return "", None


def parse_duration_value(label: str, value: Any) -> tuple[int | None, str]:
    text = normalize_text(f"{label} {display_text(value, 80)}")
    amount = parse_int_value(value)
    if amount is None:
        amount = parse_int_value(text)
    if amount is None or amount <= 0:
        return None, "days"
    if any(term in text for term in ("ano", "anos", "year", "years")):
        return amount, "years"
    if any(term in text for term in ("mes", "meses", "month", "months")):
        return amount, "months"
    if any(term in text for term in ("semana", "semanas", "week", "weeks")):
        return amount, "weeks"
    return amount, "days"


def add_calendar_duration(start: date, amount: int, unit: str) -> date:
    if unit == "days":
        return start + timedelta(days=amount - 1)
    if unit == "weeks":
        return start + timedelta(weeks=amount) - timedelta(days=1)
    months = amount * 12 if unit == "years" else amount
    month_index = start.month - 1 + months
    year = start.year + month_index // 12
    month = month_index % 12 + 1
    day = min(start.day, calendar.monthrange(year, month)[1])
    return date(year, month, day) - timedelta(days=1)


def subtract_calendar_duration(end: date, amount: int, unit: str) -> date:
    if unit == "days":
        return end - timedelta(days=amount - 1)
    if unit == "weeks":
        return end - timedelta(weeks=amount) + timedelta(days=1)
    months = amount * 12 if unit == "years" else amount
    month_index = end.month - 1 - months
    year = end.year + month_index // 12
    month = month_index % 12 + 1
    day = min(end.day, calendar.monthrange(year, month)[1])
    return date(year, month, day) + timedelta(days=1)


def read_project_window(wb) -> tuple[date, date, list[str]]:
    notes: list[str] = []
    start = None
    end = None
    duration_amount = None
    duration_unit = "days"

    if "Datos" in wb.sheetnames:
        ws = wb["Datos"]
        start = parse_date_value(find_value_near_label(ws, ("fecha", "inicio")))
        end = parse_date_value(find_value_near_label(ws, ("fecha", "final")))
        duration_label, duration_value = find_label_and_value(ws, ("duracion",))
        duration_amount, duration_unit = parse_duration_value(duration_label, duration_value)

    if start and end:
        if end < start:
            start, end = end, start
            notes.append("Fechas de Datos venían invertidas; se ordenaron.")
        notes.append(
            "Calendario delimitado por Fecha de Inicio y Fecha Final; "
            "Fecha Final tuvo prioridad sobre Duración."
        )
    elif start and duration_amount and not end:
        end = add_calendar_duration(start, duration_amount, duration_unit)
        notes.append(
            f"Fecha Final calculada desde Duración ({duration_amount} {duration_unit})."
        )
    elif end and duration_amount and not start:
        start = subtract_calendar_duration(end, duration_amount, duration_unit)
        notes.append(
            f"Fecha de Inicio calculada hacia atrás desde Fecha Final "
            f"({duration_amount} {duration_unit})."
        )
    elif start and end and end < start:
        start, end = end, start
        notes.append("Fechas de Datos venían invertidas; se ordenaron.")

    if not start:
        today = date.today()
        start = date(today.year, today.month, 1)
        notes.append("No se encontró fecha de inicio; se usó una ventana visual temporal.")
    if not end:
        end = start + timedelta(days=DEFAULT_WINDOW_DAYS - 1)
        notes.append(
            f"No se encontró fecha final; se usó una ventana visual de "
            f"{DEFAULT_WINDOW_DAYS} días."
        )
    return start, end, notes


def select_budget_sheet(wb) -> str:
    names = wb.sheetnames
    normalized = {normalize_text(name): name for name in names}
    flexio = [
        name
        for name in names
        if "presupuesto" in normalize_text(name)
        and "flexio" in normalize_text(name)
    ]
    if flexio:
        return flexio[0]
    if "presupuesto" in normalized:
        return normalized["presupuesto"]
    general = [name for name in names if "presupuesto" in normalize_text(name) and "general" in normalize_text(name)]
    if general:
        return general[0]
    pres = [name for name in names if "pres" in normalize_text(name)]
    if pres:
        return pres[0]
    for name in names:
        if normalize_text(name) not in {"datos", "destinatarios", "flujo de caja"}:
            return name
    return names[0]


def requires_general_budget_label(wb, selected_sheet: str) -> bool:
    selected = normalize_text(selected_sheet)
    if (
        "presupuesto" in selected
        and any(term in selected for term in ("flexio", "general"))
    ):
        return False
    reliable_sheets: list[str] = []
    for name in wb.sheetnames:
        if normalize_text(name) in {
            "datos",
            "destinatarios",
            "flujo de caja",
        }:
            continue
        try:
            candidate = detect_header_candidate(wb[name])
        except BudgetExtractionError:
            continue
        mapping = candidate.mapping
        support = sum(
            bool(mapping.get(field_name))
            for field_name in (
                "unidad",
                "cantidad",
                "costo_unitario",
                "costo_total",
            )
        )
        if mapping.get("actividad") and support >= 2:
            reliable_sheets.append(name)
    if "pres" not in selected and len(reliable_sheets) > 1:
        return True
    budget_sheets = [
        name
        for name in wb.sheetnames
        if "pres" in normalize_text(name)
        and normalize_text(name)
        not in {"datos", "destinatarios", "flujo de caja"}
    ]
    return len(budget_sheets) > 1


def is_non_empty(value: Any) -> bool:
    return display_text(value) != ""


def detect_header_row(ws) -> int:
    keyword_groups = (
        ("item", "partida", "no"),
        ("cc", "codigo", "c.c"),
        ("actividad", "descripcion", "descripción", "detalle", "concepto", "material"),
        ("cantidad", "cant", "qty"),
        ("unidad", "und", "um"),
        ("total", "subtotal", "monto", "precio", "costo"),
    )
    best_row = 1
    best_score = -1.0
    for row_idx in range(1, min(ws.max_row, 60) + 1):
        values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, ws.max_column + 1)]
        non_empty = [value for value in values if is_non_empty(value)]
        if len(non_empty) < 2:
            continue
        score = len(non_empty)
        for value in non_empty:
            text = normalize_text(value)
            if any(any(term in text for term in group) for group in keyword_groups):
                score += 4
            if isinstance(value, (int, float)):
                score -= 1
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def find_column(headers: list[str], terms: tuple[str, ...]) -> int | None:
    for idx, header in enumerate(headers, start=1):
        text = normalize_text(header)
        if any(term in text for term in terms):
            return idx
    return None


def find_item_column(headers: list[str]) -> int | None:
    for idx, header in enumerate(headers, start=1):
        text = normalize_text(header).replace(".", "")
        if text in {"item", "partida", "nro", "numero", "no"}:
            return idx
    return find_column(headers, ("item", "partida"))


def find_unit_column(headers: list[str]) -> int | None:
    for idx, header in enumerate(headers, start=1):
        text = normalize_text(header)
        if any(term in text for term in ("costo", "precio", "unitario", "c unidad", "p unitario", "c. unidad")):
            continue
        if any(term in text for term in ("unidad", "und", "um", "u/m")):
            return idx
    return None


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def is_numeric_unit(value: Any) -> bool:
    if is_number(value):
        return True
    text = display_text(value, 40).replace(",", ".")
    return bool(text and re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text))


def clean_unit_value(value: Any) -> Any:
    return None if is_numeric_unit(value) else value


def item_depth(value: Any) -> int:
    text = display_text(value, 40).strip(". ")
    if not text:
        return 0
    return len([part for part in re.split(r"[.\s]+", text) if part])


def classify_row(description: str, item: Any, quantity: Any, unit: Any, amount: Any) -> tuple[str, bool, str]:
    text = normalize_text(description)
    row_text = normalize_text(" ".join(display_text(value, 80) for value in (description, item, quantity, unit, amount)))
    has_quantity = is_non_empty(quantity)
    has_unit = is_non_empty(unit)
    has_amount = is_non_empty(amount)

    if is_month_unit(unit) and text:
        return "ACTIVIDAD_PROBABLE", False, "Unidad mes: se conserva para cronograma de obra."
    if is_resident_engineer_row(description, item, quantity, unit, amount):
        return "NO_CRONOGRAMA_PROBABLE", True, "Ingeniero residente obra es control/administracion, no actividad cronogramable."
    if not text and has_amount:
        return "AMBIGUO", True, "Tiene monto pero no descripcion clara."
    if re.match(r"^(costo|precio)\s*/", text):
        return "NO_CRONOGRAMA_PROBABLE", False, "Resumen de costo o precio unitario, no proceso cronogramable."
    if any(term in row_text for term in TOTAL_TERMS):
        return "NO_CRONOGRAMA_PROBABLE", True, "Subtotal, total, impuesto, margen, precio o resumen financiero."
    if any(term in row_text for term in INDIRECT_TERMS) and not any(term in row_text for term in ("instalacion", "montaje", "construccion")):
        return "NO_CRONOGRAMA_PROBABLE", True, "Indirecto, alquiler, equipo, campamento o administracion; validar antes de cronograma."
    if is_number(unit):
        return "AMBIGUO", True, "La columna Unidad contiene un numero; posible mapeo irregular del presupuesto."
    if item_depth(item) <= 2 and not has_quantity and not has_unit and text:
        return "AGRUPADOR", False, "Agrupador preservado para trazabilidad."
    if any(term in row_text for term in RESOURCE_TERMS) and not any(term in row_text for term in WORK_TERMS):
        return "NO_CRONOGRAMA_PROBABLE", True, "Parece insumo o recurso, no actividad ejecutable."
    if any(term in row_text for term in WORK_TERMS) or (has_quantity and has_unit and text):
        return "ACTIVIDAD_PROBABLE", False, "Actividad o paquete de trabajo probable."
    if text:
        return "AMBIGUO", True, "No hay evidencia suficiente para clasificar como actividad."
    return "AMBIGUO", True, "Fila sin descripcion util."


def value_at(ws, row_idx: int, col_idx: int | None) -> Any:
    if not col_idx:
        return None
    return ws.cell(row_idx, col_idx).value


def fallback_description(ws, row_idx: int, preferred_col: int | None, ignored_cols: set[int]) -> str:
    if preferred_col:
        text = display_text(ws.cell(row_idx, preferred_col).value, 220)
        if text:
            return text
    candidates: list[str] = []
    for col_idx in range(1, ws.max_column + 1):
        if col_idx in ignored_cols:
            continue
        value = ws.cell(row_idx, col_idx).value
        text = display_text(value, 220)
        if text and not is_number(value):
            candidates.append(text)
    return max(candidates, key=len) if candidates else ""


def resolve_budget_columns(headers: list[str]) -> dict[str, int | None]:
    return {
        "item": find_item_column(headers),
        "code": find_column(headers, ("cc", "c.c", "codigo", "centro")),
        "description": find_column(
            headers,
            ("actividad", "descripcion", "detalle", "concepto", "material", "partida"),
        ),
        "quantity": find_column(headers, ("cantidad", "cant", "qty", "volumen", "area")),
        "unit": find_unit_column(headers),
        "amount": find_column(headers, ("total", "subtotal", "monto", "importe", "valor")),
    }


def is_budget_header_row(headers: list[str], columns: dict[str, int | None]) -> bool:
    desc_col = columns["description"]
    unit_col = columns["unit"]
    qty_col = columns["quantity"]
    if not desc_col or not unit_col or not qty_col:
        return False
    desc_header = normalize_text(headers[desc_col - 1])
    unit_header = normalize_text(headers[unit_col - 1])
    qty_header = normalize_text(headers[qty_col - 1])
    return (
        any(term in desc_header for term in ("actividad", "descripcion", "detalle", "concepto", "partida"))
        and any(term in unit_header for term in ("unidad", "und", "u/m"))
        and any(term in qty_header for term in ("cantidad", "cant", "qty"))
    )


def extract_budget_rows(ws, header_row: int) -> list[BudgetRow]:
    headers = [display_text(ws.cell(header_row, col_idx).value, 80) for col_idx in range(1, ws.max_column + 1)]
    item_col = find_item_column(headers)
    code_col = find_column(headers, ("cc", "c.c", "codigo", "código", "centro"))
    desc_col = find_column(headers, ("actividad", "descripcion", "descripción", "detalle", "concepto", "material", "partida"))
    qty_col = find_column(headers, ("cantidad", "cant", "qty", "volumen", "area", "área"))
    unit_col = find_unit_column(headers)
    amount_col = find_column(headers, ("total", "subtotal", "monto", "importe", "valor"))
    ignored_cols = {col for col in (item_col, code_col, qty_col, unit_col, amount_col) if col}
    columns = resolve_budget_columns(headers)

    rows: list[BudgetRow] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, ws.max_column + 1)]
        if not any(is_non_empty(value) for value in values):
            continue
        row_headers = [display_text(value, 80) for value in values]
        row_columns = resolve_budget_columns(row_headers)
        if is_budget_header_row(row_headers, row_columns):
            columns = row_columns
            continue
        item_col = columns["item"]
        code_col = columns["code"]
        desc_col = columns["description"]
        qty_col = columns["quantity"]
        unit_col = columns["unit"]
        amount_col = columns["amount"]
        ignored_cols = {col for col in (item_col, code_col, qty_col, unit_col, amount_col) if col}
        description = fallback_description(ws, row_idx, desc_col, ignored_cols)
        if not description:
            continue
        item = value_at(ws, row_idx, item_col)
        code = value_at(ws, row_idx, code_col)
        quantity = value_at(ws, row_idx, qty_col)
        raw_unit = value_at(ws, row_idx, unit_col)
        unit = clean_unit_value(raw_unit)
        amount = value_at(ws, row_idx, amount_col)
        line_type, requires_review, reason = classify_row(description, item, quantity, unit, amount)
        if is_numeric_unit(raw_unit):
            requires_review = True
            reason = (
                "El valor numerico detectado en la columna Unidad se omitio; "
                "revisar el mapeo de columnas del presupuesto."
            )
        rows.append(
            BudgetRow(
                source_row=row_idx,
                line_type=line_type,
                item=item,
                code=code,
                description=description,
                quantity=quantity,
                unit=unit,
                amount=amount,
                requires_review=requires_review,
                review_reason=reason,
            )
        )
    return rows


def row_to_llm_candidate(row: BudgetRow) -> dict[str, Any]:
    return {
        "candidate_id": f"R{row.source_row}",
        "source_row": row.source_row,
        "local_kind": row.line_type,
        "item": display_text(row.item, 80),
        "cc": display_text(row.code, 80),
        "description": row.description,
        "quantity": display_text(row.quantity, 80),
        "unit": display_text(row.unit, 80),
        "subtotal": display_text(row.amount, 80),
        "local_reason": row.review_reason,
    }


def apply_llm_plan(
    rows: list[BudgetRow],
    *,
    workbook_name: str,
    sheet_name: str,
    header_row: int,
    options: LlmOptions,
    budget_structure: str = "activity_based",
) -> tuple[list[BudgetRow], list[str]]:
    if not options.enabled:
        return rows, ["LLM actividad/cronograma: desactivado; se uso clasificacion local."]
    if not rows:
        return rows, ["LLM actividad/cronograma: sin filas candidatas."]

    candidates = [row_to_llm_candidate(row) for row in rows]
    try:
        plan = request_llm_plan(
            workbook_name=workbook_name,
            sheet_name=sheet_name,
            header_row=header_row,
            candidates=candidates,
            model=options.model or None,
            budget_structure=budget_structure,
        )
    except Exception as exc:
        return rows, [f"LLM actividad/cronograma fallo; se uso clasificacion local. Error: {exc}"]

    source_by_id = {f"R{row.source_row}": row for row in rows}
    bucket_order = {"PRELIMINARES": 0, "FABRICA": 1, "CAMPO": 2, "ACABADOS": 3, "NO_CRONOGRAMA": 4}
    planned: list[tuple[int, int, int, BudgetRow]] = []
    missing = 0
    for candidate_id, source in source_by_id.items():
        decision = plan.decisions.get(candidate_id)
        if not decision:
            missing += 1
            planned.append(
                (
                    4,
                    999999,
                    source.source_row,
                    BudgetRow(
                        source_row=source.source_row,
                        line_type="NO_CRONOGRAMA",
                        item=source.item,
                        code=source.code,
                        description=source.description,
                        quantity=source.quantity,
                        unit=source.unit,
                        amount=source.amount,
                        requires_review=True,
                        review_reason="LLM no devolvio decision para esta fila; requiere revision.",
                    ),
                )
            )
            continue

        include = decision.include_in_cronograma and decision.planning_bucket != "NO_CRONOGRAMA"
        line_type = decision.planning_bucket if include else "NO_CRONOGRAMA"
        description = decision.actividad_normalizada or source.description
        review = decision.confidence != "ALTA" or not include
        reason = f"LLM {decision.confidence}: {decision.reason or 'Sin razon.'}"
        planned.append(
            (
                bucket_order.get(decision.planning_bucket, 4),
                decision.orden_cronologico,
                source.source_row,
                BudgetRow(
                    source_row=source.source_row,
                    line_type=line_type,
                    item=source.item,
                    code=source.code,
                    description=description,
                    quantity=source.quantity,
                    unit=source.unit,
                    amount=source.amount,
                    requires_review=review,
                    review_reason=reason,
                ),
            )
        )

    notes = [
        "LLM actividad/cronograma: activo.",
        f"LLM filas evaluadas={len(rows)}; decisiones={len(plan.decisions)}; faltantes={missing}.",
        f"LLM uso tokens={plan.usage.get('total_tokens', 0)} en {plan.usage.get('batches', 0)} lote(s).",
    ]
    notes.extend(f"LLM advertencia: {warning}" for warning in plan.warnings[:5])
    return [item[3] for item in sorted(planned, key=lambda item: (item[0], item[1], item[2]))], notes


def is_chronogram_row(row: BudgetRow) -> bool:
    if is_month_unit(row.unit):
        return True
    if is_resident_engineer_row(row.description, row.item, row.code):
        return False
    if row.line_type in CATEGORY_ORDER:
        return True
    if row.line_type in {"NO_CRONOGRAMA", "NO_CRONOGRAMA_PROBABLE"}:
        return False
    if any(term in normalize_text(row.line_type) for term in ("total", "resource", "admin", "indirect")):
        return False
    return row.line_type in {"ACTIVIDAD_PROBABLE", "AGRUPADOR"}


def category_for_row(row: BudgetRow) -> str:
    if is_month_unit(row.unit):
        return "CAMPO"
    if row.line_type in CATEGORY_ORDER:
        return row.line_type
    text = normalize_text(" ".join(display_text(value, 120) for value in (row.description, row.item, row.code)))
    if any(term in text for term in ("pintura", "pasteo", "acabado", "limpieza final", "terminacion")):
        return "ACABADOS"
    if any(term in text for term in ("fabricacion", "produccion", "planta", "formaleta", "molde", "columna", "panel", "pared")):
        return "FABRICA"
    if any(term in text for term in ("instalacion", "montaje", "campo", "obra", "fundacion", "losa", "vaciado", "grua")):
        return "CAMPO"
    return "PRELIMINARES"


def filter_cronogram_rows(rows: list[BudgetRow]) -> tuple[list[BudgetRow], list[str]]:
    kept = [row for row in rows if is_chronogram_row(row)]
    removed = [row for row in rows if not is_chronogram_row(row)]
    notes = [f"Filas fuera de cronograma omitidas del Gantt: {len(removed)}."]
    if removed:
        sample = ", ".join(str(row.source_row) for row in removed[:12])
        notes.append(f"Primeras filas omitidas: {sample}.")
    return kept, notes


def style_base_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 13
    ws.column_dimensions["F"].width = 13
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 42


def write_daily_headers(ws, start: date, end: date) -> int:
    dark_fill = PatternFill("solid", fgColor="1F3864")
    month_fill = PatternFill("solid", fgColor="244B6B")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    headers = [
        "Fuente_Fila",
        "Actividad",
        "Cantidad",
        "Unidad",
        "Fecha_Inicio",
        "Fecha_Fin",
        "Estado_Planificacion",
        "Comentarios",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(HEADER_ROW, col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF" if col_idx >= START_DATE_COL else "000000")
        cell.fill = dark_fill if col_idx >= START_DATE_COL else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    days = max(1, (end - start).days + 1)
    month_start_col = DAILY_START_COL
    current_month = start.month
    current_year = start.year
    for offset in range(days):
        current = start + timedelta(days=offset)
        col_idx = DAILY_START_COL + offset
        ws.column_dimensions[get_column_letter(col_idx)].width = 3.2
        day_cell = ws.cell(HEADER_ROW, col_idx)
        day_cell.value = current
        day_cell.number_format = "d"
        day_cell.font = Font(bold=True, size=8)
        day_cell.alignment = Alignment(horizontal="center", vertical="center")
        day_cell.border = Border(
            left=Side(style="thin", color="D9E2EA"),
            right=Side(style="thin", color="D9E2EA"),
            top=Side(style="thin", color="D9E2EA"),
            bottom=Side(style="thin", color="D9E2EA"),
        )
        weekday = ws.cell(WEEKDAY_ROW, col_idx)
        weekday.value = WEEKDAYS_ES[current.weekday()]
        weekday.font = Font(size=7, color="4B5563")
        weekday.alignment = Alignment(horizontal="center", vertical="center")

        month_changed = current.month != current_month or current.year != current_year
        is_last_day = offset == days - 1
        if month_changed or is_last_day:
            month_end_col = col_idx - 1 if month_changed else col_idx
            if month_end_col >= month_start_col:
                if month_end_col > month_start_col:
                    ws.merge_cells(
                        start_row=9,
                        start_column=month_start_col,
                        end_row=9,
                        end_column=month_end_col,
                    )
                month_cell = ws.cell(9, month_start_col)
                month_cell.value = f"{MONTHS_ES[current_month]} {current_year}"
                month_cell.fill = month_fill
                month_cell.font = Font(bold=True, color="FFFFFF", size=9)
                month_cell.alignment = Alignment(horizontal="center", vertical="center")
            if month_changed:
                month_start_col = col_idx
                current_month = current.month
                current_year = current.year

    return DAILY_START_COL + days - 1


def add_validations(ws, activity_rows: list[int]) -> None:
    if not activity_rows:
        return
    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2035,12,31)",
        allow_blank=True,
    )
    date_validation.error = "Ingrese una fecha valida."
    date_validation.errorTitle = "Fecha invalida"
    ws.add_data_validation(date_validation)
    for row_idx in activity_rows:
        date_validation.add(f"E{row_idx}:F{row_idx}")

    status_validation = DataValidation(
        type="list",
        formula1='"Pendiente,Validado,Requiere Ajuste,Excluido"',
        allow_blank=True,
    )
    ws.add_data_validation(status_validation)
    for row_idx in activity_rows:
        status_validation.add(f"G{row_idx}")


def add_bar_formatting(ws, activity_rows: list[int], calendar_end_col: int) -> None:
    if not activity_rows:
        return
    first_col = get_column_letter(DAILY_START_COL)
    last_col = get_column_letter(calendar_end_col)
    bar_fill = PatternFill(fill_type="solid", start_color="2F80ED", end_color="2F80ED")
    invalid_fill = PatternFill(fill_type="solid", start_color="F8D7DA", end_color="F8D7DA")
    thin = Side(style="thin", color="E3E8EF")

    for row_idx in activity_rows:
        for col_idx in range(DAILY_START_COL, calendar_end_col + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row_range = f"{first_col}{row_idx}:{last_col}{row_idx}"
        formula = (
            f'=AND($E{row_idx}<>"",$F{row_idx}<>"",'
            f'{first_col}$10>=$E{row_idx},'
            f'{first_col}$10<=$F{row_idx})'
        )
        ws.conditional_formatting.add(row_range, FormulaRule(formula=[formula], fill=bar_fill))

        ws.conditional_formatting.add(
            f"E{row_idx}:F{row_idx}",
            FormulaRule(
                formula=[f'=AND($E{row_idx}<>"",$F{row_idx}<>"",$F{row_idx}<$E{row_idx})'],
                fill=invalid_fill,
            ),
        )


def write_budget_rows(ws, rows: list[BudgetRow], calendar_end_col: int) -> tuple[int, list[int]]:
    border = Border(
        left=Side(style="thin", color="B7C9D8"),
        right=Side(style="thin", color="B7C9D8"),
        top=Side(style="thin", color="B7C9D8"),
        bottom=Side(style="thin", color="B7C9D8"),
    )
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    data_end = DATA_START_ROW - 1
    activity_rows: list[int] = []

    def write_section(category: str, section_row: int) -> None:
        for col_idx in range(1, calendar_end_col + 1):
            cell = ws.cell(section_row, col_idx)
            cell.fill = section_fill
            cell.font = Font(bold=True, color="1F3864")
            cell.border = border
        ws.merge_cells(
            start_row=section_row,
            start_column=1,
            end_row=section_row,
            end_column=calendar_end_col,
        )
        cell = ws.cell(section_row, 1)
        cell.value = category
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        ws.row_dimensions[section_row].height = 22

    grouped: dict[str, list[BudgetRow]] = {category: [] for category in CATEGORY_ORDER}
    for row in rows:
        grouped.setdefault(category_for_row(row), []).append(row)

    row_idx = DATA_START_ROW
    for category in CATEGORY_ORDER:
        category_rows = grouped.get(category, [])
        if not category_rows:
            continue

        write_section(category, row_idx)
        data_end = row_idx
        row_idx += 1

        for row in category_rows:
            activity_rows.append(row_idx)
            values = [
                row.source_row,
                row.description,
                row.quantity,
                row.unit,
                None,
                None,
                "Pendiente",
                row.review_reason if row.requires_review else "",
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row_idx, col_idx)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=col_idx in {2, 8})
                if row.requires_review:
                    cell.fill = review_fill
            ws.cell(row_idx, 5).number_format = "dd/mm/yyyy"
            ws.cell(row_idx, 6).number_format = "dd/mm/yyyy"
            data_end = row_idx
            row_idx += 1

    extra_categories = sorted(category for category in grouped if category not in CATEGORY_ORDER and grouped[category])
    for category in extra_categories:
        write_section(category, row_idx)
        data_end = row_idx
        row_idx += 1
        for row in grouped[category]:
            activity_rows.append(row_idx)
            values = [row.source_row, row.description, row.quantity, row.unit, None, None, "Pendiente", row.review_reason]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row_idx, col_idx)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=col_idx in {2, 8})
            ws.cell(row_idx, 5).number_format = "dd/mm/yyyy"
            ws.cell(row_idx, 6).number_format = "dd/mm/yyyy"
            data_end = row_idx
            row_idx += 1

    if not rows:
        data_end = DATA_START_ROW
        ws.cell(DATA_START_ROW, 2).value = "No se detectaron actividades cronogramables en el presupuesto. Requiere revision manual."
        ws.cell(DATA_START_ROW, 8).value = "El archivo no contiene un bloque presupuestario claro o el LLM excluyo todo como no cronogramable."
    return data_end, activity_rows


def write_metadata_sheet(wb, identity: ProjectIdentity, source_file: str, selected_sheet: str, notes: list[str]) -> None:
    ws = wb.create_sheet("Datos")
    rows = [
        ("ProyectoID", identity.project_id),
        ("NombreProyecto", identity.project_name),
        ("ArchivoPresupuesto", source_file),
        ("HojaPresupuestoUsada", selected_sheet),
        ("Notas", " | ".join(notes)),
        ("EstadoGantt", PROJECT_STATUS_DEFAULT),
    ]
    for row_idx, (label, value) in enumerate(rows, start=1):
        ws.cell(row_idx, 1).value = label
        ws.cell(row_idx, 1).font = Font(bold=True)
        ws.cell(row_idx, 2).value = value
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    status_validation = DataValidation(
        type="list",
        formula1='"En progreso,En revisión inicial"',
        allow_blank=False,
    )
    ws.add_data_validation(status_validation)
    status_validation.formula1 = f'"{",".join(PROJECT_STATUS_VALUES)}"'
    status_validation.add("B6")
    ws.freeze_panes = None


GANTT_FIXED_HEADERS = [
    "Actividad",
    "CC",
    "Fecha de Inicio",
    "Fecha de Fin",
    "Estatus",
    "Unidad",
    "Cantidad",
    "Costo Unitario",
    "Costo Total",
]


def additional_cost_columns(
    extraction: BudgetExtractionResult,
) -> list[BudgetCostColumn]:
    # Las columnas financieras adicionales permanecen disponibles en la
    # extracción reutilizable y en la hoja fuente, pero no pertenecen al
    # cronograma. El Gantt solo expone Costo Unitario y Costo Total.
    return []


def gantt_cost_groups(extraction: BudgetExtractionResult) -> dict[str, str]:
    # Rótulos situados encima del encabezado real (por ejemplo "Área total")
    # son auxiliares del presupuesto y no deben aparecer sobre el Gantt.
    return {}


def gantt_headers(
    has_cc: bool,
    has_item: bool = False,
    extra_cost_columns: list[BudgetCostColumn] | None = None,
) -> list[str]:
    headers = [
        header
        for header in GANTT_FIXED_HEADERS
        if has_cc or header != "CC"
    ]
    if has_item:
        headers.insert(0, "Ítem")
    return headers


def _reset_controlled_gantt_sheet(ws) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    if ws.max_column:
        ws.delete_cols(1, ws.max_column)
    ws.data_validations.dataValidation = []
    ws.conditional_formatting._cf_rules.clear()
    ws.auto_filter.ref = None
    ws.freeze_panes = None
    ws.row_dimensions.clear()
    ws.column_dimensions.clear()


def _prepare_gantt_sheet(wb):
    if "Gantt" in wb.sheetnames:
        ws = wb["Gantt"]
        _reset_controlled_gantt_sheet(ws)
        return ws
    if "Gantt_Diario" in wb.sheetnames:
        ws = wb["Gantt_Diario"]
        ws.title = "Gantt"
        _reset_controlled_gantt_sheet(ws)
        return ws
    return wb.create_sheet("Gantt", 0)


def _calendar_headers(
    ws,
    start: date,
    contractual_end: date,
    display_end: date,
    start_column: int,
) -> int:
    days = (display_end - start).days + 1
    if days < 1:
        raise GanttReviewRequiredError("La ventana del proyecto tiene fechas inválidas.")
    if start_column + days - 1 > 16384:
        raise GanttReviewRequiredError(
            "La ventana del proyecto excede el máximo de columnas de Excel."
        )

    month_fill = PatternFill("solid", fgColor="244B6B")
    extension_month_fill = PatternFill("solid", fgColor="C00000")
    day_fill = PatternFill("solid", fgColor="D9EAF7")
    extension_day_fill = PatternFill("solid", fgColor="F4CCCC")
    thin = Side(style="thin", color="D9E2EA")
    month_start = start_column
    active_month = (start.year, start.month)
    for offset in range(days):
        current = start + timedelta(days=offset)
        column = start_column + offset
        ws.column_dimensions[get_column_letter(column)].width = 3.2
        day_cell = ws.cell(HEADER_ROW, column, current)
        day_cell.number_format = "d"
        day_cell.font = Font(bold=True, size=8, color="1F2937")
        in_extension = current > contractual_end
        day_cell.fill = extension_day_fill if in_extension else day_fill
        day_cell.alignment = Alignment(horizontal="center", vertical="center")
        day_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        weekday_cell = ws.cell(WEEKDAY_ROW, column, WEEKDAYS_ES[current.weekday()])
        weekday_cell.font = Font(size=7, color="4B5563")
        weekday_cell.fill = extension_day_fill if in_extension else day_fill
        weekday_cell.alignment = Alignment(horizontal="center", vertical="center")
        weekday_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        next_month = (current.year, current.month)
        next_day = current + timedelta(days=1)
        closes_month = offset == days - 1 or (next_day.year, next_day.month) != next_month
        if closes_month:
            if column > month_start:
                ws.merge_cells(
                    start_row=HEADER_ROW - 1,
                    start_column=month_start,
                    end_row=HEADER_ROW - 1,
                    end_column=column,
                )
            month_cell = ws.cell(HEADER_ROW - 1, month_start)
            month_cell.value = f"{MONTHS_ES[active_month[1]]} {active_month[0]}"
            month_cell.fill = (
                extension_month_fill
                if date(active_month[0], active_month[1], 1) > contractual_end
                else month_fill
            )
            month_cell.font = Font(bold=True, color="FFFFFF", size=9)
            month_cell.alignment = Alignment(horizontal="center", vertical="center")
            month_start = column + 1
            active_month = (next_day.year, next_day.month)
    return start_column + days - 1


def _base_column_indexes(headers: list[str]) -> dict[str, int]:
    return {header: index for index, header in enumerate(headers, start=1)}


def _link_source_formulas(
    extraction: BudgetExtractionResult,
    source_ws,
) -> None:
    escaped_sheet = source_ws.title.replace("'", "''")
    linked = 0
    for row in extraction.rows:
        if row.row_type != "activity":
            continue
        for field_name in ("cantidad", "costo_unitario", "costo_total"):
            if getattr(row, field_name) is not None:
                continue
            coordinate = row.source_cells.get(field_name)
            if not coordinate:
                continue
            source_value = source_ws[coordinate].value
            if not (isinstance(source_value, str) and source_value.startswith("=")):
                continue
            setattr(row, field_name, f"='{escaped_sheet}'!{coordinate}")
            row.warnings.append(
                f"{field_name} no tenía valor cacheado; conserva la fórmula como último "
                f"recurso mediante vínculo a {source_ws.title}!{coordinate}."
            )
            linked += 1
    if linked:
        extraction.assessment.warnings.append(
            f"Se usaron {linked} vínculos de fórmula solo donde el presupuesto no tenía "
            "un valor numérico cacheado."
        )


def _source_cell_for_field(source_ws, source_row, field_name: str | None):
    if not field_name:
        return None
    coordinate = source_row.source_cells.get(field_name)
    return source_ws[coordinate] if coordinate else None


def _copy_source_appearance(target, source, *, copy_number_format: bool = False) -> None:
    if source is None:
        return
    target.fill = copy(source.fill)
    target.font = copy(source.font)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    if copy_number_format:
        target.number_format = source.number_format


def _write_gantt_rows(
    ws,
    source_ws,
    extraction: BudgetExtractionResult,
    headers: list[str],
    calendar_start_column: int,
    contractual_end_column: int,
    calendar_end_column: int,
) -> tuple[list[int], int]:
    indexes = _base_column_indexes(headers)
    extra_costs = additional_cost_columns(extraction)
    extra_cost_by_header = {
        column.output_header: column
        for column in extra_costs
    }
    border = Border(
        left=Side(style="thin", color="B7C9D8"),
        right=Side(style="thin", color="B7C9D8"),
        top=Side(style="thin", color="B7C9D8"),
        bottom=Side(style="thin", color="B7C9D8"),
    )
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    activity_rows: list[int] = []
    output_row = DATA_START_ROW
    source_default_height = source_ws.sheet_format.defaultRowHeight or 15

    for source_row in extraction.rows:
        source_height = (
            source_ws.row_dimensions[source_row.row_number].height
            or source_default_height
        )
        if source_row.row_type == "spacer":
            ws.row_dimensions[output_row].height = source_height
            output_row += 1
            continue
        if source_row.row_type == "section":
            source_activity_cell = _source_cell_for_field(
                source_ws,
                source_row,
                "actividad",
            )
            for column in range(1, calendar_end_column + 1):
                cell = ws.cell(output_row, column)
                cell.fill = section_fill
                cell.font = Font(bold=True, color="1F3864")
                cell.border = border
                if source_activity_cell is not None:
                    cell.fill = copy(source_activity_cell.fill)
                    cell.font = copy(source_activity_cell.font)
            ws.merge_cells(
                start_row=output_row,
                start_column=1,
                end_row=output_row,
                end_column=calendar_end_column,
            )
            label = source_row.actividad
            if display_text(source_row.cc):
                label = f"{display_text(source_row.cc)} - {label}"
            cell = ws.cell(output_row, 1, label)
            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
                wrap_text=True,
            )
            ws.row_dimensions[output_row].height = source_height
            output_row += 1
            continue

        activity_rows.append(output_row)
        values = {
            "Ítem": source_row.item,
            "Actividad": source_row.actividad,
            "CC": source_row.cc,
            "Fecha de Inicio": None,
            "Fecha de Fin": None,
            "Estatus": None,
            "Unidad": source_row.unidad,
            "Cantidad": source_row.cantidad,
            "Costo Unitario": source_row.costo_unitario,
            "Costo Total": source_row.costo_total,
        }
        values.update(
            {
                column.output_header: source_row.extra_costs.get(column.key)
                for column in extra_costs
            }
        )
        if (
            values["Costo Total"] is None
            and values["Cantidad"] is not None
            and values["Costo Unitario"] is not None
        ):
            quantity_column = get_column_letter(indexes["Cantidad"])
            unit_cost_column = get_column_letter(indexes["Costo Unitario"])
            values["Costo Total"] = (
                f"={quantity_column}{output_row}*{unit_cost_column}{output_row}"
            )
            if source_row.row_number not in extraction.assessment.calculated_cost_rows:
                extraction.assessment.calculated_cost_rows.append(source_row.row_number)
            source_row.warnings.append(
                "Costo Total calculado en el Gantt como Cantidad × Costo Unitario."
            )
        for header, column in indexes.items():
            cell = ws.cell(output_row, column, values.get(header))
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=header == "Actividad",
                indent=source_row.level if header == "Actividad" else 0,
            )
            field_name = {
                "Ítem": "item",
                "Actividad": "actividad",
                "CC": "cc",
                "Unidad": "unidad",
                "Cantidad": "cantidad",
                "Costo Unitario": "costo_unitario",
                "Costo Total": "costo_total",
            }.get(header)
            if header in extra_cost_by_header:
                field_name = extra_cost_by_header[header].key
            source_cell = (
                _source_cell_for_field(source_ws, source_row, field_name)
                if field_name
                else _source_cell_for_field(source_ws, source_row, "actividad")
            )
            _copy_source_appearance(
                cell,
                source_cell,
                copy_number_format=header
                in {"Cantidad", "Costo Unitario", "Costo Total"}
                or header in extra_cost_by_header,
            )
        for header in ("Fecha de Inicio", "Fecha de Fin"):
            date_cell = ws.cell(output_row, indexes[header])
            date_cell.number_format = "dd/mm/yyyy"
            date_cell.fill = PatternFill("solid", fgColor="D9EAF7")
        quantity_cell = ws.cell(output_row, indexes["Cantidad"])
        if quantity_cell.number_format == "General":
            quantity_cell.number_format = "#,##0.00"
        for header in ("Costo Unitario", "Costo Total"):
            cost_cell = ws.cell(output_row, indexes[header])
            if cost_cell.number_format == "General":
                cost_cell.number_format = (
                    '[$B/.-180A] #,##0.00;[Red]-[$B/.-180A] #,##0.00'
                )
        for header in extra_cost_by_header:
            cost_cell = ws.cell(output_row, indexes[header])
            if cost_cell.number_format == "General":
                role = extra_cost_by_header[header].role
                cost_cell.number_format = (
                    "0.00%"
                    if role == "margen"
                    else '[$B/.-180A] #,##0.00;[Red]-[$B/.-180A] #,##0.00'
                )

        calendar_border = Side(style="thin", color="E3E8EF")
        source_activity_cell = _source_cell_for_field(
            source_ws,
            source_row,
            "actividad",
        )
        for column in range(calendar_start_column, calendar_end_column + 1):
            calendar_cell = ws.cell(output_row, column)
            if column > contractual_end_column:
                calendar_cell.fill = PatternFill("solid", fgColor="FCE8E6")
                calendar_cell.font = Font(color="9C0006")
            elif source_activity_cell is not None:
                calendar_cell.fill = copy(source_activity_cell.fill)
                calendar_cell.font = copy(source_activity_cell.font)
            calendar_cell.border = Border(
                left=calendar_border,
                right=calendar_border,
                top=calendar_border,
                bottom=calendar_border,
            )
        ws.row_dimensions[output_row].height = source_height
        output_row += 1
    return activity_rows, output_row - 1


def _add_gantt_validations_and_bars(
    ws,
    activity_rows: list[int],
    headers: list[str],
    calendar_start_column: int,
    calendar_end_column: int,
) -> None:
    if not activity_rows:
        return
    indexes = _base_column_indexes(headers)
    start_column = get_column_letter(indexes["Fecha de Inicio"])
    end_column = get_column_letter(indexes["Fecha de Fin"])
    status_column = get_column_letter(indexes["Estatus"])
    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2000,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
    )
    date_validation.error = "Ingrese una fecha válida."
    date_validation.errorTitle = "Fecha inválida"
    ws.add_data_validation(date_validation)
    status_validation = DataValidation(
        type="list",
        formula1='"Pendiente,En progreso,En revisión inicial"',
        allow_blank=True,
    )
    ws.add_data_validation(status_validation)

    first_calendar = get_column_letter(calendar_start_column)
    last_calendar = get_column_letter(calendar_end_column)
    bar_fill = PatternFill(fill_type="solid", start_color="2F80ED", end_color="2F80ED")
    invalid_fill = PatternFill(fill_type="solid", start_color="F8D7DA", end_color="F8D7DA")
    for row_number in activity_rows:
        date_validation.add(f"{start_column}{row_number}:{end_column}{row_number}")
        status_validation.add(f"{status_column}{row_number}")
        calendar_range = f"{first_calendar}{row_number}:{last_calendar}{row_number}"
        ws.conditional_formatting.add(
            calendar_range,
            FormulaRule(
                formula=[
                    f'=AND(${start_column}{row_number}<>"",'
                    f'${end_column}{row_number}<>"",'
                    f'{first_calendar}${HEADER_ROW}>=${start_column}{row_number},'
                    f'{first_calendar}${HEADER_ROW}<=${end_column}{row_number})'
                ],
                fill=bar_fill,
            ),
        )
        ws.conditional_formatting.add(
            f"{start_column}{row_number}:{end_column}{row_number}",
            FormulaRule(
                formula=[
                    f'=AND(${start_column}{row_number}<>"",'
                    f'${end_column}{row_number}<>"",'
                    f'${end_column}{row_number}<${start_column}{row_number})'
                ],
                fill=invalid_fill,
            ),
        )


def _style_gantt_sheet(
    ws,
    identity: ProjectIdentity,
    source_file_name: str,
    headers: list[str],
    cost_groups: dict[str, str],
    calendar_end_column: int,
    data_end_row: int,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    widths = {
        "Actividad": 55,
        "CC": 16,
        "Fecha de Inicio": 15,
        "Fecha de Fin": 15,
        "Estatus": 20,
        "Unidad": 12,
        "Cantidad": 14,
        "Costo Unitario": 17,
        "Costo Total": 17,
        "Ítem": 14,
    }
    header_fill = PatternFill("solid", fgColor="1F3864")
    border = Border(
        left=Side(style="thin", color="B7C9D8"),
        right=Side(style="thin", color="B7C9D8"),
        top=Side(style="thin", color="B7C9D8"),
        bottom=Side(style="thin", color="B7C9D8"),
    )
    for column, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(column)].width = widths.get(header, 18)
        cell = ws.cell(HEADER_ROW, column, header)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        subheader = ws.cell(WEEKDAY_ROW, column)
        subheader.fill = PatternFill("solid", fgColor="EAF2F8")
        subheader.border = border

    grouped_columns = [
        (column, cost_groups.get(header, ""))
        for column, header in enumerate(headers, start=1)
        if cost_groups.get(header, "")
    ]
    group_fill = PatternFill("solid", fgColor="244B6B")
    index = 0
    while index < len(grouped_columns):
        start_column, label = grouped_columns[index]
        end_column = start_column
        index += 1
        while (
            index < len(grouped_columns)
            and grouped_columns[index][1] == label
            and grouped_columns[index][0] == end_column + 1
        ):
            end_column = grouped_columns[index][0]
            index += 1
        if end_column > start_column:
            ws.merge_cells(
                start_row=HEADER_ROW - 1,
                start_column=start_column,
                end_row=HEADER_ROW - 1,
                end_column=end_column,
            )
        group_cell = ws.cell(HEADER_ROW - 1, start_column, label)
        group_cell.fill = group_fill
        group_cell.font = Font(bold=True, color="FFFFFF", size=9)
        group_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    title_rows = (
        (1, "AUTOSYS - GANTT WORKING", 15, True),
        (2, f"Proyecto: {identity.project_id} - {identity.project_name}", 11, True),
        (3, f"Presupuesto fuente: {source_file_name}", 10, False),
        (
            4,
            "Las fechas por actividad están vacías. El calendario se pinta al completar "
            "Fecha de Inicio y Fecha de Fin.",
            10,
            False,
        ),
    )
    for row_number, text, size, bold in title_rows:
        if calendar_end_column > 1:
            ws.merge_cells(
                start_row=row_number,
                start_column=1,
                end_row=row_number,
                end_column=calendar_end_column,
            )
        cell = ws.cell(row_number, 1, text)
        cell.font = Font(bold=bold, size=size, color="1F3864")
        cell.alignment = Alignment(horizontal="left", vertical="center")

    status_label = ws[PROJECT_STATUS_LABEL_CELL]
    status_label.value = PROJECT_STATUS_LABEL
    status_label.fill = PatternFill("solid", fgColor="1F3864")
    status_label.font = Font(bold=True, color="FFFFFF")
    status_label.alignment = Alignment(horizontal="left", vertical="center")
    status_label.border = border

    status_value = ws[PROJECT_STATUS_VALUE_CELL]
    status_value.value = PROJECT_STATUS_DEFAULT
    status_value.fill = PatternFill("solid", fgColor="D9EAF7")
    status_value.font = Font(bold=True, color="1F3864")
    status_value.alignment = Alignment(horizontal="center", vertical="center")
    status_value.border = border

    project_status_validation = DataValidation(
        type="list",
        formula1='"En progreso,En revisión inicial"',
        allow_blank=False,
    )
    project_status_validation.error = (
        "Seleccione En progreso o En revisión inicial."
    )
    project_status_validation.errorTitle = "Estado inválido"
    project_status_validation.formula1 = f'"{",".join(PROJECT_STATUS_VALUES)}"'
    project_status_validation.error = "Seleccione Actual, En Progreso o Entregar."
    ws.add_data_validation(project_status_validation)
    project_status_validation.add(PROJECT_STATUS_VALUE_CELL)
    ws.row_dimensions[6].height = 24
    ws.row_dimensions[HEADER_ROW].height = 34
    ws.row_dimensions[WEEKDAY_ROW].height = 18
    for row_number in range(DATA_START_ROW, data_end_row + 1):
        if ws.row_dimensions[row_number].height is None:
            ws.row_dimensions[row_number].height = 21


def _metadata_sheet(wb, identity: ProjectIdentity, source_file: str, source_sheet: str, notes: list[str]) -> None:
    if "Datos" in wb.sheetnames:
        ws = wb["Datos"]
    else:
        ws = wb.create_sheet("Datos")
    existing: dict[str, int] = {}
    status_found = False
    for row_number in range(1, ws.max_row + 1):
        for column_number in range(1, ws.max_column + 1):
            label = normalize_text(
                ws.cell(row_number, column_number).value
            )
            if column_number == 1 and label:
                existing[label] = row_number
            if label in {
                "estadogantt",
                "estado gantt",
                "statusgantt",
                "status gantt",
                "status",
            }:
                ws.cell(row_number, column_number + 1).value = PROJECT_STATUS_DEFAULT
                status_found = True
    additions = [
        ("ProyectoID", identity.project_id),
        ("NombreProyecto", identity.project_name),
        ("ArchivoPresupuesto", source_file),
        ("HojaPresupuestoUsada", source_sheet),
        ("NotasGeneradorGantt", " | ".join(notes)),
    ]
    if not status_found:
        additions.append(("EstadoGantt", PROJECT_STATUS_DEFAULT))
    row_number = max(1, ws.max_row + 1)
    for label, value in additions:
        if normalize_text(label) in existing:
            continue
        ws.cell(row_number, 1, label).font = Font(bold=True)
        ws.cell(row_number, 2, value)
        row_number += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 28)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 90)
    ws.freeze_panes = None


def _version_baseline_sheet(
    wb,
    gantt_ws,
    headers: list[str],
    activity_rows: list[int],
    contractual_end: date,
) -> None:
    if VERSION_BASELINE_SHEET in wb.sheetnames:
        wb.remove(wb[VERSION_BASELINE_SHEET])
    ws = wb.create_sheet(VERSION_BASELINE_SHEET)
    ws.append(["AUTOSYS_VERSION_BASELINE", "Valor"])
    ws.append(["Fecha Final contractual", contractual_end])
    indexes = _base_column_indexes(headers)
    for header, column in indexes.items():
        normalized_header = normalize_text(header)
        if not any(
            term in normalized_header
            for term in ("costo total", "precio total")
        ):
            continue
        if any(term in normalized_header for term in ("utilidad", "margen")):
            continue
        total = 0.0
        for row_number in activity_rows:
            value = gantt_ws.cell(row_number, column).value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                total += float(value)
                continue
            if (
                isinstance(value, str)
                and value.startswith("=")
                and "Cantidad" in indexes
                and "Costo Unitario" in indexes
            ):
                quantity = gantt_ws.cell(row_number, indexes["Cantidad"]).value
                unit_cost = gantt_ws.cell(
                    row_number,
                    indexes["Costo Unitario"],
                ).value
                if isinstance(quantity, (int, float)) and isinstance(
                    unit_cost,
                    (int, float),
                ):
                    total += float(quantity) * float(unit_cost)
        ws.append([header, total])
    ws.sheet_state = "hidden"


def _assessment_sheet(
    wb,
    extraction: BudgetExtractionResult,
    validation: dict[str, Any] | None = None,
) -> None:
    if "Assessment" in wb.sheetnames:
        index = wb.sheetnames.index("Assessment")
        wb.remove(wb["Assessment"])
        ws = wb.create_sheet("Assessment", index)
    else:
        ws = wb.create_sheet("Assessment")
    assessment = extraction.assessment
    summary = [
        ("Archivo fuente", extraction.source_file),
        ("Hoja usada", extraction.source_sheet),
        ("Fila de encabezado", extraction.header_row),
        ("Filas preservadas", len(extraction.rows)),
        (
            "Filas de actividad",
            sum(1 for row in extraction.rows if row.row_type == "activity"),
        ),
        (
            "Filas agrupadoras",
            sum(1 for row in extraction.rows if row.row_type == "section"),
        ),
        (
            "Filas separadoras",
            sum(1 for row in extraction.rows if row.row_type == "spacer"),
        ),
        ("Columnas financieras preservadas", len(extraction.cost_columns)),
        ("Columnas faltantes", ", ".join(assessment.missing_columns) or "Ninguna"),
        ("LLM usado para mapping", "Sí" if assessment.llm_mapping_used else "No"),
        (
            "Validación final",
            json.dumps(validation or {"status": "pending"}, ensure_ascii=False, default=str),
        ),
    ]
    ws.append(["Assessment del generador de Gantt", "Valor"])
    for label, value in summary:
        ws.append([label, value])

    ws.append([])
    ws.append(["Campo", "Columna", "Encabezado fuente", "Alias usado", "Score"])
    for field_name, details in assessment.detected_columns.items():
        ws.append(
            [
                field_name,
                details.get("column"),
                details.get("header"),
                assessment.aliases_used.get(field_name, ""),
                details.get("score"),
            ]
        )

    ws.append([])
    ws.append(
        [
            "Columna financiera",
            "Índice fuente",
            "Grupo",
            "Rol",
            "Encabezado Gantt",
        ]
    )
    for column in extraction.cost_columns:
        ws.append(
            [
                column.header,
                column.column,
                column.group,
                column.role,
                column.output_header,
            ]
        )

    ws.append([])
    ws.append(["Fila ignorada", "Razón", "Actividad"])
    for item in assessment.ignored_rows:
        ws.append([item.get("row_number"), item.get("reason"), item.get("activity", "")])

    ws.append([])
    ws.append(["Fila ambigua", "Actividad", "Razón"])
    for item in assessment.ambiguous_rows:
        ws.append([item.get("row_number"), item.get("activity"), item.get("reason")])

    ws.append([])
    ws.append(["Costos calculados", "Filas fuente"])
    ws.append(
        [
            len(assessment.calculated_cost_rows),
            ", ".join(str(value) for value in assessment.calculated_cost_rows),
        ]
    )
    ws.append([])
    ws.append(["Advertencias"])
    for warning in assessment.warnings:
        ws.append([warning])
    for extracted_row in extraction.rows:
        for warning in extracted_row.warnings:
            ws.append([f"Fila {extracted_row.row_number}: {warning}"])

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 14
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    ws.freeze_panes = None


def validate_generated_gantt(
    output_path: Path,
    extraction: BudgetExtractionResult,
    has_cc: bool,
) -> dict[str, Any]:
    if not output_path.exists() or output_path.stat().st_size <= 0:
        raise GanttReviewRequiredError("El archivo Gantt no existe o quedó vacío.")
    wb = load_workbook(output_path, data_only=False, read_only=False)
    try:
        if "Gantt" not in wb.sheetnames:
            raise GanttReviewRequiredError("El archivo generado no contiene la hoja Gantt.")
        ws = wb["Gantt"]
        expected_headers = gantt_headers(
            has_cc,
            bool(extraction.column_mapping.get("item")),
            additional_cost_columns(extraction),
        )
        actual_headers = [
            display_text(ws.cell(HEADER_ROW, column).value, 80)
            for column in range(1, len(expected_headers) + 1)
        ]
        if actual_headers != expected_headers:
            raise GanttReviewRequiredError(
                f"Encabezados Gantt inválidos. Esperados={expected_headers}; "
                f"detectados={actual_headers}."
            )
        merged_rows = {
            merged_range.min_row
            for merged_range in ws.merged_cells.ranges
            if merged_range.min_row >= DATA_START_ROW
        }
        output_activity_rows = [
            row_number
            for row_number in range(DATA_START_ROW, ws.max_row + 1)
            if row_number not in merged_rows and display_text(ws.cell(row_number, 1).value)
        ]
        expected_activities = sum(
            1 for row in extraction.rows if row.row_type == "activity"
        )
        if len(output_activity_rows) != expected_activities:
            raise GanttReviewRequiredError(
                f"Cantidad de actividades inconsistente: esperadas={expected_activities}, "
                f"generadas={len(output_activity_rows)}."
            )
        indexes = _base_column_indexes(expected_headers)
        date_cells_blank = all(
            ws.cell(row_number, indexes["Fecha de Inicio"]).value in (None, "")
            and ws.cell(row_number, indexes["Fecha de Fin"]).value in (None, "")
            for row_number in output_activity_rows
        )
        if not date_cells_blank:
            raise GanttReviewRequiredError(
                "El generador introdujo fechas por actividad; deben quedar vacías."
            )
        source_has_costs = any(
            row.costo_unitario is not None or row.costo_total is not None
            for row in extraction.rows
            if row.row_type == "activity"
        )
        output_has_costs = any(
            ws.cell(row_number, indexes["Costo Unitario"]).value is not None
            or ws.cell(row_number, indexes["Costo Total"]).value is not None
            for row_number in output_activity_rows
        )
        if source_has_costs and not output_has_costs:
            raise GanttReviewRequiredError(
                "El presupuesto tenía costos, pero el Gantt quedó sin costos."
            )
        buffer = io.BytesIO()
        wb.save(buffer)
        if buffer.tell() <= 0:
            raise GanttReviewRequiredError(
                "El workbook no pudo volver a guardarse correctamente."
            )
        return {
            "status": "ok",
            "sheet": "Gantt",
            "headers": expected_headers,
            "activity_rows": len(output_activity_rows),
            "source_has_costs": source_has_costs,
            "output_has_costs": output_has_costs,
            "dates_blank": date_cells_blank,
            "file_size": output_path.stat().st_size,
            "reopen_and_save": True,
        }
    finally:
        wb.close()


def build_gantt_workbook(
    input_path: Path,
    output_path: Path,
    source_file_name: str,
    llm_options: LlmOptions | None = None,
) -> BuildResult:
    options = llm_options or LlmOptions(enabled=False)
    mapping_notes: list[str] = []

    try:
        value_wb = load_workbook(input_path, data_only=True, read_only=False)
    except Exception as exc:
        raise GanttReviewRequiredError(
            "ARCHIVO_ROTO_CAUSA_DESCONOCIDA: Excel no pudo abrir el presupuesto "
            f"({type(exc).__name__})."
        ) from exc
    try:
        selected_sheet = select_budget_sheet(value_wb)
        source_ws = value_wb[selected_sheet]
        header_candidate = detect_header_candidate(source_ws)
        samples_by_column = {
            column: [
                source_ws.cell(row_number, column).value
                for row_number in range(
                    header_candidate.row_number + 1,
                    min(source_ws.max_row, header_candidate.row_number + 40) + 1,
                )
                if display_text(source_ws.cell(row_number, column).value)
            ][:8]
            for column in range(1, source_ws.max_column + 1)
        }

        def llm_mapper(
            sheet_name: str,
            headers: list[Any],
            missing_fields: tuple[str, ...],
        ) -> dict[str, int | None]:
            if not options.enabled:
                mapping_notes.append(
                    "Mapping LLM no ejecutado porque OPENAI_ACTIVITY_PLANNER_MODE no está live."
                )
                return {}
            try:
                result = request_llm_column_mapping(
                    workbook_name=source_file_name,
                    sheet_name=sheet_name,
                    headers=headers,
                    missing_fields=missing_fields,
                    samples_by_column=samples_by_column,
                    model=options.model or None,
                )
                mapping_notes.append(
                    "LLM usado solamente para validar o completar el mapping de columnas."
                )
                print(
                    "LLM column validation completed for "
                    f"{sheet_name}: requested={','.join(missing_fields)}"
                )
                return result
            except Exception as exc:
                mapping_notes.append(
                    f"Fallback LLM de columnas falló; no se forzó mapping: {exc}"
                )
                print(
                    "WARNING: LLM column validation failed; deterministic mapping "
                    f"was retained: {exc}"
                )
                return {}

        extraction = extract_budget_structure(
            source_ws,
            source_file=source_file_name,
            source_sheet=selected_sheet,
            column_mapper=llm_mapper,
            enforce_quality_gate=True,
            require_general_budget_label=requires_general_budget_label(
                value_wb,
                selected_sheet,
            ),
        )
        window_start, window_end, notes = read_project_window(value_wb)
    except BudgetExtractionError as exc:
        raise GanttReviewRequiredError(str(exc)) from exc
    finally:
        value_wb.close()

    extraction.assessment.warnings.extend(mapping_notes)
    notes.extend(mapping_notes)
    notes.append(
        "La lógica LLM de clasificación se conserva para otros módulos, "
        "pero no filtra ni reordena las filas del Gantt."
    )
    identity = derive_project_identity(source_file_name)
    has_cc = bool(extraction.column_mapping.get("cc"))
    has_item = bool(extraction.column_mapping.get("item"))
    extra_costs = additional_cost_columns(extraction)
    headers = gantt_headers(has_cc, has_item, extra_costs)
    cost_groups = gantt_cost_groups(extraction)

    wb = load_workbook(input_path, data_only=False, read_only=False, keep_links=True)
    try:
        _link_source_formulas(extraction, wb[selected_sheet])
        ws = _prepare_gantt_sheet(wb)
        calendar_start_column = len(headers) + 1
        display_end = window_end + timedelta(days=CALENDAR_EXTENSION_DAYS)
        contractual_end_column = (
            calendar_start_column + (window_end - window_start).days
        )
        calendar_end_column = _calendar_headers(
            ws,
            window_start,
            window_end,
            display_end,
            calendar_start_column,
        )
        activity_rows, data_end_row = _write_gantt_rows(
            ws,
            wb[selected_sheet],
            extraction,
            headers,
            calendar_start_column,
            contractual_end_column,
            calendar_end_column,
        )
        _style_gantt_sheet(
            ws,
            identity,
            source_file_name,
            headers,
            cost_groups,
            calendar_end_column,
            data_end_row,
        )
        _add_gantt_validations_and_bars(
            ws,
            activity_rows,
            headers,
            calendar_start_column,
            calendar_end_column,
        )
        _metadata_sheet(
            wb,
            identity,
            source_file_name,
            selected_sheet,
            notes,
        )
        _version_baseline_sheet(
            wb,
            ws,
            headers,
            activity_rows,
            window_end,
        )
        _assessment_sheet(wb, extraction)
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
    finally:
        wb.close()

    validation = validate_generated_gantt(output_path, extraction, has_cc)
    extraction.assessment.validation = validation
    validated_wb = load_workbook(output_path, data_only=False, read_only=False)
    try:
        _assessment_sheet(validated_wb, extraction, validation)
        validated_wb.save(output_path)
    finally:
        validated_wb.close()
    validation = validate_generated_gantt(output_path, extraction, has_cc)

    return BuildResult(
        output_path=output_path,
        selected_sheet=selected_sheet,
        header_row=extraction.header_row,
        rows_written=sum(1 for row in extraction.rows if row.row_type == "activity"),
        review_rows=len(extraction.assessment.ambiguous_rows),
        window_start=window_start,
        window_end=window_end,
        notes=notes,
        validation=validation,
    )
