from __future__ import annotations

import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl.utils.cell import range_boundaries


ColumnMapper = Callable[[str, list[Any], tuple[str, ...]], dict[str, int | None]]

FIELDS = (
    "item",
    "actividad",
    "cc",
    "unidad",
    "cantidad",
    "costo_unitario",
    "costo_total",
)

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "item": (
        "item",
        "items",
        "ítem",
        "ítems",
        "número item",
        "numero item",
        "partida",
    ),
    "actividad": (
        "actividad",
        "actividades",
        "material",
        "materiales",
        "descripcion",
        "concepto",
        "detalle",
        "tarea",
        "nombre actividad",
        "rubro",
    ),
    "cc": (
        "cc",
        "c c",
        "centro de costo",
        "centro costo",
        "codigo",
        "cod",
        "id partida",
        "capitulo",
        "familia",
        "partida",
        "rubro",
        "item",
    ),
    "unidad": (
        "unidad",
        "und",
        "unid",
        "u m",
        "um",
        "medida",
        "unidad medida",
    ),
    "cantidad": (
        "cantidad",
        "cant",
        "qty",
        "q",
        "volumen",
    ),
    "costo_unitario": (
        "c unitario",
        "costo unitario",
        "precio unitario",
        "unitario",
        "p u",
        "pu",
        "precio u",
        "costo u",
        "valor unitario",
        "cu",
    ),
    "costo_total": (
        "c total",
        "costo total",
        "precio total",
        "total",
        "importe",
        "subtotal",
        "valor total",
        "monto",
        "total costo",
        "ct",
    ),
}

SUPPORT_FIELDS = ("unidad", "cantidad", "costo_unitario", "costo_total")
TOTAL_LABELS = (
    "total",
    "subtotal",
    "gran total",
    "impuesto",
    "itbms",
    "resumen",
    "utilidad",
    "descuento",
)
NOTE_LABELS = ("nota", "notas", "observacion", "observaciones", "aclaracion")
GENERAL_BUDGET_LABEL = "presupuesto general"


class BudgetExtractionError(RuntimeError):
    def __init__(self, message: str, assessment: "BudgetAssessment") -> None:
        super().__init__(message)
        self.assessment = assessment


def _reject(assessment: "BudgetAssessment", code: str, message: str) -> None:
    assessment.validation = {
        "status": "rejected",
        "code": code,
        "message": message,
    }
    assessment.warnings.append(message)
    raise BudgetExtractionError(f"{code}: {message}", assessment)


@dataclass
class BudgetExtractedRow:
    row_number: int
    row_type: str
    level: int
    actividad: str
    item: Any = None
    cc: Any = None
    unidad: Any = None
    cantidad: Any = None
    costo_unitario: Any = None
    costo_total: Any = None
    extra_costs: dict[str, Any] = field(default_factory=dict)
    source_cells: dict[str, str] = field(default_factory=dict)
    raw_values: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass
class BudgetAssessment:
    detected_columns: dict[str, dict[str, Any]] = field(default_factory=dict)
    missing_columns: list[str] = field(default_factory=list)
    aliases_used: dict[str, str] = field(default_factory=dict)
    ignored_rows: list[dict[str, Any]] = field(default_factory=list)
    ambiguous_rows: list[dict[str, Any]] = field(default_factory=list)
    calculated_cost_rows: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    header_score: float = 0.0
    llm_mapping_used: bool = False
    validation: dict[str, Any] = field(default_factory=dict)
    column_profiles: dict[str, dict[str, Any]] = field(default_factory=dict)
    llm_validation: dict[str, Any] = field(default_factory=dict)


@dataclass
class BudgetCostColumn:
    key: str
    column: int
    header: str
    group: str
    role: str
    output_header: str


@dataclass
class BudgetExtractionResult:
    source_file: str
    source_sheet: str
    header_row: int
    column_mapping: dict[str, int | None]
    cost_columns: list[BudgetCostColumn]
    rows: list[BudgetExtractedRow]
    assessment: BudgetAssessment

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class HeaderCandidate:
    row_number: int
    headers: list[Any]
    mapping: dict[str, int | None]
    scores: dict[str, float]
    score: float


def normalize_text(value: Any) -> str:
    text = str(value or "").replace("\r", " ").replace("\n", " ").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.replace("&", " y ")
    text = re.sub(r"[/\\|_.:;()\[\]{}\-]+", " ", text)
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def display_text(value: Any, max_length: int = 300) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()[:max_length]


def _alias_score(header: Any, aliases: tuple[str, ...]) -> tuple[float, str]:
    normalized = normalize_text(header)
    if not normalized:
        return 0.0, ""
    best_score = 0.0
    best_alias = ""
    header_tokens = set(normalized.split())
    for position, alias in enumerate(aliases):
        alias_normalized = normalize_text(alias)
        alias_tokens = set(alias_normalized.split())
        score = 0.0
        if normalized == alias_normalized:
            score = 120.0 - min(position, 20)
        elif alias_tokens and alias_tokens == header_tokens:
            score = 112.0 - min(position, 20)
        elif len(alias_normalized) >= 3 and re.search(
            rf"(?:^|\s){re.escape(alias_normalized)}(?:$|\s)", normalized
        ):
            score = 92.0 - min(position, 20)
        elif len(alias_tokens) > 1 and alias_tokens.issubset(header_tokens):
            score = 82.0 - min(position, 20)
        if score > best_score:
            best_score = score
            best_alias = alias
    return best_score, best_alias


def resolve_column_mapping(
    headers: list[Any],
) -> tuple[dict[str, int | None], dict[str, float], dict[str, str]]:
    candidates: dict[str, list[tuple[float, int, str]]] = {}
    for field_name in FIELDS:
        options: list[tuple[float, int, str]] = []
        for column_index, header in enumerate(headers, start=1):
            score, alias = _alias_score(header, FIELD_ALIASES[field_name])
            if score:
                options.append((score, column_index, alias))
        candidates[field_name] = sorted(
            options,
            key=lambda option: (-option[0], option[1]),
        )

    mapping = {field_name: None for field_name in FIELDS}
    scores = {field_name: 0.0 for field_name in FIELDS}
    aliases: dict[str, str] = {}
    used_columns: set[int] = set()
    # Activity and commercial values are less replaceable than the optional CC.
    for field_name in (
        "actividad",
        "item",
        "unidad",
        "cantidad",
        "costo_unitario",
        "costo_total",
        "cc",
    ):
        for score, column_index, alias in candidates[field_name]:
            if column_index in used_columns:
                continue
            mapping[field_name] = column_index
            scores[field_name] = score
            aliases[field_name] = alias
            used_columns.add(column_index)
            break
    return mapping, scores, aliases


def _candidate_score(
    mapping: dict[str, int | None],
    scores: dict[str, float],
    headers: list[Any],
) -> float:
    recognized = sum(1 for value in mapping.values() if value)
    activity_bonus = 35 if mapping["actividad"] else 0
    support_bonus = 15 * sum(1 for field_name in SUPPORT_FIELDS if mapping[field_name])
    non_empty = sum(1 for value in headers if display_text(value))
    return recognized * 20 + activity_bonus + support_bonus + sum(scores.values()) / 25 + min(non_empty, 12)


def detect_header_candidate(ws, scan_limit: int = 50) -> HeaderCandidate:
    best = HeaderCandidate(1, [], {field_name: None for field_name in FIELDS}, {}, -1.0)
    max_row = min(max(1, ws.max_row), scan_limit)
    for row_number in range(1, max_row + 1):
        headers = [
            ws.cell(row_number, column_number).value
            for column_number in range(1, ws.max_column + 1)
        ]
        if sum(1 for value in headers if display_text(value)) < 2:
            continue
        mapping, scores, _ = resolve_column_mapping(headers)
        score = _candidate_score(mapping, scores, headers)
        if score > best.score:
            best = HeaderCandidate(row_number, headers, mapping, scores, score)
    return best


def _is_reliable_mapping(mapping: dict[str, int | None]) -> bool:
    return bool(mapping["actividad"]) and sum(
        1 for field_name in SUPPORT_FIELDS if mapping[field_name]
    ) >= 2


def _apply_llm_mapping(
    mapping: dict[str, int | None],
    proposed: dict[str, int | None],
    headers: list[Any],
    profiles_by_column: dict[int, dict[str, Any]] | None = None,
) -> bool:
    changed = False
    used = {value for value in mapping.values() if value}
    for field_name in FIELDS:
        if mapping[field_name] or field_name not in proposed:
            continue
        try:
            column_index = int(proposed[field_name] or 0)
        except (TypeError, ValueError):
            continue
        if (
            column_index < 1
            or column_index > len(headers)
            or column_index in used
            or not display_text(headers[column_index - 1])
            or (
                profiles_by_column is not None
                and not _role_matches_profile(
                    field_name,
                    profiles_by_column[column_index],
                )
            )
        ):
            continue
        mapping[field_name] = column_index
        used.add(column_index)
        changed = True
    return changed


def coerce_number(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    text = display_text(value, 80)
    if not text or text.startswith("="):
        return None
    cleaned = re.sub(r"[^0-9,.\-+()]", "", text)
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = cleaned.strip("()")
    if not re.search(r"\d", cleaned):
        return None
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        tail = cleaned.rsplit(",", 1)[-1]
        cleaned = cleaned.replace(",", ".") if len(tail) <= 2 else cleaned.replace(",", "")
    try:
        result = float(cleaned)
    except ValueError:
        return None
    if negative:
        result *= -1
    return int(result) if result.is_integer() else result


KNOWN_UNITS = {
    "%",
    "dia",
    "dias",
    "gal",
    "gl",
    "global",
    "hr",
    "hrs",
    "kg",
    "lb",
    "m",
    "m2",
    "m3",
    "mes",
    "meses",
    "ml",
    "pie",
    "pies",
    "pza",
    "ud",
    "und",
    "unidad",
    "unidades",
    "viaje",
    "viajes",
}


def _item_code(value: Any) -> bool:
    text = display_text(value, 80).strip().rstrip(".")
    if not text:
        return False
    return bool(re.fullmatch(r"\d{1,3}(?:\.\d{1,3}){1,6}", text))


def _column_profile(ws, column: int, start_row: int, end_row: int) -> dict[str, Any]:
    values = [
        ws.cell(row_number, column).value
        for row_number in range(start_row, end_row + 1)
        if display_text(ws.cell(row_number, column).value)
    ][:40]
    total = len(values)
    numeric = sum(coerce_number(value) is not None for value in values)
    item_codes = sum(_item_code(value) for value in values)
    text_values = sum(
        bool(re.search(r"[A-Za-zÁÉÍÓÚáéíóúÑñ]{3,}", display_text(value, 200)))
        and not _item_code(value)
        for value in values
    )
    units = sum(normalize_text(value) in KNOWN_UNITS for value in values)
    return {
        "sample_count": total,
        "numeric_ratio": numeric / total if total else 0.0,
        "item_ratio": item_codes / total if total else 0.0,
        "text_ratio": text_values / total if total else 0.0,
        "unit_ratio": units / total if total else 0.0,
        "samples": [display_text(value, 80) for value in values[:8]],
    }


def _role_matches_profile(field_name: str, profile: dict[str, Any]) -> bool:
    if not profile["sample_count"]:
        return True
    if field_name == "item":
        return profile["item_ratio"] >= 0.45
    if field_name == "actividad":
        return profile["text_ratio"] >= 0.45 and profile["item_ratio"] < 0.45
    if field_name == "unidad":
        return profile["unit_ratio"] >= 0.35
    if field_name in {"cantidad", "costo_unitario", "costo_total"}:
        return profile["numeric_ratio"] >= 0.35
    if field_name == "cc":
        samples = profile.get("samples") or []
        return any("cc" in normalize_text(value).split() for value in samples) or (
            profile["text_ratio"] >= 0.35
            or profile["item_ratio"] >= 0.35
        )
    return True


def _contains_general_budget_label(ws, scan_limit: int = 60) -> bool:
    max_row = min(ws.max_row, scan_limit)
    for row_number in range(1, max_row + 1):
        for column in range(1, ws.max_column + 1):
            if GENERAL_BUDGET_LABEL in normalize_text(
                ws.cell(row_number, column).value
            ):
                return True
    return False


def _validate_semantic_mapping(
    mapping: dict[str, int | None],
    profiles_by_column: dict[int, dict[str, Any]],
    cost_columns: list["BudgetCostColumn"],
    assessment: "BudgetAssessment",
) -> None:
    semantic_fields = ("item", "actividad", "unidad", "cantidad")
    mapped = [
        (field_name, mapping.get(field_name))
        for field_name in semantic_fields
        if mapping.get(field_name)
    ]
    by_column: dict[int, list[str]] = {}
    for field_name, column in mapped:
        by_column.setdefault(int(column), []).append(field_name)
    collisions = {
        column: fields for column, fields in by_column.items() if len(fields) > 1
    }
    if collisions:
        _reject(
            assessment,
            "ENCABEZADOS_DESORDENADOS",
            f"Una misma columna fue interpretada con roles incompatibles: {collisions}.",
        )

    requirements = {
        "actividad": ("text_ratio", 0.45, "texto de materiales/actividades"),
        "cantidad": ("numeric_ratio", 0.45, "cantidades numéricas"),
        "unidad": ("unit_ratio", 0.30, "unidades reconocibles"),
    }
    if mapping.get("item"):
        requirements["item"] = ("item_ratio", 0.40, "códigos jerárquicos como 1.1")

    failures: list[str] = []
    for field_name, (metric, minimum, expected) in requirements.items():
        column = mapping.get(field_name)
        if not column:
            failures.append(f"{field_name}: columna ausente")
            continue
        profile = profiles_by_column[int(column)]
        if (
            profile["sample_count"] >= 2
            and float(profile.get(metric) or 0.0) < minimum
        ):
            failures.append(
                f"{field_name}: columna {column} no contiene principalmente {expected} "
                f"({metric}={float(profile.get(metric) or 0.0):.0%})"
            )
    if failures:
        _reject(
            assessment,
            "CONTENIDO_DE_COLUMNAS_INVALIDO",
            "; ".join(failures),
        )

    if not cost_columns:
        _reject(
            assessment,
            "COLUMNAS_DE_COSTO_AUSENTES",
            "No se encontró ninguna columna confiable de costo o precio.",
        )
    invalid_costs = []
    for cost_column in cost_columns:
        profile = profiles_by_column[cost_column.column]
        if (
            profile["sample_count"] >= 2
            and float(profile.get("numeric_ratio") or 0.0) < 0.35
        ):
            invalid_costs.append(
                f"{cost_column.header} (columna {cost_column.column}, "
                f"numeric_ratio={float(profile.get('numeric_ratio') or 0.0):.0%})"
            )
    if invalid_costs:
        _reject(
            assessment,
            "CONTENIDO_DE_COSTOS_INVALIDO",
            "Estas columnas financieras no contienen principalmente números: "
            + "; ".join(invalid_costs),
        )


def _financial_role(header: Any) -> str:
    text = normalize_text(header)
    if not text:
        return ""
    has_cost = "costo" in text.split() or text.startswith("c total")
    has_price = (
        "precio" in text.split()
        or "presio" in text.split()
        or text in {"p u", "pu"}
    )
    if "margen" in text:
        return "margen"
    if "utilidad" in text:
        return "utilidad_total"
    if has_price and ("unitario" in text or text in {"p u", "pu"}):
        return "precio_unitario"
    if has_price and "total" in text:
        return "precio_total"
    if has_cost and "unitario" in text:
        return "costo_unitario"
    if has_cost and "total" in text:
        return "costo_total"
    return ""


def _column_group_label(ws, header_row: int, column: int) -> str:
    candidates: list[tuple[int, str]] = []
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_col <= column <= merged.max_col
            and merged.max_row < header_row
            and header_row - merged.max_row <= 4
        ):
            label = display_text(ws.cell(merged.min_row, merged.min_col).value, 120)
            if label:
                candidates.append((merged.max_row, label))
    if candidates:
        return max(candidates, key=lambda item: item[0])[1]
    for row_number in range(header_row - 1, max(0, header_row - 4), -1):
        label = display_text(ws.cell(row_number, column).value, 120)
        if label:
            return label
    return ""


def detect_cost_columns(
    ws,
    header_row: int,
    headers: list[Any],
    mapping: dict[str, int | None],
) -> list[BudgetCostColumn]:
    columns: list[BudgetCostColumn] = []
    used_output_headers: set[str] = set()
    for column, header_value in enumerate(headers, start=1):
        if column == mapping.get("costo_unitario"):
            role = "costo_unitario"
        elif column == mapping.get("costo_total"):
            role = "costo_total"
        else:
            role = _financial_role(header_value)
        if not role:
            continue
        header = display_text(header_value, 120)
        group = _column_group_label(ws, header_row, column)
        if column == mapping.get("costo_unitario"):
            output_header = "Costo Unitario"
        elif column == mapping.get("costo_total"):
            output_header = "Costo Total"
        else:
            output_header = header
            if group:
                output_header = f"{header} — {group}"
        base_output_header = output_header
        suffix = 2
        while normalize_text(output_header) in used_output_headers:
            output_header = f"{base_output_header} ({suffix})"
            suffix += 1
        used_output_headers.add(normalize_text(output_header))
        columns.append(
            BudgetCostColumn(
                key=f"cost_{column}",
                column=column,
                header=header,
                group=group,
                role=role,
                output_header=output_header,
            )
        )
    return columns


def _value(ws, row_number: int, mapping: dict[str, int | None], field_name: str) -> Any:
    column_index = mapping.get(field_name)
    return ws.cell(row_number, column_index).value if column_index else None


def _row_raw_values(ws, row_number: int, headers: list[Any]) -> dict[str, Any]:
    raw: dict[str, Any] = {}
    for column_number in range(1, ws.max_column + 1):
        value = ws.cell(row_number, column_number).value
        if value is None:
            continue
        header = display_text(headers[column_number - 1], 80) or f"Columna {column_number}"
        raw[f"{column_number}:{header}"] = value
    return raw


def _row_level(
    ws,
    row_number: int,
    item: Any,
    cc: Any,
    activity_column: int,
) -> int:
    item_text = display_text(item, 80).strip(" .-")
    if item_text:
        parts = [part for part in item_text.split(".") if part.strip()]
        if parts:
            return min(max(len(parts) - 1, 0), 6)
    cc_text = display_text(cc, 80).strip(" .-")
    if cc_text:
        parts = [part for part in re.split(r"[.\-/]+", cc_text) if part.strip()]
        if len(parts) > 1:
            return min(len(parts) - 1, 6)
    indent = int(ws.cell(row_number, activity_column).alignment.indent or 0)
    return min(max(indent, 0), 6)


def _financial_value(value: Any) -> Any:
    number = coerce_number(value)
    if number is not None:
        return number
    text = display_text(value, 120)
    if text and not text.startswith("="):
        return text
    return None


def _row_is_merged(ws, row_number: int) -> bool:
    return any(cell_range.min_row <= row_number <= cell_range.max_row for cell_range in ws.merged_cells.ranges)


def _total_or_note_kind(activity: str) -> str:
    text = normalize_text(activity)
    if any(re.match(rf"^{re.escape(term)}(?:\s|$)", text) for term in TOTAL_LABELS):
        return "total"
    if re.match(
        r"^(?:costo|precio)\s*(?:m2|m3|kg|lb|ml|m|unidad|und|u|gal|hora|dia|mes)$",
        text,
    ):
        return "cost_summary"
    if any(re.match(rf"^{re.escape(term)}(?:\s|$)", text) for term in NOTE_LABELS):
        return "note"
    return ""


def _explicit_table_end_row(ws, header_row: int) -> int | None:
    references: list[str] = []
    if ws.auto_filter.ref:
        references.append(str(ws.auto_filter.ref))
    references.extend(str(table.ref) for table in ws.tables.values())
    matching_ends: list[int] = []
    for reference in references:
        try:
            _, min_row, _, max_row = range_boundaries(reference)
        except (TypeError, ValueError):
            continue
        if min_row <= header_row <= max_row:
            matching_ends.append(max_row)
    return max(matching_ends) if matching_ends else None


def _numeric_only_label(value: str) -> bool:
    return coerce_number(value) is not None and not re.search(
        r"[A-Za-zÁÉÍÓÚáéíóúÑñ]",
        value,
    )


def extract_budget_structure(
    ws,
    *,
    source_file: str | Path,
    source_sheet: str | None = None,
    scan_limit: int = 50,
    column_mapper: ColumnMapper | None = None,
    enforce_quality_gate: bool = False,
) -> BudgetExtractionResult:
    candidate = detect_header_candidate(ws, scan_limit=scan_limit)
    mapping = dict(candidate.mapping)
    assessment = BudgetAssessment(header_score=candidate.score)
    if enforce_quality_gate and not _contains_general_budget_label(ws):
        _reject(
            assessment,
            "SIN_PRESUPUESTO_GENERAL",
            "No se encontró un bloque identificado como Presupuesto General; "
            "el archivo parece dividido o incompleto.",
        )
    if enforce_quality_gate:
        for row_number in range(candidate.row_number + 1, ws.max_row + 1):
            values = [
                ws.cell(row_number, column).value
                for column in range(1, ws.max_column + 1)
            ]
            repeated_mapping, _, _ = resolve_column_mapping(values)
            if _is_reliable_mapping(repeated_mapping):
                _reject(
                    assessment,
                    "ENCABEZADOS_REPETIDOS",
                    f"Se encontró otra tabla de presupuesto en la fila {row_number}; "
                    "el archivo parece dividido o tiene encabezados desordenados.",
                )
    profile_end_row = min(
        ws.max_row,
        _explicit_table_end_row(ws, candidate.row_number)
        or candidate.row_number + 40,
    )
    profiles_by_column = {
        column: _column_profile(
            ws,
            column,
            candidate.row_number + 1,
            profile_end_row,
        )
        for column in range(1, len(candidate.headers) + 1)
    }
    for field_name, column_index in list(mapping.items()):
        if not column_index:
            continue
        profile = profiles_by_column[column_index]
        assessment.column_profiles[field_name] = {
            "column": column_index,
            **profile,
        }
        if not _role_matches_profile(field_name, profile):
            if (
                field_name == "item"
                and not mapping.get("actividad")
                and profile["text_ratio"] >= 0.45
            ):
                mapping["actividad"] = column_index
                mapping["item"] = None
                assessment.warnings.append(
                    f"La columna {column_index} parecía Ítem por encabezado, pero su "
                    "contenido es textual; se reutilizó como Actividad."
                )
                continue
            assessment.warnings.append(
                f"La columna {column_index} no pasó la validación de contenido para "
                f"{field_name}; se solicitó fallback controlado."
            )
            mapping[field_name] = None

    preliminary_cost_columns = detect_cost_columns(
        ws,
        candidate.row_number,
        candidate.headers,
        mapping,
    )
    missing = tuple(field_name for field_name in FIELDS if not mapping[field_name])
    complex_mapping = bool(mapping.get("item")) or len(preliminary_cost_columns) > 2
    requested_fields = tuple(FIELDS) if complex_mapping else missing
    if requested_fields and column_mapper:
        proposed = column_mapper(
            source_sheet or ws.title,
            candidate.headers,
            requested_fields,
        )
        assessment.llm_mapping_used = bool(proposed)
        changed = _apply_llm_mapping(
            mapping,
            proposed,
            candidate.headers,
            profiles_by_column,
        )
        assessment.llm_validation = {
            "requested_fields": list(requested_fields),
            "proposed_mapping": proposed,
            "mapping_changed": changed,
            "mode": "complex_validation" if complex_mapping else "missing_fields",
        }

    if enforce_quality_gate:
        _validate_semantic_mapping(
            mapping,
            profiles_by_column,
            detect_cost_columns(
                ws,
                candidate.row_number,
                candidate.headers,
                mapping,
            ),
            assessment,
        )

    _, scores, aliases = resolve_column_mapping(candidate.headers)
    if not _is_reliable_mapping(mapping):
        assessment.missing_columns = [
            field_name for field_name in FIELDS if not mapping[field_name]
        ]
        assessment.warnings.append(
            "No se identificó un encabezado confiable: se requiere Actividad y al menos "
            "dos columnas entre Unidad, Cantidad, Costo Unitario y Costo Total."
        )
        raise BudgetExtractionError(assessment.warnings[-1], assessment)

    for field_name, column_index in mapping.items():
        if not column_index:
            continue
        header = display_text(candidate.headers[column_index - 1], 120)
        assessment.detected_columns[field_name] = {
            "column": column_index,
            "header": header,
            "score": scores.get(field_name, 0.0),
        }
        assessment.aliases_used[field_name] = aliases.get(field_name, header)
    assessment.missing_columns = [
        field_name for field_name in FIELDS if not mapping[field_name]
    ]
    cost_columns = detect_cost_columns(
        ws,
        candidate.row_number,
        candidate.headers,
        mapping,
    )

    rows: list[BudgetExtractedRow] = []
    current_mapping = dict(mapping)
    current_headers = list(candidate.headers)
    explicit_end_row = _explicit_table_end_row(ws, candidate.row_number)
    scan_end_row = explicit_end_row or ws.max_row
    consecutive_non_table_rows = 0
    for row_number in range(candidate.row_number + 1, scan_end_row + 1):
        values = [
            ws.cell(row_number, column_number).value
            for column_number in range(1, ws.max_column + 1)
        ]
        if not any(display_text(value) for value in values):
            assessment.ignored_rows.append({"row_number": row_number, "reason": "empty"})
            rows.append(
                BudgetExtractedRow(
                    row_number=row_number,
                    row_type="spacer",
                    level=0,
                    actividad="",
                )
            )
            consecutive_non_table_rows += 1
            if explicit_end_row is None and consecutive_non_table_rows >= 8:
                break
            continue

        row_mapping, _, _ = resolve_column_mapping(values)
        if _is_reliable_mapping(row_mapping):
            if enforce_quality_gate:
                _reject(
                    assessment,
                    "ENCABEZADOS_REPETIDOS",
                    f"Se encontró otra tabla de presupuesto en la fila {row_number}; "
                    "el archivo parece dividido o tiene encabezados desordenados.",
                )
            current_mapping = row_mapping
            current_headers = list(values)
            assessment.ignored_rows.append(
                {
                    "row_number": row_number,
                    "reason": "repeated_header",
                    "mapping": dict(row_mapping),
                }
            )
            consecutive_non_table_rows = 0
            continue

        activity_column = int(current_mapping["actividad"] or 0)
        activity = display_text(
            _value(ws, row_number, current_mapping, "actividad"),
            500,
        )
        if not activity:
            assessment.ignored_rows.append(
                {"row_number": row_number, "reason": "missing_activity"}
            )
            consecutive_non_table_rows += 1
            if explicit_end_row is None and consecutive_non_table_rows >= 8:
                break
            continue
        if _numeric_only_label(activity):
            assessment.ignored_rows.append(
                {
                    "row_number": row_number,
                    "reason": "numeric_activity",
                    "activity": activity,
                }
            )
            consecutive_non_table_rows += 1
            if explicit_end_row is None and consecutive_non_table_rows >= 8:
                break
            continue

        excluded_kind = _total_or_note_kind(activity)
        if excluded_kind:
            assessment.ignored_rows.append(
                {
                    "row_number": row_number,
                    "reason": excluded_kind,
                    "activity": activity,
                }
            )
            consecutive_non_table_rows += 1
            continue

        item = _value(ws, row_number, current_mapping, "item")
        cc = _value(ws, row_number, current_mapping, "cc")
        unit = _value(ws, row_number, current_mapping, "unidad")
        quantity = coerce_number(
            _value(ws, row_number, current_mapping, "cantidad")
        )
        unit_cost = coerce_number(
            _value(ws, row_number, current_mapping, "costo_unitario")
        )
        total_cost = coerce_number(
            _value(ws, row_number, current_mapping, "costo_total")
        )
        extra_costs = {
            cost_column.key: _financial_value(
                ws.cell(row_number, cost_column.column).value
            )
            for cost_column in cost_columns
            if cost_column.column
            not in {
                current_mapping.get("costo_unitario"),
                current_mapping.get("costo_total"),
            }
        }
        warnings: list[str] = []
        if total_cost is None and quantity is not None and unit_cost is not None:
            total_cost = quantity * unit_cost
            assessment.calculated_cost_rows.append(row_number)
            warnings.append("Costo Total calculado como Cantidad × Costo Unitario.")
        elif total_cost is not None and quantity is not None and unit_cost is not None:
            calculated = quantity * unit_cost
            tolerance = max(1.0, abs(total_cost) * 0.02)
            if abs(total_cost - calculated) > tolerance:
                warnings.append(
                    "Costo Total fuente difiere más de 2% de Cantidad × Costo Unitario; "
                    "se conservó el valor fuente."
                )

        has_commercial_value = any(
            display_text(value)
            for value in (
                unit,
                _value(ws, row_number, current_mapping, "cantidad"),
                _value(ws, row_number, current_mapping, "costo_unitario"),
                _value(ws, row_number, current_mapping, "costo_total"),
                *extra_costs.values(),
            )
        )
        row_type = "activity" if has_commercial_value else "section"
        if not has_commercial_value and not (_row_is_merged(ws, row_number) or cc):
            assessment.ambiguous_rows.append(
                {
                    "row_number": row_number,
                    "activity": activity,
                    "reason": "Fila sin unidad, cantidad ni costos; se preservó como sección.",
                }
            )
            warnings.append("Fila preservada como sección por falta de valores comerciales.")

        rows.append(
            BudgetExtractedRow(
                row_number=row_number,
                row_type=row_type,
                level=_row_level(ws, row_number, item, cc, activity_column),
                actividad=activity,
                item=item,
                cc=cc,
                unidad=unit,
                cantidad=quantity,
                costo_unitario=unit_cost,
                costo_total=total_cost,
                extra_costs=extra_costs,
                source_cells={
                    field_name: ws.cell(row_number, column_index).coordinate
                    for field_name, column_index in current_mapping.items()
                    if column_index
                }
                | {
                    cost_column.key: ws.cell(
                        row_number,
                        cost_column.column,
                    ).coordinate
                    for cost_column in cost_columns
                },
                raw_values=_row_raw_values(ws, row_number, current_headers),
                warnings=warnings,
            )
        )
        consecutive_non_table_rows = 0

    while rows and rows[-1].row_type == "spacer":
        rows.pop()

    if not any(row.row_type == "activity" for row in rows):
        assessment.warnings.append(
            "El presupuesto no produjo filas de actividad con valores comerciales."
        )
        raise BudgetExtractionError(assessment.warnings[-1], assessment)

    if not mapping["cc"]:
        assessment.warnings.append(
            "No se detectó una columna CC independiente; el Gantt omitirá CC."
        )
    return BudgetExtractionResult(
        source_file=str(source_file),
        source_sheet=source_sheet or ws.title,
        header_row=candidate.row_number,
        column_mapping=mapping,
        cost_columns=cost_columns,
        rows=rows,
        assessment=assessment,
    )
